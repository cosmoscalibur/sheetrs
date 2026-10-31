#!/usr/bin/env python3
"""Generate tests/assets/semantics.{xlsx,ods}.

The workbook holds one formula per row of the `Calc` sheet over fixed data in
`Data`. openpyxl writes the formulas without values; LibreOffice then
recalculates and stores them (XLSX and ODS), so the cached results are an
independent oracle for the engine (see tests/semantics/oracle.rs).

Also generates tests/assets/loan.{xlsx,ods}, the example model used by the
CLI tests and the README.

Usage: scripts/gen_semantics_fixture.py [TAG ...]   (needs openpyxl and soffice; with tags, only the
fixtures of those locales, see LOCALE_FIXTURES)
"""
import os
import re
import subprocess
import tempfile
from datetime import date

import openpyxl

SOFFICE = os.environ.get("SOFFICE", "soffice")  # the reference build: LibreOffice 26.8
OUT = os.path.join(os.path.dirname(__file__), "..", "tests", "assets")

FORMULAS = [
    # aggregates
    "SUM(Data!A1:A8)", "AVERAGE(Data!A1:A8)", "MIN(Data!A1:A8)", "MAX(Data!A1:A8)",
    "COUNT(Data!A1:A8)", "COUNTA(Data!B1:B9)", "MEDIAN(Data!A1:A8)",
    "LARGE(Data!A1:A8,2)", "SMALL(Data!A1:A8,3)", "PRODUCT(Data!A1:A3)",
    "SUMIF(Data!B1:B8,\"apple\",Data!A1:A8)", "COUNTIF(Data!B1:B8,\"a*\")",
    "COUNTIF(Data!A1:A8,\">4\")", "COUNTIF(Data!A1:A8,\"<>5\")",
    "SUMIFS(Data!A1:A8,Data!B1:B8,\"apple\",Data!A1:A8,\">=5\")",
    "AVERAGEIF(Data!B1:B8,\"banana\",Data!A1:A8)",
    "COUNTIFS(Data!B1:B8,\"apple\",Data!A1:A8,\"<6\")",
    "_xlfn.MAXIFS(Data!A1:A8,Data!B1:B8,\"apple\")",
    "_xlfn.MINIFS(Data!A1:A8,Data!B1:B8,\"apple\")",
    "SUMPRODUCT(Data!A1:A8,Data!A1:A8)",
    "SUMPRODUCT((Data!B1:B8=\"apple\")*(Data!A1:A8))",
    "COUNTBLANK(Data!B1:B9)", "SUM(Data!A:A)", "SUMSQ(Data!A1:A3)",
    "STDEV(Data!A1:A8)", "VAR(Data!A1:A8)",
    # math
    "ROUND(Data!A1/3,2)", "ROUND(2.5,0)", "ROUND(-2.5,0)", "ROUND(1.005,2)", "ROUND(1234.567,-2)",
    "ROUNDUP(2.341,2)", "ROUNDDOWN(-2.349,2)", "TRUNC(8.97)", "TRUNC(8.97,1)",
    "INT(-3.2)", "MOD(-7,3)", "MOD(7,-3)", "POWER(2,10)", "SQRT(16)", "ABS(-4)", "SIGN(-3)",
    "EXP(1)", "LN(10)", "LOG(100,10)", "LOG(8,2)", "LOG10(1000)", "CEILING(4.3,0.5)",
    "FLOOR(4.3,0.5)", "PI()", "EVEN(3)", "ODD(2)", "EVEN(-3)", "FACT(5)", "DEGREES(PI())",
    "SIN(PI()/2)", "ATAN2(1,1)",
    # logic and info
    "IF(Data!A1>3,\"big\",\"small\")", "IF(Data!A2>3,\"big\",\"small\")", "IFERROR(1/0,\"err\")",
    "_xlfn.IFNA(NA(),\"na\")", "AND(TRUE,1,Data!A1>1)", "OR(FALSE,0,Data!A1<1)", "NOT(Data!A1>1)",
    "_xlfn.XOR(TRUE,TRUE)", "CHOOSE(2,\"a\",\"b\",\"c\")", "_xlfn.IFS(Data!A1<1,\"x\",Data!A1>1,\"y\")",
    "_xlfn.SWITCH(Data!A1,5,\"five\",3,\"three\",\"other\")",
    "ISNUMBER(Data!A1)", "ISTEXT(Data!B1)", "ISBLANK(Data!B9)", "ISERROR(1/0)", "ISNA(NA())",
    "ISEVEN(4)", "ISODD(4)", "ISLOGICAL(TRUE)", "N(TRUE)",
    # dynamic-array era (LibreOffice 24.8 and later)
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.UNIQUE(Data!B1:B8))", "COUNTA(_xlfn.UNIQUE(Data!B1:B8))",
    "SUM(_xlfn.UNIQUE(Data!A1:A8))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.UNIQUE(Data!B1:B8,FALSE,TRUE))",
    "_xlfn.LET(_xlpm.x,2,_xlpm.y,_xlpm.x*3,_xlpm.x+_xlpm.y)", "_xlfn.LET(_xlpm.x,Data!A1,_xlpm.x*2)",
    "_xlfn.LET(_xlpm.x,1,_xlfn.LET(_xlpm.x,_xlpm.x+10,_xlpm.x*2))",
    # currency text: the locale's symbol and negative pattern
    "DOLLAR(1234.567,2)", "DOLLAR(-1234.567,2)", "DOLLAR(1234.567)", "DOLLAR(1234.567,0)", "DOLLAR(1234.567,-2)",
    "DOLLAR(0.5,1)", "DOLLAR(0,2)", "DOLLAR(1E6,0)", "DOLLAR(\"12\",1)", "DOLLAR(TRUE)", "DOLLAR(-0.004,2)",
    "DOLLAR(99.995,2)", "DOLLAR(1234.5678,3)",
    # what a cell holds: known when the workbook is read
    "_xlfn.ISFORMULA(A1)", "_xlfn.ISFORMULA(Z999)",
    "_xlfn.FORMULATEXT(A1)", "_xlfn.FORMULATEXT(A2)", "_xlfn.FORMULATEXT(A12)",
    # text
    "LEN(Data!B2)", "UPPER(Data!B2)", "LOWER(\"ABC\")", "PROPER(\"hELLO wORLD\")", "TRIM(\"  a   b  \")",
    "LEFT(Data!B2,3)", "RIGHT(Data!B2,2)", "MID(Data!B2,2,3)", "FIND(\"n\",Data!B2)",
    "SEARCH(\"AN\",Data!B2)", "SEARCH(\"b*a\",Data!B2)", "SUBSTITUTE(\"a-b-c\",\"-\",\"+\")",
    "SUBSTITUTE(\"a-b-c\",\"-\",\"+\",2)", "REPLACE(\"abcdef\",2,3,\"X\")", "REPT(\"ab\",3)",
    "EXACT(\"a\",\"A\")", "CONCATENATE(\"a\",1,\"b\")", "_xlfn.CONCAT(\"a\",Data!A1)",
    "_xlfn.TEXTJOIN(\"-\",TRUE,\"a\",\"\",\"b\")", "VALUE(\"12.5\")", "VALUE(\"50%\")",
    "TEXT(1234.567,\"#,##0.00\")", "TEXT(0.256,\"0.0%\")", "\"[\"&TEXT(5,\"000\")&\"]\"",
    "TEXT(Data!C1,\"yyyy-mm-dd\")", "TEXT(Data!C1,\"dddd\")", "TEXT(Data!C1,\"mmm d, yyyy\")",
    "TEXT(-3.5,\"0.0;(0.0)\")", "CHAR(65)", "CODE(\"A\")", "FIXED(1234.567,1)",
    "\"[\"&(0.1+0.2)&\"]\"", "\"[\"&(1/3)&\"]\"", "\"[\"&100000000000000000000&\"]\"",
    "\"[\"&0.00001&\"]\"", "\"[\"&1234567.891&\"]\"", "\"[\"&-0.5&\"]\"",
    # lookup
    "VLOOKUP(\"cherry\",Data!B1:C8,2,FALSE)", "VLOOKUP(\"zzz\",Data!B1:C8,2,FALSE)",
    "VLOOKUP(4,Data!A1:C8,3,TRUE)", "HLOOKUP(\"y\",Data!E1:G2,2,FALSE)",
    # an empty cell as the searched value finds nothing in LibreOffice
    "VLOOKUP(Data!B9,Data!A1:C8,2,FALSE)", "MATCH(Data!B9,Data!A1:A8,0)", "MATCH(Data!B9,Data!A1:A8,1)",
    "HLOOKUP(Data!B9,Data!E1:G2,2,FALSE)", "LOOKUP(Data!B9,Data!A1:A8,Data!A1:A8)",
    "COUNTIF(Data!A1:A8,Data!B9)", "MATCH(\"\",Data!B1:B9,0)",
    "SEARCH(\"\",\"abc\")", "FIND(\"\",\"abc\")", "SEARCH(Data!B9,\"abc\")",
    "MATCH(\"cherry\",Data!B1:B8,0)", "MATCH(\"ch*\",Data!B1:B8,0)", "MATCH(6,Data!H1:H5,1)",
    "MATCH(6,Data!I1:I5,-1)", "INDEX(Data!A1:A8,3)", "INDEX(Data!A1:C8,2,1)",
    "SUM(INDEX(Data!A1:C8,0,1))", "SUM(INDEX(Data!A1:C8,2,0))",
    "_xlfn.XLOOKUP(\"cherry\",Data!B1:B8,Data!A1:A8)", "_xlfn.XLOOKUP(\"q\",Data!B1:B8,Data!A1:A8,\"none\")",
    "ROW(Data!A5)", "COLUMN(Data!C1)", "ROWS(Data!A1:C8)", "COLUMNS(Data!A1:C8)",
    "ADDRESS(2,3)", "ADDRESS(2,3,4)", "OFFSET(Data!A1,2,0)", "SUM(OFFSET(Data!A1,0,0,3,1))",
    "INDIRECT(\"Data!A3\")", "SUM(INDIRECT(\"Data!A1:A3\"))", "INDIRECT(\"Data!B\"&2)",
    # dates
    "DATE(2024,1,31)", "DATE(2024,14,1)", "DATE(2024,3,0)", "DATE(1900,3,1)", "DATE(2024,2,29)+1",
    "YEAR(Data!C1)", "MONTH(Data!C1)", "DAY(Data!C1)", "WEEKDAY(Data!C1)", "WEEKDAY(Data!C1,2)",
    "EDATE(DATE(2024,1,31),1)", "EDATE(DATE(2024,3,31),-1)", "EOMONTH(Data!C1,0)", "EOMONTH(Data!C1,2)",
    "_xlfn.DAYS(Data!C2,Data!C1)", "DATEDIF(Data!C1,Data!C2,\"M\")", "DATEDIF(Data!C1,Data!C2,\"D\")",
    "DATEDIF(DATE(2020,1,15),DATE(2024,1,14),\"Y\")", "HOUR(0.75)", "MINUTE(0.7512)", "SECOND(0.7512)",
    "TIME(13,30,0)", "DATEVALUE(\"2024-03-01\")",
    # finance
    "PMT(0.01,12,-1000)", "PMT(0.05/12,360,200000)", "FV(0.05,10,-100)", "PV(0.05,10,-100)",
    "NPV(0.1,100,200,300)", "IRR(Data!J1:J5)",
    # operators and coercion
    "-2^2", "2^-1", "2^3^2", "10%*3", "\"3\"+1", "TRUE+1", "\"a\"&\"b\"", "1=\"1\"", "\"a\"<\"B\"",
    "\"abc\"=\"ABC\"", "1<\"a\"", "(1+2)*3", "7/2", "-(-3)",
    "Data!A1:A3*2", "SUM(Data!A1:A3*2)", "SUM(Data!A1:A3>3)",
    # errors
    "1/0", "SQRT(-1)", "NA()", "\"a\"+1", "0^0", "LN(0)", "VALUE(\"abc\")", "MATCH(99,Data!A1:A8,0)",
    "INDEX(Data!A1:A8,99)", "SUMPRODUCT(LEN(Data!B1:B3))", "SUMPRODUCT(--(Data!A1:A8>4))", "Data!A1+Data!B1", "#REF!+1", "ISERROR(Data!A1/0)",
]


# A formula that ends in a blank cell: rows below refer to the one at row `n + 1`.
_n = len(FORMULAS)
FORMULAS += [
    "Data!Z99", f"ISBLANK(A{_n + 1})", f"A{_n + 1}=\"\"", f"A{_n + 1}=0",
    "Data!Z99&\"x\"", "IF(TRUE,Data!Z99)", "ISBLANK(Data!Z99)", "Data!Z99=0", "Data!Z99=\"\"",
    "SUM(Data!Z99,1)", "LEN(Data!Z99)", "ISNUMBER(A%d)" % (_n + 1),
]


FORMULAS += [
    # math (batch 2)
    "ASIN(0.5)", "ACOS(0.5)", "SINH(1)", "COSH(1)", "TANH(0.5)", "ASINH(1)", "ACOSH(2)", "ATANH(0.5)",
    "_xlfn.COT(1)", "_xlfn.COTH(1)", "_xlfn.ACOT(2)", "_xlfn.ACOTH(2)", "_xlfn.SEC(1)", "_xlfn.CSC(1)",
    "_xlfn.SECH(1)", "_xlfn.CSCH(1)", "SQRTPI(2)", "ASIN(2)", "ACOSH(0.5)",
    "GCD(12,18)", "GCD(Data!A1:A3)", "LCM(4,6)", "LCM(Data!A1:A3)", "GCD(-1,2)", "QUOTIENT(7,2)", "QUOTIENT(-7,2)",
    "QUOTIENT(1,0)", "MROUND(10,3)", "MROUND(-10,-3)", "MROUND(5,-2)", "MROUND(1.3,0.2)",
    "_xlfn.CEILING.MATH(4.3)", "_xlfn.CEILING.MATH(-4.3)", "_xlfn.CEILING.MATH(-4.3,1,1)",
    "_xlfn.CEILING.MATH(4.3,0.5)", "_xlfn.FLOOR.MATH(4.3)", "_xlfn.FLOOR.MATH(-4.3)",
    "_xlfn.FLOOR.MATH(-4.3,1,1)", "_xlfn.FLOOR.MATH(4.3,0.5)", "_xlfn.CEILING.PRECISE(4.3,-2)",
    "_xlfn.FLOOR.PRECISE(-4.3,2)", "ISO.CEILING(4.3,2)",
    "COMBIN(5,2)", "COMBIN(5,6)", "_xlfn.COMBINA(3,2)", "PERMUT(5,2)", "_xlfn.PERMUTATIONA(3,2)",
    "FACTDOUBLE(7)", "FACTDOUBLE(8)", "MULTINOMIAL(2,3,4)",
    "_xlfn.BITAND(12,10)", "_xlfn.BITOR(12,10)", "_xlfn.BITXOR(12,10)", "_xlfn.BITLSHIFT(4,2)",
    "_xlfn.BITRSHIFT(16,2)", "ERF(1)", "ERF(0.5,1)", "_xlfn.ERF.PRECISE(1)", "ERFC(1)",
    "_xlfn.ERFC.PRECISE(1)", "ERFC(3)", "ERF(-0.5)",
    # statistics (batch 2)
    "_xlfn.STDEV.S(Data!A1:A8)", "_xlfn.STDEV.P(Data!A1:A8)", "STDEVP(Data!A1:A8)",
    "_xlfn.VAR.S(Data!A1:A8)", "_xlfn.VAR.P(Data!A1:A8)", "VARP(Data!A1:A8)",
    "STDEVA(Data!A1:B3)", "STDEVPA(Data!A1:B3)", "VARA(Data!A1:B3)", "VARPA(Data!A1:B3)",
    "AVERAGEA(Data!A1:B3)", "MAXA(Data!A1:B3)", "MINA(Data!A1:B3)",
    "DEVSQ(Data!A1:A8)", "AVEDEV(Data!A1:A8)", "GEOMEAN(Data!A1:A8)", "HARMEAN(Data!A1:A8)",
    "SKEW(Data!A1:A8)", "_xlfn.SKEW.P(Data!A1:A8)", "KURT(Data!A1:A8)", "TRIMMEAN(Data!A1:A8,0.25)",
    "MODE(Data!K1:K8)", "_xlfn.MODE.SNGL(Data!K1:K8)", "MODE(Data!A1:A8)",
    "RANK(8,Data!A1:A8)", "RANK(8,Data!A1:A8,1)", "_xlfn.RANK.EQ(4,Data!K1:K8)",
    "_xlfn.RANK.AVG(4,Data!K1:K8)", "_xlfn.RANK.AVG(5,Data!K1:K8,1)", "RANK(6,Data!A1:A8)",
    "PERCENTILE(Data!A1:A8,0.3)", "_xlfn.PERCENTILE.INC(Data!A1:A8,0.9)",
    "_xlfn.PERCENTILE.EXC(Data!A1:A8,0.3)", "_xlfn.PERCENTILE.EXC(Data!A1:A8,0.05)",
    "QUARTILE(Data!A1:A8,1)", "_xlfn.QUARTILE.INC(Data!A1:A8,3)", "_xlfn.QUARTILE.EXC(Data!A1:A8,1)",
    "QUARTILE(Data!A1:A8,5)", "PERCENTRANK(Data!A1:A8,5)", "_xlfn.PERCENTRANK.INC(Data!A1:A8,4.5)",
    "_xlfn.PERCENTRANK.EXC(Data!A1:A8,5)", "PERCENTRANK(Data!A1:A8,5,5)", "PERCENTRANK(Data!A1:A8,10)",
    "CORREL(Data!A1:A8,Data!K1:K8)", "PEARSON(Data!A1:A8,Data!K1:K8)", "RSQ(Data!K1:K8,Data!A1:A8)",
    "SLOPE(Data!K1:K8,Data!A1:A8)", "INTERCEPT(Data!K1:K8,Data!A1:A8)",
    "FORECAST(10,Data!K1:K8,Data!A1:A8)", "_xlfn.FORECAST.LINEAR(10,Data!K1:K8,Data!A1:A8)",
    "STEYX(Data!K1:K8,Data!A1:A8)", "COVAR(Data!A1:A8,Data!K1:K8)",
    "_xlfn.COVARIANCE.P(Data!A1:A8,Data!K1:K8)", "_xlfn.COVARIANCE.S(Data!A1:A8,Data!K1:K8)",
    "NORMSDIST(1)", "_xlfn.NORM.S.DIST(1,TRUE)", "_xlfn.NORM.S.DIST(1,FALSE)", "NORMDIST(110,100,15,TRUE)",
    "_xlfn.NORM.DIST(110,100,15,FALSE)", "NORMSINV(0.975)", "_xlfn.NORM.S.INV(0.3)", "NORMINV(0.9,100,15)",
    "_xlfn.NORM.INV(0.001,0,1)", "NORMSINV(0)", "STANDARDIZE(5,3,2)", "_xlfn.GAUSS(1)", "_xlfn.PHI(1)",
    "BINOMDIST(3,10,0.4,FALSE)", "_xlfn.BINOM.DIST(3,10,0.4,TRUE)", "POISSON(3,2.5,FALSE)",
    "_xlfn.POISSON.DIST(3,2.5,TRUE)", "EXPONDIST(1,2,TRUE)", "_xlfn.EXPON.DIST(1,2,FALSE)",
    # domain errors: which kind does each application report?
    "ACOS(2)", "ATANH(1)", "_xlfn.ACOTH(0.5)", "SQRTPI(-1)", "FACT(-1)", "FACT(171)", "COMBIN(-1,1)",
    "PERMUT(2,3)", "FACTDOUBLE(-2)", "MULTINOMIAL(-1,2)", "LCM(-1)", "_xlfn.BITAND(-1,1)",
    "_xlfn.BITLSHIFT(1,100)", "EXP(1000)", "POWER(10,400)", "SINH(1000)", "COSH(1000)", "LN(-1)",
    "LOG(-1)", "LOG(10,1)", "TRIMMEAN(Data!A1:A8,1)", "GEOMEAN(Data!H1:H5,Data!J1:J5)",
    "HARMEAN(0,1)", "CORREL(Data!A1,Data!K1)", "_xlfn.STDEV.S(Data!A1)", "_xlfn.PERCENTILE.INC(Data!A1:A8,1.5)",
    "_xlfn.QUARTILE.EXC(Data!A1:A8,0)", "_xlfn.PERCENTILE.EXC(Data!A1:A8,0.99)", "NORMDIST(1,0,0,TRUE)",
    "NORMINV(1.5,0,1)", "STANDARDIZE(1,0,0)", "BINOMDIST(11,10,0.5,TRUE)", "POISSON(-1,2,TRUE)",
    "EXPONDIST(-1,2,TRUE)", "SLOPE(Data!A1,Data!K1)", "SKEW(Data!A1:A2)", "KURT(Data!A1:A3)", "RANK(1,Data!B1:B8)",
    "_xlfn.CEILING.MATH(1,-0.5)", "MROUND(-5,2)",
]


FORMULAS += [
    # dates (batch 3)
    "DAYS360(DATE(2024,1,30),DATE(2024,3,31))", "DAYS360(DATE(2024,2,29),DATE(2024,3,31))",
    "DAYS360(DATE(2024,1,31),DATE(2024,2,29),TRUE)", "DAYS360(Data!C1,Data!C8)",
    "DAYS360(DATE(2023,2,28),DATE(2023,3,31))", "DAYS360(DATE(2024,3,31),DATE(2024,1,15))",
    "WORKDAY(DATE(2024,1,15),10)", "WORKDAY(DATE(2024,1,15),-5)", "WORKDAY(DATE(2024,1,15),10,Data!C2:C3)",
    "WORKDAY(DATE(2024,1,13),1)", "WORKDAY(DATE(2024,1,15),0)", "WORKDAY.INTL(DATE(2024,1,15),10,11)",
    "WORKDAY.INTL(DATE(2024,1,15),10,\"0000011\")", "WORKDAY.INTL(DATE(2024,1,15),10,\"1111111\")",
    "NETWORKDAYS(DATE(2024,1,15),DATE(2024,2,16))", "NETWORKDAYS(DATE(2024,2,16),DATE(2024,1,15))",
    "NETWORKDAYS(DATE(2024,1,15),DATE(2024,3,31),Data!C2:C3)", "NETWORKDAYS.INTL(DATE(2024,1,15),DATE(2024,2,16),7)",
    "NETWORKDAYS.INTL(DATE(2024,1,15),DATE(2024,2,16),\"0010001\")",
    "YEARFRAC(DATE(2024,1,15),DATE(2024,8,22))", "YEARFRAC(DATE(2023,1,31),DATE(2024,3,1),0)",
    "YEARFRAC(DATE(2023,1,31),DATE(2024,3,1),1)", "YEARFRAC(DATE(2023,1,31),DATE(2024,3,1),2)",
    "YEARFRAC(DATE(2023,1,31),DATE(2024,3,1),3)", "YEARFRAC(DATE(2023,1,31),DATE(2024,3,1),4)",
    "YEARFRAC(DATE(2020,2,29),DATE(2024,2,28),1)", "YEARFRAC(DATE(2024,1,1),DATE(2024,12,31),1)",
    "YEARFRAC(DATE(2023,3,1),DATE(2024,2,29),1)", "YEARFRAC(DATE(2023,3,1),DATE(2024,2,29),5)",
    "WEEKNUM(DATE(2024,1,1))", "WEEKNUM(DATE(2024,12,31))", "WEEKNUM(DATE(2024,1,1),2)", "WEEKNUM(DATE(2021,1,3),11)",
    "WEEKNUM(DATE(2021,1,3),21)", "WEEKNUM(DATE(2024,1,1),17)", "WEEKNUM(DATE(2024,1,1),21)", "WEEKNUM(DATE(2024,1,1),5)",
    "_xlfn.ISOWEEKNUM(DATE(2021,1,3))", "_xlfn.ISOWEEKNUM(DATE(2020,12,31))", "_xlfn.ISOWEEKNUM(DATE(2024,12,30))",
    "TIMEVALUE(\"13:30\")", "TIMEVALUE(\"1:30 PM\")", "TIMEVALUE(\"2024-01-15 06:00\")", "TIMEVALUE(\"abc\")",
    # finance (batch 3)
    "NPER(0.01,-100,1000)", "NPER(0,-100,1000)", "NPER(0.05/12,-200,5000,0,1)", "NPER(0.01,-100,1000,500)",
    "RATE(60,-200,10000)", "RATE(60,-200,10000,0,1)", "RATE(10,0,-1000,2000)", "RATE(12,-100,1000,0,0,0.05)",
    "IPMT(0.01,1,12,1000)", "IPMT(0.01,5,12,1000)", "IPMT(0.01,5,12,1000,0,1)", "IPMT(0.01,1,12,1000,0,1)",
    "PPMT(0.01,5,12,1000)", "PPMT(0.01,1,12,1000,500,1)", "IPMT(0.01,13,12,1000)", "ISPMT(0.01,3,12,1000)",
    "CUMIPMT(0.01,12,1000,1,6,0)", "CUMPRINC(0.01,12,1000,1,6,0)", "CUMIPMT(0.01,12,1000,3,12,1)",
    "CUMPRINC(0.01,12,1000,1,12,1)", "CUMIPMT(0.01,12,1000,0,6,0)",
    "SLN(10000,1000,5)", "SYD(10000,1000,5,2)", "DDB(10000,1000,5,2)", "DDB(10000,1000,5,5)", "DDB(10000,1000,5,3,1.5)",
    "DB(10000,1000,5,1)", "DB(10000,1000,5,3)", "DB(10000,1000,5,6,6)", "DB(10000,1000,5,1,6)",
    "MIRR(Data!J1:J5,0.1,0.12)", "XNPV(0.1,Data!J1:J5,Data!C1:C5)", "XIRR(Data!J1:J5,Data!C1:C5)",
    "XIRR(Data!J1:J5,Data!C1:C5,0.5)", "EFFECT(0.1,4)", "NOMINAL(0.1,4)", "_xlfn.PDURATION(0.05,100,200)",
    "_xlfn.RRI(10,100,200)", "FVSCHEDULE(1000,{0.1,0.2})", "DOLLARDE(1.02,16)", "DOLLARFR(1.125,16)", "DOLLARDE(1.5,8)",
    # matrices, lookup, types (batch 3)
    "SUMPRODUCT(TRANSPOSE(Data!A1:A3)*{1,2,3})", "SUMPRODUCT(MMULT(TRANSPOSE(Data!H1:I5),Data!H1:I5))",
    "MDETERM(Data!H1:I2)", "SUMPRODUCT(MINVERSE(Data!H1:I2)*{1,2;3,4})", "SUMPRODUCT(_xlfn.MUNIT(3)*{1,2,3;4,5,6;7,8,9})",
    "SUMXMY2(Data!A1:A8,Data!K1:K8)", "SUMX2MY2(Data!A1:A8,Data!K1:K8)", "SUMX2PY2(Data!A1:A8,Data!K1:K8)",
    "MDETERM(Data!A1:B2)", "SUMPRODUCT(MINVERSE({1,2;2,4}))",
    "LOOKUP(6,Data!H1:H5)", "LOOKUP(6,Data!H1:H5,Data!I1:I5)", "LOOKUP(0,Data!H1:H5)", "LOOKUP(4,Data!H1:J5)",
    "TYPE(Data!A1)", "TYPE(Data!B1)", "TYPE(TRUE)", "TYPE(1/0)", "TYPE(Data!A1:A3)",
    # text (batch 3)
    "CLEAN(\"a\"&CHAR(10)&\"b\"&CHAR(7))", "_xlfn.UNICODE(\"é\")", "_xlfn.UNICHAR(233)", "_xlfn.UNICODE(\"\")",
    "ROMAN(1994)", "ROMAN(499,1)", "ROMAN(499,2)", "ROMAN(499,3)", "ROMAN(499,4)", "ROMAN(3999)", "ROMAN(0)",
    "ROMAN(4000)", "ROMAN(45,TRUE)", "_xlfn.ARABIC(\"MCMXCIV\")", "_xlfn.ARABIC(\"-IV\")", "_xlfn.ARABIC(\"\")",
    "_xlfn.ARABIC(\"ABC\")", "_xlfn.BASE(255,16)", "_xlfn.BASE(7,2,8)", "_xlfn.BASE(-1,2)",
    "_xlfn.DECIMAL(\"FF\",16)", "_xlfn.DECIMAL(\"102\",2)",
    # engineering (batch 3)
    "DEC2BIN(10)", "DEC2BIN(-1)", "DEC2BIN(10,8)", "DEC2BIN(600)", "DEC2OCT(64)", "DEC2HEX(255)", "DEC2HEX(-1)",
    "DEC2HEX(255,6)", "BIN2DEC(\"1010\")", "BIN2DEC(\"1111111111\")", "BIN2OCT(\"1111111111\")",
    "BIN2HEX(\"11111111\")", "OCT2DEC(\"777\")", "OCT2BIN(\"7\")", "OCT2HEX(\"777\")", "HEX2DEC(\"FF\")",
    "HEX2DEC(\"FFFFFFFFFF\")", "HEX2BIN(\"1F\")", "HEX2OCT(\"FF\")", "DELTA(2,2)", "DELTA(1,2)", "GESTEP(3,2)",
    "GESTEP(1)", "HEX2DEC(\"GG\")",
]


FORMULAS += [
    # subtotals over hidden rows (rows 7-8 of Data) and nested subtotals (batch 4)
    "SUBTOTAL(9,Data!A1:A8)", "SUBTOTAL(109,Data!A1:A8)", "SUBTOTAL(1,Data!A1:A8)", "SUBTOTAL(101,Data!A1:A8)",
    "SUBTOTAL(2,Data!A1:A8)", "SUBTOTAL(102,Data!A1:A8)", "SUBTOTAL(3,Data!B1:B8)", "SUBTOTAL(103,Data!B1:B8)",
    "SUBTOTAL(4,Data!A1:A8)", "SUBTOTAL(105,Data!A1:A8)", "SUBTOTAL(5,Data!A1:A8)", "SUBTOTAL(6,Data!A1:A3)",
    "SUBTOTAL(7,Data!A1:A8)", "SUBTOTAL(107,Data!A1:A8)", "SUBTOTAL(8,Data!A1:A8)", "SUBTOTAL(10,Data!A1:A8)",
    "SUBTOTAL(11,Data!A1:A8)", "SUBTOTAL(9,Data!L1:L6)", "SUBTOTAL(12,Data!A1:A3)", "SUBTOTAL(9,Data!A1:A2,Data!A4:A5)",
    "_xlfn.AGGREGATE(9,6,Data!A1:A8)", "_xlfn.AGGREGATE(9,5,Data!A1:A8)", "_xlfn.AGGREGATE(9,4,Data!A1:A8)", "_xlfn.AGGREGATE(4,6,Data!A1:A8)",
    "_xlfn.AGGREGATE(12,5,Data!A1:A8)", "_xlfn.AGGREGATE(14,5,Data!A1:A8,2)", "_xlfn.AGGREGATE(15,6,Data!A1:A8,2)",
    "_xlfn.AGGREGATE(16,6,Data!A1:A8,0.5)", "_xlfn.AGGREGATE(17,6,Data!A1:A8,1)", "_xlfn.AGGREGATE(9,6,Data!L1:L6)",
    "_xlfn.AGGREGATE(9,0,Data!L1:L6)", "_xlfn.AGGREGATE(1,6,Data!A1:A8)", "_xlfn.AGGREGATE(2,5,Data!A1:A8)", "_xlfn.AGGREGATE(22,6,Data!A1:A8)",
]


FORMULAS += [
    # a formula that ends in a blank cell (row `_n + 1` above): how do the rest of the functions see it?
    f"COUNTBLANK(A{_n + 1}:A{_n + 1})", f"COUNTA(A{_n + 1}:A{_n + 1})", f"COUNT(A{_n + 1}:A{_n + 1})",
    f"SUM(A{_n + 1}:A{_n + 1})", f"A{_n + 1}&\"x\"", f"LEN(A{_n + 1})", f"ISTEXT(A{_n + 1})",
    f"IF(A{_n + 1}=\"\",\"empty\",\"not\")", f"COUNTIF(A{_n + 1}:A{_n + 1},\"\")", f"COUNTIF(A{_n + 1}:A{_n + 1},0)",
    f"N(A{_n + 1})", f"T(A{_n + 1})", f"MAX(A{_n + 1}:A{_n + 1})", f"AVERAGE(A{_n + 1},1)", f"A{_n + 1}+1",
    f"MATCH(\"\",A{_n + 1}:A{_n + 1},0)", f"ISERROR(A{_n + 1})", f"ISNA(A{_n + 1})", f"A{_n + 1}<1", f"A{_n + 1}<\"a\"",
]


# Number and date format codes of TEXT. The same matrix is evaluated in an English
# and a Spanish (es-CO) LibreOffice: format codes are read in the user's locale
# (`#,##0` is `1.235` in both, `#,##0;(#,##0)` is `1,235` and `1234,568`), with the
# codes LibreOffice reads in English form whatever the locale.
_NUMS = ["1234.5678", "-1234.5678", "0", "0.5", "1000000", "12345678.9", "0.000123"]
_NUM_FORMATS = [
    "0", "0.00", "#,##0", "#,##0.00", "#.##", "0.0%", "0.00E+00", "# ?/?", "# ??/??", "General",
    "[$$-409]#,##0.00", "[$€-407]#,##0.00", "[$$-en-US]#,##0", "$#,##0", "#,##0;(#,##0)",
    "#,##0;[Red]-#,##0", "[>100]0.0;0", "0,", "0.0,,", "00000", "\"abc\"0", "0\" units\"",
    "#,##0_);(#,##0)", "#.##0", "#.###.##0", "0,00", "#.##0,00", "@", "0.0E+0", "##0.0E+0", "\\$0",
    "000-00", "#,###", "0.0", "0.000", "#.#", "00.00", "0.#", "#,#", "#,##", "0,0", "0,000",
    "#,##0.0", "0.0.0", "0,0,0", "0.00%", "#,##0.00_);(#,##0.00)", "$#,##0.00", "#,##0.00;-#,##0.00",
    "#,##0.000", "0.00 \"x\"", "0.000%", "0.0000", "0.00;0.00", "#,##0.00;[Red](#,##0.00)", "0.00E+000",
    "0.000E+00", "#,##0.00E+00",
]
_DATES = ["45292.5", "45658.75123", "0.5", "1.999", "61", "36526.12345", "0.999999", "-1.5"]
_DATE_FORMATS = [
    "yyyy-mm-dd", "dd/mm/yyyy", "d-mmm-yy", "mmmm", "mmm", "dddd", "ddd", "h:mm", "hh:mm:ss", "h:mm AM/PM",
    "hh:mm:ss.000", "hh:mm:ss.0", "[h]:mm", "[m]", "[s]", "mm:ss", "yy", "yyyy", "m/d/yy", "d", "dd", "mmmmm",
    "yyyy-mm-dd hh:mm", "AM/PM", "h A/P", "mm", "m", "mm-dd", "dd/mm/yyyy hh:mm:ss", "[hh]:mm:ss",
    "yyyy\"/\"mm", "ddd, d mmm yyyy", "mm/dd/yy;@", "[$-409]dddd", "[$-40A]dddd", "h:mm:ss;@", "[Red]yyyy",
    "yyyy [Blue]", "0.0 \"h\"", "d \"de\" mmmm \"de\" yyyy", "hh\"h\"mm", "nn", "NN", "DDDD", "MMMM", "YYYY",
    "H:MM", "ss", "s.00", "[h]", "m:ss", "mm:ss.0", "yyyy-mm-dd\"T\"hh:mm:ss",
]


FORMULAS += [
    # distributions (batch 4): gamma, chi-square, t, F, beta, log-normal, Weibull and the discrete ones
    "GAMMALN(5)", "GAMMALN(0.5)", "GAMMALN(100.3)", "GAMMALN(0)", "GAMMALN(-1)", "_xlfn.GAMMALN.PRECISE(10.5)",
    "_xlfn.GAMMALN.PRECISE(-2)", "_xlfn.GAMMA(5)", "_xlfn.GAMMA(0.5)", "_xlfn.GAMMA(-1.5)", "_xlfn.GAMMA(0)",
    "_xlfn.GAMMA(-2)", "_xlfn.GAMMA(171)", "_xlfn.GAMMA(172)", "_xlfn.GAMMA(1E-300)",
    "FISHER(0.75)", "FISHER(-0.3)", "FISHER(0)", "FISHER(1)", "FISHER(2)", "FISHERINV(1)", "FISHERINV(-0.5)",
    "FISHERINV(30)", "FISHERINV(0)",
    "_xlfn.GAMMA.DIST(2,3,2,TRUE)", "_xlfn.GAMMA.DIST(2,3,2,FALSE)", "GAMMADIST(0.5,0.5,1,TRUE)",
    "GAMMADIST(10,2,3,FALSE)", "_xlfn.GAMMA.DIST(0,3,2,TRUE)", "_xlfn.GAMMA.DIST(0,1,2,FALSE)",
    "_xlfn.GAMMA.DIST(0,0.5,2,FALSE)", "_xlfn.GAMMA.DIST(-1,3,2,TRUE)", "_xlfn.GAMMA.DIST(2,0,2,TRUE)",
    "_xlfn.GAMMA.DIST(2,3,0,TRUE)", "_xlfn.GAMMA.DIST(500,3,2,TRUE)",
    "_xlfn.GAMMA.INV(0.5,3,2)", "GAMMAINV(0.9,0.5,1)", "_xlfn.GAMMA.INV(0.001,10,2)", "_xlfn.GAMMA.INV(0,3,2)",
    "_xlfn.GAMMA.INV(1,3,2)", "_xlfn.GAMMA.INV(1.5,3,2)", "_xlfn.GAMMA.INV(0.5,-1,2)", "_xlfn.GAMMA.INV(0.5,3,0)",
    "_xlfn.CHISQ.DIST(3,4,TRUE)", "_xlfn.CHISQ.DIST(3,4,FALSE)", "_xlfn.CHISQ.DIST.RT(3,4)", "CHIDIST(3,4)",
    "CHIDIST(0,4)", "_xlfn.CHISQ.DIST(-1,4,TRUE)", "_xlfn.CHISQ.DIST(3,0,TRUE)", "_xlfn.CHISQ.DIST(3,1E10,TRUE)",
    "_xlfn.CHISQ.DIST(3,4.9,TRUE)", "_xlfn.CHISQ.DIST(0,2,FALSE)", "_xlfn.CHISQ.DIST(0,1,FALSE)",
    "_xlfn.CHISQ.INV(0.5,4)", "_xlfn.CHISQ.INV.RT(0.05,4)", "CHIINV(0.05,4)", "CHIINV(0.5,1)",
    "_xlfn.CHISQ.INV(0,4)", "_xlfn.CHISQ.INV(1,4)", "CHIINV(1.5,4)", "CHIINV(0.5,0)", "_xlfn.CHISQ.INV(0.99,30)",
    "_xlfn.CHISQ.TEST(Data!H1:H5,Data!I1:I5)", "CHITEST(Data!H1:H5,Data!K1:K5)", "CHITEST(Data!A1:A2,Data!A1:A3)",
    "_xlfn.T.DIST(1.5,10,TRUE)", "_xlfn.T.DIST(1.5,10,FALSE)", "_xlfn.T.DIST(-1.5,10,TRUE)", "_xlfn.T.DIST(0.3,1,TRUE)",
    "_xlfn.T.DIST(2,0,TRUE)", "_xlfn.T.DIST(2,3.7,TRUE)", "_xlfn.T.DIST.RT(1.5,10)",
    "_xlfn.T.DIST.RT(-1.5,10)", "_xlfn.T.DIST.2T(1.5,10)", "_xlfn.T.DIST.2T(-1.5,10)", "_xlfn.T.DIST.2T(0,10)",
    "TDIST(1.5,10,1)", "TDIST(1.5,10,2)", "TDIST(-1.5,10,2)", "TDIST(1.5,10,3)", "TDIST(1.5,0,1)", "TDIST(1.5,10.9,2)",
    "_xlfn.T.INV(0.9,10)", "_xlfn.T.INV(0.1,10)", "_xlfn.T.INV(0.5,10)", "_xlfn.T.INV.2T(0.1,10)", "TINV(0.05,10)",
    "TINV(0.05,1)", "_xlfn.T.INV(0,10)", "_xlfn.T.INV(1,10)", "TINV(0,10)", "TINV(1.5,10)", "TINV(0.05,0)",
    "_xlfn.T.INV(0.975,2.9)",
    "_xlfn.T.TEST(Data!A1:A8,Data!K1:K8,2,1)", "_xlfn.T.TEST(Data!A1:A8,Data!K1:K8,1,2)",
    "TTEST(Data!A1:A8,Data!K1:K8,2,3)", "TTEST(Data!A1:A5,Data!K1:K8,2,2)", "TTEST(Data!A1:A5,Data!K1:K8,2,3)",
    "TTEST(Data!A1:A5,Data!K1:K8,2,1)", "TTEST(Data!A1:A8,Data!K1:K8,3,2)", "TTEST(Data!A1:A8,Data!K1:K8,2,4)",
    "_xlfn.F.DIST(2,5,8,TRUE)", "_xlfn.F.DIST(2,5,8,FALSE)", "_xlfn.F.DIST.RT(2,5,8)", "FDIST(2,5,8)",
    "_xlfn.F.DIST(0,5,8,TRUE)", "_xlfn.F.DIST(0,2,8,FALSE)", "_xlfn.F.DIST(-1,5,8,TRUE)", "_xlfn.F.DIST(2,0,8,TRUE)",
    "_xlfn.F.DIST(2,5,1E11,TRUE)", "_xlfn.F.DIST(2,5.9,8.9,TRUE)", "FDIST(0,5,8)",
    "_xlfn.F.INV(0.9,5,8)", "_xlfn.F.INV.RT(0.05,5,8)", "FINV(0.05,5,8)", "_xlfn.F.INV(0,5,8)", "_xlfn.F.INV(1,5,8)",
    "FINV(1.5,5,8)", "FINV(0.05,0,8)",
    "_xlfn.F.TEST(Data!A1:A8,Data!K1:K8)", "FTEST(Data!A1:A5,Data!K1:K8)", "FTEST(Data!A1,Data!K1:K8)",
    "_xlfn.BETA.DIST(0.4,2,3,TRUE)", "_xlfn.BETA.DIST(0.4,2,3,FALSE)", "_xlfn.BETA.DIST(5,2,3,TRUE,1,10)",
    "_xlfn.BETA.DIST(5,2,3,FALSE,1,10)", "BETADIST(0.4,2,3)", "BETADIST(5,2,3,1,10)", "_xlfn.BETA.DIST(0,2,3,TRUE)",
    "_xlfn.BETA.DIST(1,2,3,TRUE)", "_xlfn.BETA.DIST(0,1,3,FALSE)", "_xlfn.BETA.DIST(1.5,2,3,TRUE)",
    "_xlfn.BETA.DIST(0.4,0,3,TRUE)", "_xlfn.BETA.DIST(0.4,2,3,TRUE,1,1)", "BETADIST(0.4,2,3,0.5,0.3)",
    "_xlfn.BETA.INV(0.5,2,3)", "BETAINV(0.9,2,3,1,10)", "_xlfn.BETA.INV(0,2,3)", "_xlfn.BETA.INV(1,2,3)",
    "BETAINV(1.5,2,3)", "BETAINV(0.5,0,3)", "BETAINV(0.5,2,3,1,1)",
    "_xlfn.LOGNORM.DIST(4,1.3,0.7,TRUE)", "_xlfn.LOGNORM.DIST(4,1.3,0.7,FALSE)", "LOGNORMDIST(4,1.3,0.7)",
    "_xlfn.LOGNORM.DIST(0,1.3,0.7,TRUE)", "_xlfn.LOGNORM.DIST(-1,1.3,0.7,TRUE)", "_xlfn.LOGNORM.DIST(4,1.3,0,TRUE)",
    "_xlfn.LOGNORM.INV(0.5,1.3,0.7)", "LOGINV(0.9,0,1)", "LOGINV(0,0,1)", "LOGINV(1,0,1)", "LOGINV(0.5,0,0)",
    "_xlfn.WEIBULL.DIST(2,1.5,3,TRUE)", "_xlfn.WEIBULL.DIST(2,1.5,3,FALSE)", "WEIBULL(2,1.5,3,TRUE)",
    "WEIBULL(0,1.5,3,TRUE)", "WEIBULL(0,1,3,FALSE)", "WEIBULL(-1,1.5,3,TRUE)", "WEIBULL(2,0,3,TRUE)", "WEIBULL(2,1.5,0,TRUE)",
    "_xlfn.HYPGEOM.DIST(2,5,4,12,TRUE)", "_xlfn.HYPGEOM.DIST(2,5,4,12,FALSE)", "HYPGEOMDIST(2,5,4,12)",
    "HYPGEOMDIST(0,5,4,12)", "HYPGEOMDIST(6,5,4,12)", "HYPGEOMDIST(2,5,4,3)", "HYPGEOMDIST(2,5,4,2)", "HYPGEOMDIST(2.9,5.9,4.9,12.9)",
    "_xlfn.NEGBINOM.DIST(3,2,0.4,TRUE)", "_xlfn.NEGBINOM.DIST(3,2,0.4,FALSE)", "NEGBINOMDIST(3,2,0.4)",
    "NEGBINOMDIST(0,1,0.4)", "NEGBINOMDIST(-1,2,0.4)", "NEGBINOMDIST(3,0,0.4)", "NEGBINOMDIST(3,2,0)", "NEGBINOMDIST(3,2,1)",
    "_xlfn.BINOM.INV(10,0.4,0.5)", "CRITBINOM(10,0.4,0.9)", "CRITBINOM(10,0.4,0)", "CRITBINOM(10,0.4,1)", "CRITBINOM(10,0.4,1.5)",
    "CRITBINOM(-1,0.4,0.5)", "CRITBINOM(10,1.5,0.5)", "CRITBINOM(10.9,0.4,0.5)", "CRITBINOM(1000,0.5,0.999)",
    "_xlfn.CONFIDENCE.NORM(0.05,2.5,50)", "CONFIDENCE(0.05,2.5,50)", "_xlfn.CONFIDENCE.T(0.05,2.5,50)",
    "_xlfn.CONFIDENCE.T(0.05,2.5,1)", "_xlfn.CONFIDENCE.T(0.05,2.5,2)", "CONFIDENCE(0,2.5,50)", "CONFIDENCE(1,2.5,50)",
    "CONFIDENCE(0.05,0,50)", "CONFIDENCE(0.05,2.5,0.5)", "CONFIDENCE(0.05,2.5,50.9)",
    "_xlfn.Z.TEST(Data!A1:A8,4)", "ZTEST(Data!A1:A8,6,2)", "ZTEST(Data!A1:A8,6,0)", "ZTEST(Data!A1,6)", "ZTEST(Data!H1:H5,2)",
    "PROB(Data!H1:H5,Data!I1:I5,3,7)", "PROB(Data!H1:H5,Data!I1:I5,3)", "PROB(Data!H1:H5,Data!I1:I5,10)",
    "PROB(Data!H1:H5,Data!A1:A5,3)", "PROB(Data!H1:H5,Data!A1:A5,3,7)", "PROB(Data!H1:H3,Data!I1:I5,3)",
    "PROB(Data!H1:H5,Data!I1:I5,7,3)", "PROB(Data!H1:H5,Data!K1:K5,3,7)",
    "B(10,0.4,3)", "B(10,0.4,3,5)", "B(10,0.4,5,3)", "B(10,0.4,11)", "B(10,1.5,3)", "B(10,0.4,3,11)", "B(10,0.4,-1)",
]


FORMULAS += [
    # distributions (batch 4b): large degrees of freedom, boundaries and the arrays of PROB
    # More than 1E7 degrees of freedom or shape (T.DIST(2,1E8) is off by 2E-9 already): LibreOffice's result is off (T.DIST(2,1E9) by 3E-9,
    # F.DIST(2,5,1E9) by 4E-8, BETA.DIST(0.5,1E9,1E9) by 1E-7) and the build stops (`notes::BETA_LARGE`).
    *[f"_xlfn.T.DIST(2,{d},TRUE)" for d in ("100", "1000", "1E4", "1E5", "1E6", "1E7")],
    *[f"_xlfn.CHISQ.DIST({d},{d},TRUE)" for d in ("1000", "1E5", "1E7", "1E9", "1E10")],
    *[f"_xlfn.F.DIST(2,{d},{d},TRUE)" for d in ("1000", "1E5", "1E7")],
    "_xlfn.F.DIST(2,5,1E10,TRUE)", "_xlfn.CHISQ.DIST(3,10000000001,TRUE)", "CHIDIST(3,1E10)", "CHIDIST(3,10000000001)", "_xlfn.CHISQ.DIST.RT(3,10000000001)", "_xlfn.CHISQ.INV(0.5,10000000001)",
    "CHIINV(0.5,10000000001)", "_xlfn.CHISQ.INV(0.5,1E10)", "_xlfn.CHISQ.INV.RT(0.5,1E10)",
    *[f"_xlfn.GAMMA.DIST({d},{d},1,TRUE)" for d in ("1000", "1E5", "1E7", "1E9")],
    *[f"_xlfn.BETA.DIST(0.5,{d},{d},TRUE)" for d in ("1000", "1E6", "1E7")],
    "CHIINV(1,4)", "CHIINV(0,4)", "TINV(1,10)", "TINV(0.5,10)", "FINV(1,5,8)", "_xlfn.F.INV.RT(1,5,8)", "_xlfn.F.INV.RT(0,5,8)",
    "_xlfn.BETA.DIST(0,0.5,3,FALSE)", "_xlfn.BETA.DIST(1,2,1,FALSE)", "_xlfn.LOGNORM.DIST(0,1.3,0.7,FALSE)",
    "_xlfn.LOGNORM.DIST(-1,1.3,0.7,FALSE)", "WEIBULL(0,0.5,3,FALSE)", "WEIBULL(0,2,3,FALSE)",
    "_xlfn.NEGBINOM.DIST(3,0,0.4,TRUE)", "_xlfn.NEGBINOM.DIST(0,2,0.4,TRUE)", "_xlfn.NEGBINOM.DIST(3.9,2.9,0.4,FALSE)",
    "_xlfn.NEGBINOM.DIST(0,2,0.4,FALSE)", "_xlfn.NEGBINOM.DIST(0,2,1,FALSE)",
    "_xlfn.HYPGEOM.DIST(0,0,4,12,FALSE)", "_xlfn.HYPGEOM.DIST(0,0,0,12,TRUE)", "_xlfn.HYPGEOM.DIST(4,5,4,12,TRUE)",
    "_xlfn.HYPGEOM.DIST(1,5,4,12,TRUE)", "_xlfn.HYPGEOM.DIST(5,5,5,5,FALSE)",
    "B(0,0.4,0)", "B(10,0,0)", "B(10,1,10)", "B(10,0.4,0,10)", "B(10.9,0.4,3.9)",
    "_xlfn.CONFIDENCE.T(0.05,2.5,3.9)", "_xlfn.CONFIDENCE.NORM(0.05,2.5,1)", "_xlfn.CONFIDENCE.T(0.05,2.5,0.9)",
    "TTEST({1;1;1},{1;1;1},2,2)", "TTEST({1;2;3},{1;2;3},2,1)", "TTEST({1;2;3},{2;3;5},2,1)", "TTEST({1;2},{1;2;3},1,3)",
    "TTEST({1},{1;2;3},1,2)", "FTEST({1;1;1},{1;2;3})", "FTEST({1;2;3},{1;1;1})", "FTEST({1;2;3},{1;2;3.5})",
    "ZTEST({1;2;3},2)", "ZTEST({1;2;3},2,1)", "ZTEST({1;1;1},2)", "ZTEST({1},2,3)",
    "PROB({1;2;3;4},{0.1;0.2;0.3;0.4},2,3)", "PROB({1;2;3;4},{0.1;0.2;0.3;0.4},2)", "PROB({1;2;3;4},{0.1;0.2;0.3;0.4},2.5)",
    "PROB({1;2;3;4},{0.1;0.2;0.3;0.4},0,10)", "PROB({1;2;3;4},{0.1;0.2;0.3;0.4},3,2)", "PROB({1;2;3;4},{0.1;0.2;0.3;0.5},2,3)",
    "PROB({1;2;3;4},{0.1;0.2;0.3;0.3},2,3)", "PROB({1;2;3;4},{-0.1;0.4;0.3;0.4},2,3)", "PROB({1;2;3;4},{0.1;0.2;0.3},2,3)",
    "PROB({1;2},{0.5;0.5},1.5,2.5)", "PROB({1;2;3;4},{0.1;0.2;0.3;0.4},4,4)", "PROB({1;2;3;4},{0.1;0.2;0.3;0.4},1)",
    "CHITEST({1;2;3},{1;2;3})", "CHITEST({1,2;3,4},{2,2;2,3})", "CHITEST({1,2;3,4},{1,0;3,4})", "CHITEST({5},{5})",
    "_xlfn.GAMMALN.PRECISE(0.1)", "GAMMALN(1E10)", "_xlfn.GAMMA(1E-5)", "_xlfn.GAMMA(-0.5)", "_xlfn.GAMMA(-170.5)",
    "_xlfn.GAMMA(100.5)", "_xlfn.GAMMA(170.5)", "_xlfn.GAMMA(171.5)", "FISHER(0.999999)", "FISHERINV(-30)",
    "_xlfn.GAMMA.INV(0.5,0.01,1)", "_xlfn.GAMMA.INV(0.999999,50,2)", "_xlfn.CHISQ.INV(0.000001,3)", "_xlfn.CHISQ.INV.RT(0.000001,3)",
    "_xlfn.T.INV(0.000001,3)", "_xlfn.T.INV(0.999999,3)", "_xlfn.T.INV.2T(0.000001,3)", "_xlfn.F.INV(0.000001,3,4)",
    "_xlfn.F.INV.RT(0.000001,3,4)", "_xlfn.BETA.INV(0.000001,2,3)", "_xlfn.BETA.INV(0.999999,0.5,0.5)",
]


FORMULAS += [
    # dynamic arrays and text (batch 5)
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.SEQUENCE(5))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.SEQUENCE(2,3))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.SEQUENCE(3,1,10,-2))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.SEQUENCE(2,2,0.5,0.25))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.SEQUENCE(2.9,1.9))",
    "_xlfn.SEQUENCE(0)", "_xlfn.SEQUENCE(1,0)", "_xlfn.SEQUENCE(-1)", "_xlfn.SEQUENCE(\"a\")",
    "ROWS(_xlfn.SEQUENCE(4,3))", "COLUMNS(_xlfn.SEQUENCE(4,3))", "SUM(_xlfn.SEQUENCE(100))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.SEQUENCE(3,2,5,0))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.SEQUENCE(1,3,TRUE))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.SEQUENCE(2,,3))", "_xlfn.XMATCH(5,Data!A1:A8)", "_xlfn.XMATCH(6,Data!A1:A8)",
    "_xlfn.XMATCH(6,Data!A1:A8,-1)", "_xlfn.XMATCH(6,Data!A1:A8,1)", "_xlfn.XMATCH(0,Data!A1:A8,-1)",
    "_xlfn.XMATCH(100,Data!A1:A8,1)", "_xlfn.XMATCH(\"a*\",Data!B1:B8,2)", "_xlfn.XMATCH(\"*e\",Data!B1:B8,2)",
    "_xlfn.XMATCH(\"apple\",Data!B1:B8,0,-1)", "_xlfn.XMATCH(\"APPLE\",Data!B1:B8)",
    "_xlfn.XMATCH(\"fig\",Data!B1:B8,0,1)", "_xlfn.XMATCH(3,{1;2;3;4},0,2)", "_xlfn.XMATCH(3,{4;3;2;1},0,-2)",
    "_xlfn.XMATCH(2.5,{1;2;3;4},1,2)", "_xlfn.XMATCH(2.5,{1;2;3;4},-1,2)", "_xlfn.XMATCH(3,{1,2,3,4})",
    "_xlfn.XMATCH(3,{1,2;3,4})", "_xlfn.XMATCH(5,Data!A1:A8,3)", "_xlfn.XMATCH(5,Data!A1:A8,0,0)",
    "_xlfn.XMATCH(5,Data!A1:A8,0,3)", "_xlfn.XMATCH(5,Data!A1)", "_xlfn.XMATCH(TRUE,{1;TRUE;3})",
    "_xlfn.XMATCH(\"5\",Data!A1:A8)", "_xlfn.XMATCH(\"?e*\",Data!B1:B8,2)",
    "_xlfn.XMATCH(\"a~*\",{\"a*\";\"ab\"},2)", "_xlfn.XMATCH(5,{\"a\";5})",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!A1:A8,Data!A1:A8>4))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!B1:B8,Data!A1:A8>4))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!A1:B8,Data!A1:A8>4))",
    "_xlfn._xlws.FILTER(Data!A1:A8,Data!A1:A8>100,\"none\")", "_xlfn._xlws.FILTER(Data!A1:A8,Data!A1:A8>100)",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!A1:A8,Data!K1:K8))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!A1:A8,{1;0;1}))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!E2:G2,Data!E2:G2>15))",
    "COUNTA(_xlfn._xlws.FILTER(Data!A1:A8,Data!B1:B8=\"apple\"))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!A1:A8,(Data!B1:B8=\"apple\")+(Data!A1:A8>8)))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!A1:A8,(Data!B1:B8=\"apple\")*(Data!A1:A8>4)))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!A1:A8,Data!B1:B8))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!A1:A8,{TRUE;FALSE;TRUE;FALSE;TRUE;FALSE;TRUE;FALSE}))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!A1:B3,{TRUE;FALSE;TRUE}))", "_xlfn._xlws.FILTER(Data!A1:A8,{1;0})",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!A1:A8,Data!A1:A8>4,\"x\"))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!A1:A3,{1;1;1},0))", "_xlfn._xlws.FILTER(Data!A1:A8,Data!A1:A8>100,)",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.SORT(Data!A1:A8))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.SORT(Data!A1:A8,1,-1))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.SORT(Data!B1:B8))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.SORT(Data!A1:B8,2))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.SORT(Data!A1:B8,2,-1))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.SORT(Data!A1:B8,1,1,FALSE))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.SORT(Data!E2:G2,1,-1,TRUE))", "_xlfn._xlws.SORT(Data!A1:B8,3)",
    "_xlfn._xlws.SORT(Data!A1:A8,1,0)", "_xlfn._xlws.SORT(Data!A1:A8,0)", "_xlfn._xlws.SORT(Data!A1:A8,1,2)",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.SORT({3;\"a\";TRUE;2;\"B\";1}))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.SORT({3;\"b\";\"A\";\"a\";\"B\"}))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.SORT({3;1;2},1,-1))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.SORT(Data!A1:K3,11))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.SORT({1;5;3},1,1,TRUE))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.SORT(Data!A1:A8,1,-1,FALSE))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.SORT(Data!B1:B8,1,-1))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.SORT({\"é\";\"e\";\"f\";\"E\";\"z\";\"Z\"}))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.SORT({10;9;\"10\";\"9\"}))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.SORTBY(Data!B1:B8,Data!A1:A8))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.SORTBY(Data!A1:A8,Data!B1:B8,1,Data!A1:A8,-1))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.SORTBY(Data!A1:A8,Data!B1:B8,-1))", "_xlfn.SORTBY(Data!A1:A8,Data!B1:B7)",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.SORTBY(Data!A1:B8,Data!A1:A8,-1))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.SORTBY(Data!A1:A8,Data!A1:A8))", "_xlfn.SORTBY(Data!A1:A8,Data!B1:B8,0)",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.SORTBY(Data!E1:G2,{3,1,2}))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.SORTBY(Data!A1:A8,Data!B1:B8,1,Data!K1:K8))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TAKE(Data!A1:A8,3))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TAKE(Data!A1:A8,-3))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TAKE(Data!A1:B8,2,1))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TAKE(Data!A1:B8,,1))",
    "_xlfn.TAKE(Data!A1:A8,0)", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TAKE(Data!A1:A8,100))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TAKE(Data!A1:B8,2,-1))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TAKE(Data!A1:B3,-2,-1))", "_xlfn.TAKE(Data!A1:A8,2,0)",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TAKE(Data!A1:A8,2,5))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TAKE(Data!E1:G2,1))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.DROP(Data!A1:A8,5))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.DROP(Data!A1:A8,-5))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.DROP(Data!A1:B8,1,1))", "_xlfn.DROP(Data!A1:A8,8)",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.DROP(Data!A1:A8,0))", "_xlfn.DROP(Data!A1:A8,100)",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.DROP(Data!A1:B3,,1))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.DROP(Data!A1:B3,-1,-1))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.CHOOSECOLS(Data!A1:B3,2))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.CHOOSECOLS(Data!A1:B3,2,1))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.CHOOSECOLS(Data!A1:B3,-1))", "_xlfn.CHOOSECOLS(Data!A1:B3,3)",
    "_xlfn.CHOOSECOLS(Data!A1:B3,0)", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.CHOOSECOLS(Data!A1:B3,1,1,2.9))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.CHOOSEROWS(Data!A1:B8,1,-1))", "_xlfn.CHOOSEROWS(Data!A1:B8,9)",
    "_xlfn.CHOOSEROWS(Data!A1:B8,0)", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.CHOOSEROWS(Data!A1:B3,{1,3}))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.CHOOSEROWS(Data!A1:A8,-8,8))", "_xlfn.CHOOSEROWS(Data!A1:A8,\"a\")",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.CHOOSECOLS(Data!A1:G1,{7,1}))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.HSTACK(Data!A1:A3,Data!B1:B3))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.VSTACK(Data!A1:A2,Data!B1:B3))",
    "SUMPRODUCT(--ISNA(_xlfn.HSTACK(Data!A1:A3,{1;2})))", "SUMPRODUCT(--ISNA(_xlfn.VSTACK(Data!A1:B2,{1;2;3})))",
    "ROWS(_xlfn.VSTACK(Data!A1:A2,Data!B1:B3))", "COLUMNS(_xlfn.HSTACK(Data!A1:A3,Data!B1:B3,Data!A1))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.VSTACK(1,2,3))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.HSTACK(1,2,3))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.HSTACK({1;2},\"a\"))", "ROWS(_xlfn.HSTACK(Data!A1:A3,Data!A1:A8))",
    "SUMPRODUCT(--ISNA(_xlfn.EXPAND(Data!A1:A3,5)))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.EXPAND(Data!A1:A3,5,2,0))",
    "_xlfn.EXPAND(Data!A1:A3,2)", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.EXPAND(Data!A1:A3,4,,\"x\"))",
    "_xlfn.EXPAND(Data!A1:A3,5,0)", "ROWS(_xlfn.EXPAND(Data!A1:B2,3,3))", "COLUMNS(_xlfn.EXPAND(Data!A1:B2,3,3))",
    "_xlfn.EXPAND(Data!A1:A3,3,1.9)", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TOCOL(Data!A1:B3))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TOCOL(Data!A1:B3,0,TRUE))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TOROW(Data!A1:B3))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TOROW(Data!A1:B3,0,TRUE))", "ROWS(_xlfn.TOCOL(Data!A1:B3))",
    "COLUMNS(_xlfn.TOROW(Data!A1:B3))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TOCOL({1,\"\";2,3},1))",
    "_xlfn.TOCOL(Data!A1:B3,4)",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TOCOL(5))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.WRAPCOLS(Data!A1:A8,3))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.WRAPROWS(Data!A1:A8,3))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.WRAPROWS(Data!A1:A8,3,0))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.WRAPCOLS(Data!A1:A8,3,0))", "ROWS(_xlfn.WRAPCOLS(Data!A1:A8,3))",
    "COLUMNS(_xlfn.WRAPCOLS(Data!A1:A8,3))", "ROWS(_xlfn.WRAPROWS(Data!A1:A8,3))",
    "COLUMNS(_xlfn.WRAPROWS(Data!A1:A8,3))", "_xlfn.WRAPROWS(Data!A1:A8,0)", "_xlfn.WRAPROWS(Data!A1:B3,2)",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.WRAPROWS(Data!A1:A8,100))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.WRAPROWS(Data!A1:C1,2))", "SUMPRODUCT(--ISNA(_xlfn.WRAPROWS(Data!A1:A8,3)))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"a,b,c\",\",\"))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"a,b;c,d\",\",\",\";\"))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"a,,b\",\",\",,TRUE))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"a,,b\",\",\"))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"aXbxc\",\"x\",,FALSE,1))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"aXbxc\",\"x\"))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"a,b;c\",\",\",\";\",FALSE,0,\"-\"))",
    "ROWS(_xlfn.TEXTSPLIT(\"a,b;c,d\",\",\",\";\"))", "COLUMNS(_xlfn.TEXTSPLIT(\"a,b;c,d\",\",\",\";\"))",
    "SUMPRODUCT(--ISNA(_xlfn.TEXTSPLIT(\"a,b;c\",\",\",\";\")))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"a1b2c\",{\"1\",\"2\"}))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"a,b|c\",{\",\",\"|\"}))", "_xlfn.TEXTSPLIT(\"abc\",\"\")",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"abc\",\"x\"))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"\",\",\"))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"a,b\",,\",\"))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"a,b\",\",\",,FALSE,2))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(123,2))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\",a,\",\",\",,TRUE))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"a  b\",\" \"))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"a;;b,c\",\",\",\";\",TRUE))", "_xlfn.TEXTBEFORE(\"a-b-c\",\"-\")",
    "_xlfn.TEXTBEFORE(\"a-b-c\",\"-\",2)", "_xlfn.TEXTBEFORE(\"a-b-c\",\"-\",-1)",
    "_xlfn.TEXTBEFORE(\"a-b-c\",\"-\",-2)", "_xlfn.TEXTBEFORE(\"a-b-c\",\"-\",4)",
    "_xlfn.TEXTBEFORE(\"a-b-c\",\"x\",1,0,0,\"none\")", "_xlfn.TEXTBEFORE(\"A-b\",\"a\",1,1)",
    "_xlfn.TEXTBEFORE(\"A-b\",\"a\")", "_xlfn.TEXTBEFORE(\"abc\",\"c\",1,0,1)",
    "_xlfn.TEXTBEFORE(\"abc\",\"d\",1,0,1)", "_xlfn.TEXTBEFORE(\"a-b-c\",\"-\",0)",
    "_xlfn.TEXTBEFORE(\"a-b-c\",\"\")", "_xlfn.TEXTBEFORE(\"abc\",\"\",-1)",
    "_xlfn.TEXTBEFORE(\"a-b-c\",\"-\",5,0,1)", "_xlfn.TEXTBEFORE(\"a-b-c\",{\"-\",\"b\"})",
    "_xlfn.TEXTBEFORE(\"a-b-c\",\"-\",-3)", "_xlfn.TEXTBEFORE(\"a--b\",\"--\")", "_xlfn.TEXTBEFORE(12345,3)",
    "_xlfn.TEXTBEFORE(\"\",\"a\")", "_xlfn.TEXTBEFORE(\"aaa\",\"aa\",2)", "_xlfn.TEXTAFTER(\"a-b-c\",\"-\")",
    "_xlfn.TEXTAFTER(\"a-b-c\",\"-\",2)", "_xlfn.TEXTAFTER(\"a-b-c\",\"-\",-1)",
    "_xlfn.TEXTAFTER(\"a-b-c\",\"-\",-2)", "_xlfn.TEXTAFTER(\"a-b-c\",\"-\",4)",
    "_xlfn.TEXTAFTER(\"a-b-c\",\"x\",1,0,0,\"none\")", "_xlfn.TEXTAFTER(\"A-b\",\"a\",1,1)",
    "_xlfn.TEXTAFTER(\"abc\",\"c\",1,0,1)", "_xlfn.TEXTAFTER(\"abc\",\"d\",1,0,1)",
    "_xlfn.TEXTAFTER(\"a-b-c\",\"-\",0)", "_xlfn.TEXTAFTER(\"a-b-c\",\"\")", "_xlfn.TEXTAFTER(\"abc\",\"\",-1)",
    "_xlfn.TEXTAFTER(\"a-b-c\",\"-\",5,0,1)", "_xlfn.TEXTAFTER(\"a-b-c\",{\"-\",\"b\"})",
    "_xlfn.TEXTAFTER(\"a--b\",\"--\")", "_xlfn.TEXTAFTER(\"aaa\",\"aa\",2)", "ROWS(_xlfn.RANDARRAY(3))",
    "COLUMNS(_xlfn.RANDARRAY(2,4))", "SUMPRODUCT(--(_xlfn.RANDARRAY(5,5,10,20)>=10))",
    "MAX(_xlfn.RANDARRAY(100,1,5,5,TRUE))",
    "SUMPRODUCT(--(_xlfn.RANDARRAY(20,5,1,3,TRUE)<>INT(_xlfn.RANDARRAY(20,5,1,3,TRUE)*0)))", "_xlfn.RANDARRAY(0)",
    "_xlfn.RANDARRAY(2,2,5,1)", "SUMPRODUCT(--(_xlfn.RANDARRAY(10,10,1,2,TRUE)>2))", "_xlfn.ENCODEURL(\"a b&c/ñ\")",
    "_xlfn.ENCODEURL(\"-._~\")", "_xlfn.ENCODEURL(\"€\")", "_xlfn.ENCODEURL(\"\")", "_xlfn.ENCODEURL(\"100%\")", "_xlfn.ENCODEURL(\"a+b=c?d#e\")",
    "_xlfn.ENCODEURL(123.5)", "_xlfn.ENCODEURL(\"日本\")", "_xlfn.ENCODEURL(\"a,b;c:d@e!f*g(h)i'j\")", "_xlfn.ENCODEURL(\"😀\")",
    "_xlfn.ENCODEURL(\"\\\\[]{}|<>^`\")", "_xlfn.ENCODEURL(CHAR(10))", "ROT13(\"Hello, World!\")", "ROT13(\"abcxyzABCXYZ123\")",
    "ROT13(\"ñ\")", "ROT13(\"\")", "ROT13(123)",

    "INDEX(_xlfn._xlws.SORT(Data!E1:E3),1)", "INDEX(_xlfn._xlws.SORT(Data!E1:E3),2)", "INDEX(_xlfn._xlws.SORT(Data!E1:E3),3)",
    "INDEX(_xlfn._xlws.SORT(Data!E1:E3,1,-1),1)", "INDEX(_xlfn._xlws.SORT(Data!E1:E3,1,-1),2)", "INDEX(_xlfn._xlws.SORT(Data!E1:E3,1,-1),3)",
    "ISBLANK(INDEX(_xlfn._xlws.SORT(Data!E1:E3),3))", "_xlfn.TOCOL(Data!D1:D2,1)", "_xlfn.TOCOL(Data!D1:D2)", "ROWS(_xlfn.TOCOL(Data!D1:D2))",
    "_xlfn.TEXTJOIN(\",\",FALSE,_xlfn.TOCOL(Data!E1:E3))", "ROWS(_xlfn.TOCOL(Data!E1:E3,1))", "_xlfn.TOROW(Data!D1:D2,3)",
    "COUNTA(_xlfn._xlws.FILTER(Data!E1:E3,Data!E1:E3<>\"\"))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!E1:E3,Data!E1:E3=\"\"))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!A1:A3,Data!D1:D3))", "ROWS(_xlfn._xlws.FILTER(Data!A1:A3,Data!D1:D3,\"x\"))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!A1:A3,{1;\"a\";1}))",
    "_xlfn.SEQUENCE(2,2,\"a\")", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.SEQUENCE(2,2,5.5,-1.5))", "SUM(_xlfn.SEQUENCE(3,3))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TAKE(_xlfn._xlws.SORT(Data!A1:A8),3))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.DROP(_xlfn.TAKE(Data!A1:A8,5),2))",
    "_xlfn.XMATCH(\"b*\",Data!B1:B8,2,-1)", "_xlfn.XMATCH(5,Data!A1:A8,0,-1)", "_xlfn.XMATCH(\"~\",{\"a\";\"~\"},2)", "_xlfn.XMATCH(4,{5;4;3;2;1},-1,-2)",
    "_xlfn.XMATCH(4,{5;4;3;2;1},1,-2)", "_xlfn.XMATCH(4,{1;2;3;4;5},0,2)", "_xlfn.XMATCH(\"x\",{\"a\";\"b\";\"x\"},0,2)",
    "_xlfn.XMATCH(\"b\",{\"a\";\"b\";\"c\"},0,2)", "_xlfn.XMATCH(\"B\",{\"a\";\"b\";\"c\"},-1)", "_xlfn.XMATCH(\"bb\",{\"a\";\"b\";\"c\"},1)",
    "_xlfn.XLOOKUP(\"5\",Data!A1:A8,Data!B1:B8)", "_xlfn.XMATCH(\"5.0\",Data!A1:A8)", "_xlfn.XMATCH(\" 5\",Data!A1:A8)",
    "_xlfn.XMATCH(\"5\",Data!A1:A8,1)", "_xlfn.XMATCH(\"5\",Data!A1:A8,-1)", "_xlfn.XMATCH(\"5\",Data!A1:A8,2)", "_xlfn.XMATCH(\"5\",Data!A1:A8,3)",
    "_xlfn.XMATCH(5,{\"5\";\"a\"})", "_xlfn.XMATCH(\"5\",{5;6})", "_xlfn.XMATCH(\"5\",{\"a\";5},0,-1)", "_xlfn.XMATCH(\"a\",Data!A1:A8)",
    "_xlfn.XMATCH(\"1e0\",{1;2})", "_xlfn.XMATCH(TRUE,Data!A1:A8)", "MATCH(\"5\",Data!A1:A8,0)", "MATCH(\"5\",{5;6},0)", "_xlfn.XLOOKUP(\"5\",{5;6},{\"a\";\"b\"})",
    "_xlfn.XMATCH(\"5\",Data!B1:B8)", "_xlfn.XMATCH(\"apple\",Data!A1:B8)",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"aXb\",\"x\",,FALSE,2))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"aXb\",\"x\",,FALSE,-1))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TEXTSPLIT(\"aXb\",\"x\",,FALSE,5))", "_xlfn.TEXTBEFORE(\"A-b\",\"a\",1,2)", "_xlfn.TEXTBEFORE(\"A-b\",\"a\",1,-1)",
    "_xlfn.TEXTBEFORE(\"A-b\",\"a\",1,5)", "_xlfn.TEXTAFTER(\"A-b\",\"a\",1,2)", "_xlfn.TEXTBEFORE(\"a-b\",\"-\",1,0,2)", "_xlfn.TEXTBEFORE(\"a-b\",\"-\",1,0,-1)",
    "_xlfn.TEXTAFTER(\"a-b\",\"-\",1,0,2)", "_xlfn.TEXTBEFORE(\"a-b\",\"-\",1.9)", "_xlfn.TEXTBEFORE(\"a-b\",\"-\",\"x\")",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TOCOL(_xlfn.HSTACK(Data!E1:E3,Data!D1:D3),1))", "_xlfn.TEXTJOIN(\",\",FALSE,_xlfn.TOROW(_xlfn.HSTACK(Data!E1:E3,Data!D1:D3),1))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TOCOL(_xlfn.HSTACK(Data!E1:E3,Data!D1:D3),1,TRUE))",
    "ROWS(_xlfn._xlws.SORT(Data!E1:E3))", "_xlfn.TEXTJOIN(\",\",FALSE,_xlfn._xlws.SORT(_xlfn.HSTACK(Data!E1:E3,Data!A1:A3)))",
    "_xlfn._xlws.FILTER(Data!A1:A3,Data!E1:E3=10)", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!A1:A3,Data!E1:E3<>10))",
    "INDEX(_xlfn._xlws.SORT(Data!M1:M4),1)", "INDEX(_xlfn._xlws.SORT(Data!M1:M4),4)", "INDEX(_xlfn._xlws.SORT(Data!M1:M3),1)",
    "INDEX(_xlfn._xlws.SORT(_xlfn.HSTACK(Data!A1:A4,Data!M1:M4),1),1,1)", "ISERROR(INDEX(_xlfn._xlws.SORT(_xlfn.HSTACK(Data!A1:A4,Data!M1:M4),1),1,2))",
    "INDEX(_xlfn._xlws.SORT(_xlfn.HSTACK(Data!A1:A4,Data!M1:M4),1,-1),1,1)",
    "_xlfn._xlws.FILTER(Data!A1:A4,Data!M1:M4)", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!M1:M4,Data!A1:A4>3))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn._xlws.FILTER(Data!A1:A4,Data!A1:A4>3))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TOCOL(Data!M1:M4,2))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TOCOL(Data!M1:M4,3))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TOROW(Data!M1:M4,0))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TOROW(Data!M1:M4,1))",
    "COUNT(_xlfn.TOCOL(Data!M1:M4,2))", "ROWS(_xlfn.TOCOL(Data!M1:M4,2))", "ROWS(_xlfn.TOCOL(Data!M1:M4,0))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TAKE(Data!M1:M4,1))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.TAKE(Data!M1:M4,2))", "ISERROR(INDEX(_xlfn.TAKE(Data!M1:M4,2),2))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.DROP(Data!M1:M4,2))",
    "ISERROR(INDEX(_xlfn.VSTACK(Data!M1:M4,Data!A1:A2),2))", "INDEX(_xlfn.VSTACK(Data!M1:M4,Data!D1),2)", "ISERROR(INDEX(_xlfn.HSTACK(Data!M1:M4,Data!A1:A4),2,1))",
    "ISNUMBER(INDEX(_xlfn.VSTACK(Data!M1:M4,Data!D1),5))", "INDEX(_xlfn.VSTACK(Data!M1:M4,Data!D1),5)&\"|\"", "ISBLANK(INDEX(_xlfn.VSTACK(Data!M1:M4,Data!D1),5))", "_xlfn.XMATCH(3,Data!M1:M4)", "_xlfn.XMATCH(\"a\",Data!M1:M4)", "_xlfn.XMATCH(NA(),Data!M1:M4)",
    "_xlfn.XMATCH(\"A\",Data!M1:M4)",
    "ISERROR(INDEX(_xlfn._xlws.SORT(Data!M1:M4),4))", "ISBLANK(INDEX(_xlfn._xlws.SORT(Data!M1:M4),4))", "ISNA(INDEX(_xlfn._xlws.SORT(Data!M1:M4),4))",
    "_xlfn.TEXTJOIN(\",\",FALSE,_xlfn._xlws.SORT(Data!M1:M4))", "ROWS(_xlfn._xlws.SORT(Data!M1:M4))", "COUNT(_xlfn._xlws.SORT(Data!M1:M4))", 
    "ISERROR(INDEX(_xlfn._xlws.SORT(Data!M1:M4,1,-1),1))", "ISERROR(INDEX(_xlfn._xlws.SORT(Data!M1:M4,1,-1),4))", "INDEX(_xlfn._xlws.SORT(Data!M1:M4,1,-1),1)",
    "INDEX(_xlfn._xlws.SORT(Data!M1:M4,1,-1),2)",
    "INDEX(_xlfn._xlws.SORT(Data!E1:E3),3)+1", "INDEX(_xlfn._xlws.SORT(Data!E1:E3),3)&\"x\"", "ISTEXT(INDEX(_xlfn._xlws.SORT(Data!E1:E3),3))",
    "ISNUMBER(INDEX(_xlfn._xlws.SORT(Data!E1:E3),3))", "LEN(INDEX(_xlfn._xlws.SORT(Data!E1:E3),3))", "COUNTA(_xlfn._xlws.SORT(Data!E1:E3))",
    "COUNTBLANK(_xlfn._xlws.SORT(Data!E1:E3))", "INDEX(_xlfn._xlws.SORT(Data!E1:E3),3)=0", "INDEX(_xlfn._xlws.SORT(Data!E1:E3),3)=\"\"",
    "INDEX(_xlfn.TAKE(Data!E1:E3,3),3)+1", "INDEX(_xlfn.TAKE(Data!E1:E3,3),3)&\"x\"", "ISBLANK(INDEX(_xlfn.TAKE(Data!E1:E3,3),3))",
    "INDEX(_xlfn.TAKE(Data!E1:E3,3),3)", "COUNTA(_xlfn.TAKE(Data!E1:E3,3))", "ISTEXT(INDEX(_xlfn.TAKE(Data!E1:E3,3),3))",
    "INDEX(_xlfn.VSTACK(Data!E1:E3),3)", "ISBLANK(INDEX(_xlfn.VSTACK(Data!E1:E3),3))", "INDEX(_xlfn.TOCOL(Data!E1:E3),3)", "ISBLANK(INDEX(_xlfn.TOCOL(Data!E1:E3),3))",
    "INDEX(_xlfn.CHOOSEROWS(Data!E1:E3,3),1)", "INDEX(_xlfn._xlws.FILTER(Data!E1:E3,Data!A1:A3>0),3)", "ISBLANK(INDEX(_xlfn._xlws.FILTER(Data!E1:E3,Data!A1:A3>0),3))",
    "INDEX(_xlfn._xlws.SORT(_xlfn.HSTACK(Data!A1:A4,Data!M1:M4),1),2,2)", "ISERROR(INDEX(_xlfn._xlws.SORT(_xlfn.HSTACK(Data!A1:A4,Data!M1:M4),1),2,2))",
    "ISNA(INDEX(_xlfn._xlws.SORT(_xlfn.HSTACK(Data!A1:A4,Data!M1:M4),1),2,2))", "INDEX(_xlfn.SORTBY(Data!M1:M4,Data!A1:A4),2)", "ISERROR(INDEX(_xlfn.SORTBY(Data!M1:M4,Data!A1:A4),2))",
    "MATCH(\"5\",Data!A1:A8,1)", "MATCH(\"5\",Data!A1:A8,-1)", "MATCH(\"5\",Data!A1:A8)", "VLOOKUP(\"5\",Data!A1:B8,2,FALSE)", "VLOOKUP(\"5\",Data!A1:B8,2,TRUE)",
    "HLOOKUP(\"20\",Data!E2:G2,1,FALSE)", "LOOKUP(\"5\",Data!A1:A8)", "COUNTIF(Data!A1:A8,\"5\")", "_xlfn.XLOOKUP(\"5\",Data!A1:A8,Data!B1:B8,\"nf\",1)",
    "_xlfn.XLOOKUP(\"5\",Data!A1:A8,Data!B1:B8,\"nf\",2)", "_xlfn.XLOOKUP(\"5.0\",Data!A1:A8,Data!B1:B8,\"nf\")", "MATCH(\"5.0\",Data!A1:A8,0)",
    "MATCH(\"*\",Data!A1:A8,0)", "MATCH(\"?\",Data!A1:A8,0)", "VLOOKUP(\"*\",Data!A1:B8,2,FALSE)", "MATCH(\"5\",Data!K1:K8,0)"
]


def _text_rows():
    q = lambda f: f.replace('"', '""')
    rows = [f'TEXT({v},"{q(f)}")' for f in _NUM_FORMATS for v in _NUMS]
    rows += [f'TEXT({v},"{q(f)}")' for f in _DATE_FORMATS for v in _DATES]
    rows += ['TEXT("abc","0;0;0;""t:""@")', 'TEXT("abc","@@")', 'TEXT("abc","0.00")']
    return rows


FORMULAS += [
    # LibreOffice 26.8 on Linux: CHAR above 127 is U+FFFD, CODE the first UTF-8 byte
    "CHAR(65)", "CHAR(127)", "CHAR(128)", "CHAR(200)", "CHAR(255)", "CHAR(256)", "CHAR(65.9)", "CHAR(-1)",
    'CODE("A")', 'CODE("€")', 'CODE("ñ")', 'CODE("漢")', 'CODE("😀")', 'CODE("")', 'CODE("ab")', 'CODE(CHAR(200))',
    "CODE(1)", "CODE(TRUE)", "LEN(CHAR(0))", '_xlfn.UNICODE("€")', "_xlfn.UNICHAR(8364)",
    # years up to 32767
    "DATE(10000,1,1)", "DATE(12345,1,1)", "DATE(32767,12,31)", "DATE(32768,1,1)", "YEAR(3815019)", "YEAR(2958466)",
    'TEXT(3815019,"yyyy-mm-dd")', 'DATEVALUE("3/1/12345")', 'DATEVALUE("3/1/32767")', 'DATEVALUE("3/1/32768")', 'DATEVALUE("12345-3-1")',
    'TEXT(45367,"[$-en-US]dddd")', 'TEXT(45367,"[$-ZZZ]dddd")', 'TEXT(45367,"[$-]dddd")', 'TEXT(45367,"[$-409]dddd")', 'TEXT(45367,"[$-0409]dddd")',
    'TEXT(45367,"[$$-en-US]#,##0")', 'TEXT(45367,"[$-en-US]#,##0")',
]
FORMULAS += _text_rows()

# Dates with a month name, read by LibreOffice's input scanner: what it accepts depends on the
# locale's date order and words (see `parse_named_date` in sheetrs-engine/src/coerce.rs).
_NAMED_DATES_EN = [
    "Mar 1 2024", "Mar 1, 2024", "March 1, 2024", "March 1 2024", "Mar. 1, 2024", "Mar. 1 2024",
    "1-Mar-2024", "1-Mar-24", "2024-Mar-1", "Mar 2024", "March 2024", "Mar-2024", "MAR 1 2024", "mar 1 2024",
    "  Mar 1 2024  ", "Mar  1  2024", "Mar 1 2024 Monday", "Monday Mar 1 2024", "Mar 31 2024", "Mar 12 12",
    "Mar 1 24", "Mar/1/2024", "Mar 1 2024 10:30", "Mar 1, 2024 10:30",
    "1 Mar 2024", "1 March 2024", "1/Mar/2024", "2024 Mar", "2024 Mar 1", "Mar-1-2024", "1 Mar, 2024",
    "Mar 32 2024", "Feb 30 2024", "29 Feb 2023", "1st Mar 2024", "1 Mar 2024 x", "1 de marzo de 2024", "1 Mar 24",
]
_NAMED_DATES_ES = [
    "1 marzo 2024", "marzo 1 2024", "1 mar 2024", "1 mar. 2024", "1-mar-2024", "1-Mar-24", "2024-mar-1", "mar 2024",
    "marzo 2024", "Mar 1 2024", "Mar 24", "1 Mar 24", "Mar 1 24", "MAR 1 2024", "  1 marzo 2024 ",
    "1 marzo 2024 10:30", "31 mar 2024", "Mar/1/2024", "marzo 1, 2024", "1 de marzo de 2024", "March 1 2024",
    "March 2024", "1/Mar/2024", "2024 mar 1", "mar-1-2024", "32 mar 2024", "30 feb 2024",
]
FORMULAS += [f'DATEVALUE("{d}")' for d in _NAMED_DATES_EN]

# Dates, math and information functions of OpenOffice and LibreOffice (batch 6).
BATCH6 = [
    "WEEKS(45292,45299,0)", "WEEKS(45292,45299,1)", "WEEKS(45292,45298,0)", "WEEKS(45299,45292,0)", "WEEKS(45299,45292,1)",
    "WEEKS(45292,45292,0)", "WEEKS(45292,45292,1)", "WEEKS(45293,45300,1)", "WEEKS(45292.9,45299.1,0)", "WEEKS(45292,45499,1)",
    "WEEKS(45292,45299,2)", "WEEKS(45292,45299,-1)", "WEEKS(45292,45299)", "WEEKS(\"2024-01-01\",\"2024-01-15\",0)", "WEEKS(45292,45299,TRUE)",
    "WEEKS(0,7,0)", "WEEKS(-1,7,0)", "WEEKS(45292,45299,\"x\")", "WEEKS(45290,45299,1)", "WEEKS(45291,45298,1)", "WEEKS(45297,45298,1)",
    "MONTHS(45292,45323,0)", "MONTHS(45292,45322,0)", "MONTHS(45292,45323,1)", "MONTHS(45323,45292,0)", "MONTHS(45323,45292,1)",
    "MONTHS(45306,45337,0)", "MONTHS(45306,45337,1)", "MONTHS(45322,45323,1)", "MONTHS(45322,45351,0)", "MONTHS(45322,45352,0)",
    "MONTHS(45322,45351,1)", "MONTHS(45292,45292,0)", "MONTHS(45292,45292,1)", "MONTHS(45292.5,45323.5,0)", "MONTHS(45291,45292,1)",
    "MONTHS(45292,45657,0)", "MONTHS(45292,45657,1)", "MONTHS(45292,45323,2)", "MONTHS(45292,45323)", "MONTHS(45322,45351,0)",
    "MONTHS(DATE(2024,1,31),DATE(2024,2,29),0)", "MONTHS(DATE(2024,1,31),DATE(2024,3,1),0)", "MONTHS(DATE(2024,2,29),DATE(2025,2,28),0)",
    "MONTHS(DATE(2024,2,29),DATE(2025,2,28),1)", "MONTHS(DATE(2024,2,29),DATE(2025,3,1),0)",
    "YEARS(45292,45658,0)", "YEARS(45292,45657,0)", "YEARS(45292,45658,1)", "YEARS(45658,45292,0)", "YEARS(45658,45292,1)",
    "YEARS(DATE(2024,2,29),DATE(2025,2,28),0)", "YEARS(DATE(2024,2,29),DATE(2025,3,1),0)", "YEARS(DATE(2024,2,29),DATE(2028,2,29),0)",
    "YEARS(DATE(2024,2,29),DATE(2025,2,28),1)", "YEARS(DATE(2024,12,31),DATE(2025,1,1),1)", "YEARS(45292,45292,0)", "YEARS(45292,45292,1)",
    "YEARS(45292,46388,0)", "YEARS(45292,46752,1)", "YEARS(45292,45658,2)", "YEARS(45292,45658)", "YEARS(45292.5,45658.5,0)",
    "ISLEAPYEAR(45292)", "ISLEAPYEAR(45658)", "ISLEAPYEAR(DATE(1900,1,1))", "ISLEAPYEAR(DATE(2000,6,1))", "ISLEAPYEAR(DATE(2100,6,1))",
    "ISLEAPYEAR(0)", "ISLEAPYEAR(1)", "ISLEAPYEAR(-1)", "ISLEAPYEAR(60)", "ISLEAPYEAR(\"2024-03-01\")", "ISLEAPYEAR(TRUE)", "ISLEAPYEAR(\"x\")",
    "ISLEAPYEAR(2958465)", "ISLEAPYEAR(2958466)", "ISLEAPYEAR(45292.9)",
    "DAYSINMONTH(45292)", "DAYSINMONTH(DATE(2024,2,10))", "DAYSINMONTH(DATE(2023,2,10))", "DAYSINMONTH(DATE(1900,2,10))", "DAYSINMONTH(DATE(2024,4,30))",
    "DAYSINMONTH(0)", "DAYSINMONTH(-1)", "DAYSINMONTH(\"2024-02-01\")", "DAYSINMONTH(\"x\")", "DAYSINMONTH(2958465)", "DAYSINMONTH(45292.9)", "DAYSINMONTH(TRUE)",
    "DAYSINYEAR(45292)", "DAYSINYEAR(45658)", "DAYSINYEAR(DATE(1900,3,1))", "DAYSINYEAR(0)", "DAYSINYEAR(-1)", "DAYSINYEAR(\"2024-02-01\")",
    "WEEKSINYEAR(45292)", "WEEKSINYEAR(DATE(2020,6,1))", "WEEKSINYEAR(DATE(2015,6,1))", "WEEKSINYEAR(DATE(2026,6,1))", "WEEKSINYEAR(DATE(2021,6,1))",
    "WEEKSINYEAR(DATE(2004,6,1))", "WEEKSINYEAR(DATE(1900,6,1))", "WEEKSINYEAR(DATE(2099,6,1))", "WEEKSINYEAR(0)", "WEEKSINYEAR(\"2020-06-01\")", "WEEKSINYEAR(-5)",
    "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(1234.5678,2)", "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(1234.5678,6)", "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(0.00123456,3)",
    "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(-1234.5678,2)", "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(1234.5678,0)", "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(1234.5678,-1)",
    "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(0,3)", "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(1234.5678,1)", "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(1234.5678,15)",
    "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(1234.5678,20)", "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(2.5,1)", "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(0.15,1)",
    "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(1E300,2)", "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(1E-300,2)", "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(9.99,2)", "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(1234.5678,2.9)",
    "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(\"x\",2)", "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(5,\"2\")", "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(1.005,3)", "_xlfn.ORG.LIBREOFFICE.ROUNDSIG(123456789,3)",
    "SERIESSUM(2,0,1,{1,2,3})", "SERIESSUM(2,1,1,{1,2,3})", "SERIESSUM(2,0,2,{1,2,3})", "SERIESSUM(0,0,1,{1,2,3})", "SERIESSUM(0,-1,1,{1,2,3})",
    "SERIESSUM(2,0.5,1,{1,2})", "SERIESSUM(2,0,1,5)", "SERIESSUM(2,0,1,Data!A1:A3)", "SERIESSUM(2,0,1,Data!B1:B3)", "SERIESSUM(\"x\",0,1,{1})",
    "SERIESSUM(2,0,1,{1;2;3})", "SERIESSUM(2,0,1,{1,\"a\",3})", "SERIESSUM(2,0,1,{1,TRUE,3})", "SERIESSUM(2,0,1,Data!D1:D3)", "SERIESSUM(2,0,1,{1,2;3,4})",
    "SERIESSUM(-2,0,0.5,{1,2,3})", "SERIESSUM(2,-1,-1,{1,2,3})", "SERIESSUM(1,1,0,{1,2,3})", "SERIESSUM(2,0,1,Data!M1:M4)", "SERIESSUM(0,0,0,{1,2})",
    "_xlfn.ORG.LIBREOFFICE.RAWSUBTRACT(5,3)", "_xlfn.ORG.LIBREOFFICE.RAWSUBTRACT(5,3,1)", "_xlfn.ORG.LIBREOFFICE.RAWSUBTRACT(0.3,0.1,0.2)",
    "_xlfn.ORG.LIBREOFFICE.RAWSUBTRACT(1,0.9)", "_xlfn.ORG.LIBREOFFICE.RAWSUBTRACT(5)", "_xlfn.ORG.LIBREOFFICE.RAWSUBTRACT(\"x\",1)",
    "_xlfn.ORG.LIBREOFFICE.RAWSUBTRACT(Data!A1:A3,1)", "_xlfn.ORG.LIBREOFFICE.RAWSUBTRACT(1E16,1)", "1-0.9",
    "_xlfn.ORG.LIBREOFFICE.RAWSUBTRACT(0.1+0.2,0.3)", "(0.1+0.2)-0.3", "_xlfn.ORG.LIBREOFFICE.RAWSUBTRACT(1E-20,1E-21)", "_xlfn.ORG.LIBREOFFICE.RAWSUBTRACT(TRUE,1)",
    "ERRORTYPE(NA())", "ERRORTYPE(1/0)", "ERRORTYPE(Data!A1)", "ERRORTYPE(5)", "ERRORTYPE(\"x\")",
    "ERRORTYPE(Data!M2)", "ERRORTYPE(#NULL!)", "ERRORTYPE(#NUM!)", "ERRORTYPE(#DIV/0!)", "ERRORTYPE(#N/A)", "ERRORTYPE(ACOS(2))", "ERRORTYPE(1E308*10)", "ERRORTYPE(INDEX(Data!A1:A3,9))",
    "ERRORTYPE(MATCH(99,Data!A1:A3,0))", "ERRORTYPE(\"\")", "ERRORTYPE(TRUE)", "ERRORTYPE(DATE(2024,13,1))", "ERRORTYPE(INDIRECT(\"zz\"))", "ERRORTYPE(IFERROR(1/0,2))", "ERRORTYPE(ISNA(1))",
    "AREAS(Data!A1:B2)", "AREAS(Data!A1)", "AREAS(Data!A:A)", "AREAS(5)", "AREAS(Data!A1:A3)+AREAS(Data!A1:A2)", "AREAS(OFFSET(Data!A1,0,0,2,2))", "AREAS(INDEX(Data!A1:B4,1,1))",
    "AREAS(INDEX(Data!A1:B4,0,1))", "AREAS(\"x\")", "AREAS(Data!1:1)", "AREAS(NA())", "AREAS(INDIRECT(\"Data!A1:A3\"))", "AREAS(Data!A1:A3,Data!B1)",
    "_xlfn.SHEET()", "_xlfn.SHEET(Data!A1)", "_xlfn.SHEET(\"Data\")", "_xlfn.SHEET(\"nope\")", "_xlfn.SHEET(Data!A1:B2)", "_xlfn.SHEET(1)", "_xlfn.SHEET(\"data\")", "_xlfn.SHEET(\"Data!A1\")", "_xlfn.SHEET(NA())", "_xlfn.SHEET(\"\")",
    "_xlfn.SHEETS()", "_xlfn.SHEETS(Data!A1)", "_xlfn.SHEETS(Data!A1:B2)", "_xlfn.SHEETS(\"Data\")", "_xlfn.SHEETS(1)", "_xlfn.SHEETS(Data!A1:A3)+_xlfn.SHEETS()", "_xlfn.SHEETS(\"nope\")", "_xlfn.SHEETS(NA())",
    "_xlfn.SHEET(Data!A1:A3)+_xlfn.SHEET(\"Data\")", "_xlfn.SHEET(Data!M2)", "_xlfn.SHEET(INDIRECT(\"Data!A1\"))", "_xlfn.SHEET(OFFSET(Data!A1,0,0))", "_xlfn.SHEETS(INDIRECT(\"Data!A1\"))",
    "_xlfn.SHEET(\"Data!A1:B2\")", "_xlfn.SHEET(TRUE)", "SERIESSUM(0,1,1,{1,2,3})", "SERIESSUM(0,2,1,{1,2,3})", "SERIESSUM(0,0.5,1,{1,2})", "SERIESSUM(0,-2,1,{1,2,3})", "SERIESSUM(0,1,-1,{1,2,3})",
    "SERIESSUM(0,0,-1,{1,2,3})", "SERIESSUM(0,-1,1,{5})", "SERIESSUM(0,-1,-1,{1,2})", "SERIESSUM(0,0,1,{0})", "SERIESSUM(0,0,1,5)", "SERIESSUM(0,1,1,{0,0,5})",
    "SERIESSUM(0,0,0,{0})", "SERIESSUM(0,-1,0,{1,2})", "SERIESSUM(0,3,0,{1,2})", "SERIESSUM(0,-1,1,Data!D1:D3)", "SERIESSUM(0,0,1,Data!D1:D3)",
    "SERIESSUM(2,0,1,Data!A1:B3)", "SERIESSUM(2,0,1,\"3\")", "SERIESSUM(2,0,1,\"x\")", "SERIESSUM(2,0,1,TRUE)", "SERIESSUM(2,0,1,Data!H1)", "SERIESSUM(2,0,1,NA())",
    "SERIESSUM(Data!A1:A2,0,1,{1})", "SERIESSUM(2,Data!A1:A2,1,{1})", "SERIESSUM(2,0.5,0.5,{1,2,3})", "SERIESSUM(10,0,1,{1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1})",
    "MONTHS(DATE(2024,3,1),DATE(2024,1,31),0)", "MONTHS(DATE(2025,2,28),DATE(2024,2,29),0)", "YEARS(DATE(2025,2,28),DATE(2024,2,29),0)",
    "YEARS(DATE(2025,3,1),DATE(2024,2,29),1)", "MONTHS(DATE(2024,3,31),DATE(2024,2,29),0)", "MONTHS(DATE(2024,2,29),DATE(2024,3,31),0)", "WEEKS(45298,45292,0)",
    "WEEKS(45292,45298,1)", "WEEKS(45298,45292,1)", "WEEKS(45290,45298,1)", "WEEKS(45298,45290,1)", "MONTHS(DATE(2024,2,29),DATE(2024,3,29),0)", "MONTHS(DATE(2024,3,30),DATE(2024,2,29),0)",
    "MONTHS(DATE(2024,1,31),DATE(2024,2,28),0)", "MONTHS(DATE(2023,12,31),DATE(2024,2,29),0)", "YEARS(DATE(2024,3,1),DATE(2024,2,29),0)", "YEARS(DATE(2020,2,29),DATE(2024,2,28),0)",
    "WEEKS(45292,45299,0.5)", "WEEKS(45292,45299,1.9)", "MONTHS(45292,45323,0.5)", "YEARS(45292,45658,1.5)", "WEEKS(45292,45299,\"1\")", "WEEKS(\"x\",45299,0)", "WEEKS(45292,45299,NA())", "WEEKS(2958465,2958465,1)",
    "WEEKS(45292,2958466,0)", "WEEKS(-1000000,45292,0)", "MONTHS(0,45292,1)", "MONTHS(-1,1,1)", "WEEKS(-5,5,1)", "YEARS(-1,1,1)", "WEEKS(\"\",5,1)", "WEEKS(Data!D1,7,0)",
    "ISLEAPYEAR(10000000)", "ISLEAPYEAR(Data!D1)", "DAYSINMONTH(Data!D1)", "DAYSINYEAR(2958466)", "WEEKSINYEAR(2958466)",
    "ISLEAPYEAR(Data!A1:A3)", "ISLEAPYEAR(NA())", "WEEKSINYEAR(1)", "WEEKSINYEAR(60)", "WEEKSINYEAR(DATE(2009,6,1))", "WEEKSINYEAR(DATE(2010,6,1))", "WEEKSINYEAR(DATE(2016,6,1))", "WEEKSINYEAR(DATE(2032,6,1))", "WEEKSINYEAR(DATE(2037,6,1))", "WEEKSINYEAR(DATE(1976,6,1))",
    "HYPERLINK(\"http://x\",\"label\")", "HYPERLINK(\"http://x\")", "HYPERLINK(\"http://x\",5)", "HYPERLINK(\"http://x\",Data!A1)", "HYPERLINK(5,6)", "HYPERLINK(\"\",\"\")",
    "HYPERLINK(Data!B1)", "HYPERLINK(\"u\",NA())", "HYPERLINK(NA(),\"x\")", "HYPERLINK(\"a\",TRUE)", "HYPERLINK(Data!A1:A3,\"t\")", "HYPERLINK(\"u\",Data!D1)", "HYPERLINK(\"u\",\"\")", "LEN(HYPERLINK(\"u\"))",
]
FORMULAS += BATCH6

# Database functions (batch 7): the table is Data!O1:R9, the criteria blocks follow in T..AO.
_DB = "Data!$O$1:$R$9"
_DB_CRITERIA = {
    "age30": "Data!$T$1:$T$2", "and": "Data!$V$1:$W$2", "or": "Data!$Y$1:$Y$3", "orand": "Data!$AA$1:$AB$3",
    "all": "Data!$AD$1:$AD$2", "nohdr": "Data!$AF$1:$AF$2", "two": "Data!$AH$1:$AI$2", "emptytext": "Data!$AK$1:$AK$2",
    "eqnum": "Data!$AM$1:$AM$2", "ci": "Data!$AO$1:$AO$2", "star": "Data!$AQ$1:$AQ$2", "neq": "Data!$AS$1:$AS$2",
    "blankrow": "Data!$AU$1:$AU$3", "nonblank": "Data!$AW$1:$AW$2", "wide": "Data!$AY$1:$BA$3", "gus": "Data!$BC$1:$BC$2", "eve": "Data!$BE$1:$BE$2",
}
BATCH7 = []
for _fn in ["DSUM", "DCOUNT", "DCOUNTA", "DAVERAGE", "DGET", "DMAX", "DMIN", "DPRODUCT", "DSTDEV", "DSTDEVP", "DVAR", "DVARP"]:
    for _field in ['"Age"', "3", '"Score"']:
        for _k, _c in _DB_CRITERIA.items():
            BATCH7.append(f"{_fn}({_DB},{_field},{_c})")
    for _field in ['"age"', "0", "5", '"Nope"', "-1", "2.9", '"2"', "TRUE", "Data!$P$1", '"Name"', "Data!$O$2", '""', "NA()"]:
        if _field == "0" and _fn in ("DCOUNT", "DCOUNTA"):
            continue
        BATCH7.append(f"{_fn}({_DB},{_field},{_DB_CRITERIA['age30']})")
        BATCH7.append(f"{_fn}({_DB},{_field},{_DB_CRITERIA['all']})")
    # Not a table or criteria block.
    BATCH7 += [f"{_fn}({_DB},3,Data!$T$2:$T$3)", f"{_fn}({_DB},3,Data!$T$1)", f"{_fn}(Data!$O$2:$R$9,3,{_DB_CRITERIA['age30']})",
               f"{_fn}({_DB},3,5)", f"{_fn}(5,3,{_DB_CRITERIA['age30']})", f"{_fn}({_DB},3,{_DB_CRITERIA['age30']},1)", f"{_fn}({_DB},3)",
               f"{_fn}(Data!$O$1:$O$9,1,{_DB_CRITERIA['age30']})", f"{_fn}({_DB},3,Data!$AU$1:$AU$2)", f"{_fn}({_DB},3,Data!$T$1:$U$1)"]
FORMULAS += BATCH7

# Regression and frequency functions (batch 8). X2 is Data!BG1:BG8, X3 the collinear Data!BH1:BH8.
BATCH8 = []
def _cells(expr, rows, cols):
    return [f"INDEX({expr},{r},{c})" for r in range(1, rows + 1) for c in range(1, cols + 1)]
for _fn in ["LINEST", "LOGEST"]:
    for _args, _r, _c in [
        ("Data!K1:K8,Data!A1:A8", 1, 2), ("Data!K1:K8,Data!A1:A8,TRUE,TRUE", 5, 2), ("Data!K1:K8,Data!A1:A8,FALSE,TRUE", 5, 2),
        ("Data!K1:K8,Data!A1:A8,FALSE", 1, 2), ("Data!K1:K8", 1, 2), ("Data!K1:K8,,TRUE,TRUE", 5, 2),
        ("Data!K1:K8,Data!A1:A8,TRUE,FALSE", 1, 2), ("Data!K1:K8,Data!A1:A8,0,1", 5, 2),
        ("Data!K1:K8,Data!BG1:BG8", 1, 2), ("Data!K1:K8,Data!A1:A8", 1, 2),
        ("Data!K1:K8,Data!BF1:BG8,TRUE,TRUE", 5, 3), ("Data!K1:K8,Data!BF1:BG8", 1, 3), ("Data!K1:K8,Data!BF1:BG8,FALSE,TRUE", 5, 3),
        ("Data!K1:K8,Data!BF1:BH8,TRUE,TRUE", 5, 4), ("Data!K1:K8,Data!BG1:BH8,TRUE,TRUE", 5, 3),
        ("Data!K1:K8,Data!A1:A8,TRUE,TRUE", 5, 2),
        ("TRANSPOSE(Data!K1:K8),TRANSPOSE(Data!A1:A8),TRUE,TRUE", 5, 2), ("TRANSPOSE(Data!K1:K8),Data!BF1:BG8", 1, 3),
        ("Data!K1:K4,Data!A1:A4,TRUE,TRUE", 5, 2), ("Data!K1:K2,Data!A1:A2,TRUE,TRUE", 5, 2), ("Data!K1:K3,Data!A1:A3,TRUE,TRUE", 5, 2),
        ("Data!K1:K1,Data!A1:A1,TRUE,TRUE", 5, 2), ("Data!K1:K8,Data!A1:A7", 1, 2), ("Data!K1:K8,Data!A1:A9", 1, 2),
        ("Data!K1:K8,Data!A1:A8,TRUE,\"x\"", 5, 2), ("Data!K1:K8,Data!A1:A8,\"x\"", 1, 2), ("Data!K1:K8,Data!A1:A8,NA()", 1, 2),
        ("Data!B1:B8,Data!A1:A8", 1, 2), ("Data!K1:K8,Data!B1:B8", 1, 2), ("Data!K1:K8,Data!M1:M4", 1, 2), ("Data!M1:M4,Data!K1:K4", 1, 2),
        ("Data!K1:K8,Data!D1:D8", 1, 2), ("Data!D1:D8,Data!A1:A8", 1, 2), ("{1,2,3,4},{2,4,5,9}", 1, 2), ("{1;2;3;4},{2;4;5;9},TRUE,TRUE", 5, 2),
        ("{3;3;3;3},{2;4;5;9},TRUE,TRUE", 5, 2), ("Data!K1:K8,{1;1;1;1;1;1;1;1},TRUE,TRUE", 5, 2), ("{2;4;5;9},{3;3;3;3}", 1, 2),
        ("{1;2;4;8},{1;2;3;4},TRUE,TRUE", 5, 2), ("{0;1;2;3},{1;2;3;4}", 1, 2), ("{-1;2;4;8},{1;2;3;4}", 1, 2), ("{1;2;4;8},{1;2;3;4},TRUE,TRUE", 5, 2),
        ("Data!K1:K8,Data!A1:B8", 1, 3), ("Data!A1:B8,Data!K1:K8", 1, 2),
        ("Data!K1:K3,Data!BF1:BG3,TRUE,TRUE", 5, 3), ("Data!K1:K2,Data!BF1:BG2,TRUE,TRUE", 5, 3), ("Data!K1:K3,Data!BF1:BG3", 1, 3),
        ("Data!K1:K2,Data!BF1:BG2,FALSE,TRUE", 5, 3), ("{0;1;2;3},{1;2;3;4},TRUE,TRUE", 5, 2), ("{1;2;3},{1;2;3},TRUE,TRUE", 5, 2),
        ("Data!K1:K8,TRANSPOSE(Data!A1:A8),TRUE,TRUE", 5, 2), ("TRANSPOSE(Data!K1:K8),Data!A1:A8,TRUE,TRUE", 5, 2),
        ("TRANSPOSE(Data!K1:K8),TRANSPOSE(Data!BF1:BG8),TRUE,TRUE", 5, 3), ("TRANSPOSE(Data!K1:K8),TRANSPOSE(Data!BF1:BG8),FALSE", 1, 3),
        ("TRANSPOSE(Data!K1:K8)", 1, 2), ("Data!K1:K8,Data!A1:A8,TRUE,2", 5, 2), ("Data!K1:K8,Data!A1:A8,2,TRUE", 5, 2), ("Data!K1:K8,Data!A1:A8,-1,-1", 5, 2),
        ("Data!K1:K8,Data!A1:A8,,TRUE", 5, 2), ("Data!K1:K8,Data!A1:A8,TRUE,", 1, 2), ("Data!K1:K8,Data!A1:A8,TRUE,TRUE,1", 5, 2), ("Data!K1:K8,Data!A1:A8,Data!D1,TRUE", 5, 2),
        ("Data!K1:K8,Data!A1:A8,TRUE,Data!D1", 5, 2), 
        ("Data!K1:K8,Data!BF1:BG8,FALSE", 1, 3), ("Data!K1:K8,Data!BF1:BG8,FALSE,FALSE", 1, 3),
    ]:
        if _fn == "LOGEST" and _args == "{1;2;4;8},{1;2;3;4},TRUE,TRUE":
            continue  # a perfect fit in the logarithms: the statistics are rounding noise
        BATCH8 += _cells(f"{_fn}({_args})", _r, _c)
        BATCH8 += [f"ROWS({_fn}({_args}))", f"COLUMNS({_fn}({_args}))"]
for _fn in ["TREND", "GROWTH"]:
    for _args in [
        "Data!K1:K8,Data!A1:A8,{10;11;0}", "Data!K1:K8,Data!A1:A8", "Data!K1:K8,Data!A1:A8,3", "Data!K1:K8,Data!A1:A8,3,FALSE", "Data!K1:K8", "Data!K1:K8,,{9;10}",
        "Data!K1:K8,Data!A1:A8,Data!A1:A3", "Data!K1:K8,Data!BF1:BG8,Data!BF1:BG2", "Data!K1:K8,Data!BF1:BG8", "Data!K1:K8,Data!BF1:BG8,{1,2}", "Data!K1:K8,Data!BF1:BG8,{1,2;3,4}",
        "Data!K1:K8,Data!BF1:BG8,{1;2}", "TRANSPOSE(Data!K1:K8),TRANSPOSE(Data!A1:A8),{1,2,3}", "TRANSPOSE(Data!K1:K8),TRANSPOSE(Data!A1:A8),{1;2;3}", "TRANSPOSE(Data!K1:K8),TRANSPOSE(Data!A1:A8)",
        "Data!K1:K8,Data!A1:A7,3", "Data!K1:K8,Data!A1:A8,Data!B1:B3", "Data!K1:K8,Data!A1:A8,\"x\"", "Data!K1:K8,Data!A1:A8,3,\"x\"", "Data!K1:K8,Data!A1:A8,NA()",
        "Data!M1:M4,Data!K1:K4,3", "Data!K1:K4,Data!M1:M4,3", "Data!K1:K8,Data!A1:A8,Data!D1:D3", "{1;2;4;8},{1;2;3;4},{5;6}", "{-1;2;4;8},{1;2;3;4},{5;6}", "{3;3;3},{1;2;3},{4}",
        "Data!K1:K8,Data!BG1:BH8,Data!BG1:BH2", "Data!K1:K8,Data!BG1:BH8,Data!BG1:BG2", "Data!K1:K8,Data!A1:A8,{10,11;12,13}", "Data!K1:K8,Data!A1:A8,TRUE",
        "Data!K1:K8,Data!A1:A8,{10;11},TRUE", "Data!K1:K8,Data!A1:A8,{10;11},FALSE", "Data!K1:K8,Data!A1:A8,{10;11},0", "Data!K1:K8,Data!A1:A8,{10;11},,",
    ]:
        BATCH8 += [f"ROWS({_fn}({_args}))", f"COLUMNS({_fn}({_args}))"] + [f"INDEX({_fn}({_args}),{r},{c})" for r in (1, 2, 3) for c in (1, 2)]
BATCH8 += [
    "ROWS(FREQUENCY(Data!A1:A8,{3;6}))", "COLUMNS(FREQUENCY(Data!A1:A8,{3;6}))", "ROWS(FREQUENCY(Data!A1:A8,{3,6}))", "COLUMNS(FREQUENCY(Data!A1:A8,{3,6}))",
] + [f"INDEX(FREQUENCY(Data!A1:A8,{{3;6}}),{r},1)" for r in (1, 2, 3, 4)]
for _d, _b in [("Data!A1:A8", "{3;6}"), ("Data!A1:A8", "{6;3}"), ("Data!A1:A8", "{0;100}"), ("Data!A1:A8", "{5}"), ("Data!A1:A8", "Data!A1:A3"), ("Data!A1:B8", "{3;6}"),
               ("Data!K1:K8", "{4;5;9}"), ("Data!K1:K8", "Data!H1:H5"), ("Data!M1:M4", "{2}"), ("Data!A1:A8", "{\"a\"}"), ("Data!D1:D8", "{1}"), ("Data!A1:A8", "Data!D1:D3"),
               ("Data!A1:A8", "Data!M1:M4"), ("{1,2,3,4,5}", "{2;4}"), ("5", "{3}"), ("Data!A1:A8", "5"), ("Data!A1:A8", "{3,6}"), ("Data!A1:A8", "{3.5;6.5}"), ("Data!A1:A8", "{-1;100}"),
               ("Data!A1:A8", "{TRUE;6}"), ("Data!A1:A8", "{\"3\";6}"), ("{1,\"a\",2}", "{1}"), ("Data!B1:B8", "{1}"), ("Data!A1:A8", "{1;\"b\";5}"), ("Data!A1:A8", "{3;6;1}")]:
    BATCH8 += [f"ROWS(FREQUENCY({_d},{_b}))", f"COLUMNS(FREQUENCY({_d},{_b}))"] + [f"INDEX(FREQUENCY({_d},{_b}),{r},1)" for r in (1, 2, 3, 4)]
BATCH8 += [
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.MODE.MULT(Data!K1:K8))", "ROWS(_xlfn.MODE.MULT(Data!K1:K8))", "COLUMNS(_xlfn.MODE.MULT(Data!K1:K8))", "_xlfn.MODE.MULT(Data!A1:A8)",
    "INDEX(_xlfn.MODE.MULT(Data!A1:A8),1)", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.MODE.MULT({1,1,2,2,3}))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.MODE.MULT({1;1;2;2;3}))",
    "ROWS(_xlfn.MODE.MULT({1,1,2,2,3}))", "COLUMNS(_xlfn.MODE.MULT({1,1,2,2,3}))", "ROWS(_xlfn.MODE.MULT({1;1;2;2;3}))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.MODE.MULT(Data!A1:B8))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.MODE.MULT(Data!K1:K8,Data!A1:A3))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.MODE.MULT(3,3,4))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.MODE.MULT(Data!M1:M4))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.MODE.MULT(Data!B1:B8))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.MODE.MULT({3,1,3,1,2}))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.MODE.MULT({1,\"a\",1}))",
    "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.MODE.MULT(Data!K1:K8,3,3))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.MODE.MULT({5}))", "_xlfn.TEXTJOIN(\",\",TRUE,_xlfn.MODE.MULT(Data!D1:D3))", "_xlfn.MODE.MULT(NA())",
]
FORMULAS += BATCH8

# Securities and depreciation (batch 10). Dates are DATE(y,m,d) so that they are the same in any system.
def _d(y, m, d):
    return f"DATE({y},{m},{d})"
_SETTLE_MATURITY = [
    ((2024, 1, 15), (2029, 6, 30)), ((2024, 3, 31), (2025, 3, 31)), ((2024, 2, 29), (2026, 8, 31)), ((2023, 12, 31), (2024, 6, 30)),
    ((2024, 5, 15), (2024, 11, 15)), ((2024, 2, 15), (2024, 2, 29)), ((2025, 1, 31), (2030, 1, 31)), ((2024, 7, 1), (2024, 7, 1)),
    ((2024, 8, 1), (2024, 7, 1)), ((2024, 6, 30), (2025, 2, 28)), ((2024, 11, 30), (2027, 2, 28)), ((2020, 2, 29), (2024, 2, 29)),
    ((2024, 1, 1), (2024, 1, 2)), ((2024, 9, 30), (2034, 3, 31)), ((2024, 4, 30), (2024, 10, 31)), ((2023, 8, 31), (2025, 2, 28)),
    ((2024, 2, 28), (2024, 8, 31)), ((2024, 3, 15), (2024, 9, 15)), ((2024, 3, 1), (2025, 9, 1)),
]
BATCH10 = []
for (_s, _m) in _SETTLE_MATURITY:
    for _freq in (1, 2, 4):
        for _basis in (0, 1, 2, 3, 4):
            for _fn in ("COUPPCD", "COUPNCD", "COUPNUM", "COUPDAYBS", "COUPDAYS", "COUPDAYSNC"):
                BATCH10.append(f"{_fn}({_d(*_s)},{_d(*_m)},{_freq},{_basis})")
    BATCH10 += [f"COUPNUM({_d(*_s)},{_d(*_m)},2)", f"COUPDAYS({_d(*_s)},{_d(*_m)},2)"]
BATCH10 += [
    "COUPNUM(DATE(2024,1,15),DATE(2029,6,30),3,0)", "COUPNUM(DATE(2024,1,15),DATE(2029,6,30),0,0)", "COUPNUM(DATE(2024,1,15),DATE(2029,6,30),2,5)",
    "COUPNUM(DATE(2024,1,15),DATE(2029,6,30),2,-1)", "COUPNUM(DATE(2024,1,15),DATE(2029,6,30),2.9,0.9)", "COUPNUM(\"2024-01-15\",\"2029-06-30\",2,0)",
    "COUPNUM(DATE(2024,1,15),DATE(2029,6,30),2,\"x\")", "COUPNUM(DATE(2024,1,15),DATE(2029,6,30),TRUE,0)", "COUPNUM(DATE(2024,1,15),DATE(2029,6,30),2,Data!D1)",
    "COUPNUM(45306.9,47300.1,2,0)", "COUPNUM(DATE(2024,1,15),NA(),2,0)", "COUPNUM(DATE(2024,1,15),DATE(2029,6,30),2,0,1)", "COUPNUM(DATE(2024,1,15),DATE(2029,6,30))",
    "COUPPCD(DATE(2024,1,15),DATE(2029,6,30),2)", "COUPPCD(0,DATE(2029,6,30),2,0)", "COUPPCD(-5,DATE(2029,6,30),2,0)", "COUPNCD(DATE(2024,1,15),DATE(9999,12,31),2,0)",
    "COUPPCD(DATE(2024,1,15),DATE(2029,6,30),2,Data!D1)", "COUPPCD(Data!D1,DATE(2029,6,30),2,0)",
]
FORMULAS += BATCH10

_BOND_PAIRS = [((2024, 1, 15), (2029, 6, 30)), ((2024, 3, 31), (2025, 3, 31)), ((2024, 2, 29), (2026, 8, 31)), ((2024, 5, 15), (2024, 11, 15)),
               ((2024, 6, 30), (2025, 2, 28)), ((2023, 8, 31), (2025, 2, 28)), ((2024, 2, 15), (2024, 2, 29)), ((2024, 1, 1), (2034, 1, 1))]
BATCH11 = []
for (_s, _m) in _BOND_PAIRS:
    for _freq in (1, 2, 4):
        for _basis in (0, 1, 2, 3, 4):
            _a = f"{_d(*_s)},{_d(*_m)}"
            BATCH11 += [
                f"PRICE({_a},0.05,0.08,100,{_freq},{_basis})", f"PRICE({_a},0.1,0.03,105,{_freq},{_basis})", f"PRICE({_a},0,0.05,100,{_freq},{_basis})",
                f"YIELD({_a},0.05,95,100,{_freq},{_basis})", f"YIELD({_a},0.1,101.5,100,{_freq},{_basis})", f"YIELD({_a},0.05,100,100,{_freq},{_basis})",
                f"DURATION({_a},0.05,0.08,{_freq},{_basis})", f"DURATION({_a},0.1,0.03,{_freq},{_basis})", f"DURATION({_a},0,0.05,{_freq},{_basis})",
                f"MDURATION({_a},0.05,0.08,{_freq},{_basis})", f"MDURATION({_a},0.1,0.03,{_freq},{_basis})",
            ]
for (_s, _m) in _BOND_PAIRS + [((2024, 7, 1), (2024, 7, 1)), ((2024, 8, 1), (2024, 7, 1)), ((2023, 12, 31), (2024, 6, 30)), ((2024, 3, 1), (2024, 9, 1)), ((2020, 2, 29), (2024, 2, 29))]:
    for _basis in (0, 1, 2, 3, 4):
        _a = f"{_d(*_s)},{_d(*_m)}"
        BATCH11 += [
            f"DISC({_a},95,100,{_basis})", f"DISC({_a},100,100,{_basis})", f"DISC({_a},99.5,100.5,{_basis})",
            f"PRICEDISC({_a},0.05,100,{_basis})", f"PRICEDISC({_a},0.1,105,{_basis})", f"YIELDDISC({_a},95,100,{_basis})", f"YIELDDISC({_a},99.5,100.5,{_basis})",
            f"RECEIVED({_a},1000,0.05,{_basis})", f"RECEIVED({_a},5000,0.12,{_basis})", f"INTRATE({_a},1000,1050,{_basis})", f"INTRATE({_a},950,1000,{_basis})",
            f"PRICEMAT({_a},{_d(2023,10,1)},0.05,0.06,{_basis})", f"PRICEMAT({_a},{_d(2024,1,1)},0.1,0.03,{_basis})", f"YIELDMAT({_a},{_d(2023,10,1)},0.05,98,{_basis})",
            f"YIELDMAT({_a},{_d(2024,1,1)},0.1,101,{_basis})", f"ACCRINTM({_d(2023,10,1)},{_d(*_m)},0.05,1000,{_basis})", f"ACCRINTM({_d(2024,1,1)},{_d(*_s)},0.08,5000,{_basis})",
        ]
    BATCH11 += [f"TBILLPRICE({_d(*_s)},{_d(*_m)},0.05)", f"TBILLYIELD({_d(*_s)},{_d(*_m)},98)", f"TBILLEQ({_d(*_s)},{_d(*_m)},0.05)", f"TBILLEQ({_d(*_s)},{_d(*_m)},0.12)"]
for _issue, _first, _settle in [((2024, 1, 1), (2024, 7, 1), (2024, 4, 15)), ((2023, 10, 31), (2024, 4, 30), (2024, 2, 29)), ((2024, 1, 31), (2024, 7, 31), (2024, 12, 31)),
                                ((2023, 6, 30), (2024, 6, 30), (2025, 1, 15)), ((2024, 2, 29), (2024, 8, 31), (2024, 5, 31)), ((2024, 1, 1), (2024, 7, 1), (2024, 7, 1))]:
    for _freq in (1, 2, 4):
        for _basis in (0, 1, 2, 3, 4):
            _a = f"{_d(*_issue)},{_d(*_first)},{_d(*_settle)}"
            BATCH11 += [f"ACCRINT({_a},0.05,1000,{_freq},{_basis})", f"ACCRINT({_a},0.1,100,{_freq},{_basis},TRUE)", f"ACCRINT({_a},0.1,100,{_freq},{_basis},FALSE)"]
BATCH11 += [
    "PRICE(DATE(2024,1,15),DATE(2029,6,30),-0.01,0.08,100,2,0)", "PRICE(DATE(2024,1,15),DATE(2029,6,30),0.05,-0.01,100,2,0)", "PRICE(DATE(2024,1,15),DATE(2029,6,30),0.05,0.08,0,2,0)",
    "PRICE(DATE(2024,1,15),DATE(2029,6,30),0.05,0.08,-1,2,0)", "PRICE(DATE(2024,1,15),DATE(2029,6,30),0.05,0.08,100,3,0)", "PRICE(DATE(2024,1,15),DATE(2029,6,30),0.05,0.08,100,2,5)",
    "PRICE(DATE(2024,1,15),DATE(2029,6,30),0.05,0.08,100,2)", "PRICE(DATE(2024,1,15),DATE(2029,6,30),0.05,0.08,100)", "PRICE(DATE(2024,1,15),DATE(2029,6,30),0.05,0,100,2,0)",
    "PRICE(\"2024-01-15\",\"2029-06-30\",0.05,0.08,100,2,0)", "PRICE(DATE(2024,1,15),DATE(2029,6,30),\"x\",0.08,100,2,0)", "PRICE(DATE(2024,1,15),DATE(2029,6,30),0.05,0.08,100,2,Data!D1)",
    "PRICE(DATE(2024,1,15),DATE(2029,6,30),0.05,0.08,100,TRUE,0)", "PRICE(DATE(2024,1,15),DATE(2029,6,30),Data!D1,0.08,100,2,0)", "PRICE(DATE(2024,1,15),DATE(2029,6,30),0.05,NA(),100,2,0)",
    "YIELD(DATE(2024,1,15),DATE(2029,6,30),0.05,0,100,2,0)", "YIELD(DATE(2024,1,15),DATE(2029,6,30),0.05,-5,100,2,0)", "YIELD(DATE(2024,1,15),DATE(2029,6,30),0.05,95,0,2,0)", "YIELD(DATE(2024,1,15),DATE(2029,6,30),-0.05,95,100,2,0)",
    "YIELD(DATE(2024,1,15),DATE(2029,6,30),0.05,1000000,100,2,0)", "YIELD(DATE(2024,1,15),DATE(2029,6,30),0.05,1,100,2,0)", "YIELD(DATE(2024,1,15),DATE(2024,3,15),0.05,99,100,2,0)", "YIELD(DATE(2024,1,15),DATE(2024,3,15),0.05,101,100,2,0)",
    "YIELD(DATE(2024,1,15),DATE(2024,3,15),0.05,99,100,1,2)", "YIELD(DATE(2024,1,15),DATE(2024,2,15),0.05,99.9,100,4,3)", "YIELD(DATE(2024,1,15),DATE(2054,6,30),0.05,60,100,2,0)", "YIELD(DATE(2024,1,15),DATE(2029,6,30),0.5,300,100,2,0)",
    "DURATION(DATE(2024,1,15),DATE(2029,6,30),-0.05,0.08,2,0)", "DURATION(DATE(2024,1,15),DATE(2029,6,30),0.05,-0.08,2,0)", "DURATION(DATE(2024,1,15),DATE(2029,6,30),0.05,0.08,3,0)", "DURATION(DATE(2024,1,15),DATE(2029,6,30),0.05,0.08,2,6)",
    "DISC(DATE(2024,1,15),DATE(2024,6,30),0,100,0)", "DISC(DATE(2024,1,15),DATE(2024,6,30),95,0,0)", "DISC(DATE(2024,1,15),DATE(2024,6,30),-95,100,0)", "DISC(DATE(2024,1,15),DATE(2024,6,30),95,100,5)", "DISC(DATE(2024,1,15),DATE(2024,6,30),95,100)",
    "PRICEDISC(DATE(2024,1,15),DATE(2024,6,30),-0.05,100,0)", "PRICEDISC(DATE(2024,1,15),DATE(2024,6,30),0.05,0,0)", "PRICEDISC(DATE(2024,1,15),DATE(2024,6,30),0,100,0)", "PRICEDISC(DATE(2024,1,15),DATE(2025,6,30),5,100,0)",
    "YIELDDISC(DATE(2024,1,15),DATE(2024,6,30),0,100,0)", "YIELDDISC(DATE(2024,1,15),DATE(2024,6,30),95,0,0)", "RECEIVED(DATE(2024,1,15),DATE(2024,6,30),0,0.05,0)", "RECEIVED(DATE(2024,1,15),DATE(2024,6,30),1000,0,0)",
    "RECEIVED(DATE(2024,1,15),DATE(2024,6,30),1000,50,0)", "INTRATE(DATE(2024,1,15),DATE(2024,6,30),0,1000,0)", "INTRATE(DATE(2024,1,15),DATE(2024,6,30),1000,0,0)", "INTRATE(DATE(2024,1,15),DATE(2024,6,30),1000,900,0)",
    "PRICEMAT(DATE(2024,1,15),DATE(2024,6,30),DATE(2024,2,1),0.05,0.06,0)", "PRICEMAT(DATE(2024,1,15),DATE(2024,6,30),DATE(2024,1,15),0.05,0.06,0)", "PRICEMAT(DATE(2024,1,15),DATE(2024,6,30),DATE(2023,1,15),-0.05,0.06,0)",
    "PRICEMAT(DATE(2024,1,15),DATE(2024,6,30),DATE(2023,1,15),0.05,-0.06,0)", "YIELDMAT(DATE(2024,1,15),DATE(2024,6,30),DATE(2023,1,15),0.05,0,0)", "YIELDMAT(DATE(2024,1,15),DATE(2024,6,30),DATE(2023,1,15),-0.05,95,0)",
    "ACCRINTM(DATE(2024,1,15),DATE(2024,6,30),0,1000,0)", "ACCRINTM(DATE(2024,1,15),DATE(2024,6,30),0.05,0,0)", "ACCRINTM(DATE(2024,6,30),DATE(2024,1,15),0.05,1000,0)", "ACCRINTM(DATE(2024,1,15),DATE(2024,6,30),0.05,1000)",
    "ACCRINTM(DATE(2024,1,15),DATE(2024,6,30),0.05)", "ACCRINTM(DATE(2024,1,15),DATE(2024,1,15),0.05,1000,0)",
    "TBILLPRICE(DATE(2024,1,15),DATE(2025,1,16),0.05)", "TBILLPRICE(DATE(2024,1,15),DATE(2025,1,15),0.05)", "TBILLPRICE(DATE(2024,1,15),DATE(2024,6,30),0)", "TBILLPRICE(DATE(2024,1,15),DATE(2024,6,30),-0.05)",
    "TBILLPRICE(DATE(2024,1,15),DATE(2024,6,30),5)", "TBILLPRICE(DATE(2024,6,30),DATE(2024,1,15),0.05)", "TBILLYIELD(DATE(2024,1,15),DATE(2025,1,16),98)", "TBILLYIELD(DATE(2024,1,15),DATE(2024,6,30),0)",
    "TBILLYIELD(DATE(2024,1,15),DATE(2024,6,30),-5)", "TBILLYIELD(DATE(2024,1,15),DATE(2024,6,30),100)", "TBILLEQ(DATE(2024,1,15),DATE(2025,1,16),0.05)", "TBILLEQ(DATE(2024,1,15),DATE(2024,6,30),0)",
    "TBILLEQ(DATE(2024,1,15),DATE(2024,6,30),-0.05)", "TBILLEQ(DATE(2024,1,15),DATE(2024,6,30),0.5)", "TBILLEQ(DATE(2024,1,15),DATE(2024,12,30),0.5)", "TBILLEQ(DATE(2024,1,15),DATE(2024,12,30),0.05)", "TBILLEQ(DATE(2024,1,15),DATE(2024,9,30),0.08)",
    "ACCRINT(DATE(2024,1,1),DATE(2024,7,1),DATE(2024,4,15),-0.05,1000,2,0)", "ACCRINT(DATE(2024,1,1),DATE(2024,7,1),DATE(2024,4,15),0.05,0,2,0)", "ACCRINT(DATE(2024,1,1),DATE(2024,7,1),DATE(2024,4,15),0.05,-1,2,0)",
    "ACCRINT(DATE(2024,1,1),DATE(2024,7,1),DATE(2024,4,15),0.05,1000,3,0)", "ACCRINT(DATE(2024,1,1),DATE(2024,7,1),DATE(2024,4,15),0.05,1000,2,7)", "ACCRINT(DATE(2024,4,15),DATE(2024,7,1),DATE(2024,1,1),0.05,1000,2,0)",
    "ACCRINT(DATE(2024,1,1),DATE(2024,7,1),DATE(2024,4,15),0.05,1000,2)", "ACCRINT(DATE(2024,1,1),DATE(2024,7,1),DATE(2024,4,15),0.05,1000)", "ACCRINT(DATE(2024,1,1),DATE(2024,7,1),DATE(2024,4,15),0.05,1000,2,0,TRUE,1)",
    "ACCRINT(DATE(2024,1,1),DATE(2024,7,1),DATE(2024,4,15),0.05,1000,2,0,2)", "ACCRINT(DATE(2024,1,1),DATE(2024,7,1),DATE(2024,4,15),0.05,1000,2,0,0)", "ACCRINT(DATE(2024,1,1),DATE(2024,7,1),DATE(2024,4,15),0.05,1000,2,0,\"x\")",
    "ACCRINT(DATE(2024,1,1),DATE(2024,7,1),DATE(2024,4,15),0.05,1000,2,0,Data!D1)", "ACCRINT(DATE(2024,1,1),DATE(2024,7,1),DATE(2025,4,15),0.05,1000,2,0)", "ACCRINT(DATE(2024,1,1),DATE(2023,7,1),DATE(2024,4,15),0.05,1000,2,0)",
]
FORMULAS += BATCH11

# A grid of start and end dates around the ends of months for the day counts of the instruments.
_GRID_STARTS = [(2023, 12, 31), (2024, 1, 30), (2024, 1, 31), (2024, 2, 28), (2024, 2, 29), (2023, 2, 28), (2024, 5, 31), (2024, 6, 30)]
_GRID_ENDS = [(2025, 2, 28), (2025, 3, 31), (2024, 8, 31), (2025, 6, 30), (2026, 12, 31), (2027, 2, 28), (2025, 1, 30), (2024, 12, 31)]
BATCH12 = []
for _gs in _GRID_STARTS:
    for _ge in _GRID_ENDS:
        _a = f"{_d(*_gs)},{_d(*_ge)}"
        for _basis in (0, 1, 4):
            BATCH12 += [
                f"PRICEDISC({_a},0.05,100,{_basis})", f"RECEIVED({_a},1000,0.05,{_basis})", f"INTRATE({_a},1000,1050,{_basis})", f"DISC({_a},95,100,{_basis})",
                f"YIELDDISC({_a},95,100,{_basis})", f"ACCRINTM({_d(*_gs)},{_d(*_ge)},0.05,1000,{_basis})", f"PRICEMAT({_a},{_d(2023,1,31)},0.05,0.06,{_basis})",
                f"YIELDMAT({_a},{_d(2023,1,31)},0.05,98,{_basis})", f"ACCRINT({_d(*_gs)},{_d(2030,1,1)},{_d(*_ge)},0.05,1000,2,{_basis})",
                f"DURATION({_a},0.05,0.08,2,{_basis})", f"PRICE({_a},0.05,0.08,100,2,{_basis})", f"YIELD({_a},0.05,95,100,2,{_basis})",
            ]
        BATCH12 += [f"TBILLPRICE({_a},0.05)", f"TBILLEQ({_a},0.05)"]
        if not (_gs[1] == 2 and _ge[1] == 2):
            BATCH12.append(f"TBILLYIELD({_a},98)")
BATCH12 += ["TBILLPRICE(DATE(2024,1,15),DATE(2025,1,14),0.05)", "TBILLPRICE(DATE(2024,1,15),DATE(2025,1,13),0.05)", "TBILLPRICE(DATE(2023,12,31),DATE(2024,12,30),0.05)"]
FORMULAS += BATCH12

# Depreciation (batch 13) and securities with odd periods (batch 14).
BATCH13 = []
for _rate in (0.05, 0.1, 0.15, 0.2, 0.25, 0.3, 0.4, 0.5, 0.01):
    for _basis in (0, 1, 3, 4):
        for _per in (0, 1, 2, 3, 5, 8, 12, 25):
            BATCH13 += [f"AMORLINC(10000,DATE(2024,1,15),DATE(2024,12,31),500,{_per},{_rate},{_basis})",
                        f"AMORDEGRC(10000,DATE(2024,1,15),DATE(2024,12,31),500,{_per},{_rate},{_basis})"]
for _f in ("AMORLINC", "AMORDEGRC"):
    BATCH13 += [f"{_f}(10000,DATE(2024,1,15),DATE(2024,12,31),500,1,0.1,2)", f"{_f}(10000,DATE(2024,1,15),DATE(2024,12,31),500,1,0.1)", f"{_f}(10000,DATE(2024,1,15),DATE(2024,12,31),500,1,0.1,5)",
                f"{_f}(10000,DATE(2024,1,15),DATE(2024,12,31),500,-1,0.1,0)", f"{_f}(10000,DATE(2024,1,15),DATE(2024,12,31),500,1.5,0.1,0)", f"{_f}(0,DATE(2024,1,15),DATE(2024,12,31),500,1,0.1,0)",
                f"{_f}(-1,DATE(2024,1,15),DATE(2024,12,31),500,1,0.1,0)", f"{_f}(10000,DATE(2024,1,15),DATE(2024,12,31),-1,1,0.1,0)", f"{_f}(10000,DATE(2024,1,15),DATE(2024,12,31),10001,1,0.1,0)",
                f"{_f}(10000,DATE(2024,1,15),DATE(2024,12,31),500,1,0,0)", f"{_f}(10000,DATE(2024,1,15),DATE(2024,12,31),500,1,-0.1,0)", f"{_f}(10000,DATE(2024,12,31),DATE(2024,1,15),500,1,0.1,0)",
                f"{_f}(10000,DATE(2024,1,15),DATE(2024,1,15),500,1,0.1,0)", f"{_f}(10000,DATE(2024,1,15),DATE(2024,12,31),0,1,0.1,0)", f"{_f}(10000,DATE(2024,1,15),DATE(2024,12,31),10000,1,0.1,0)",
                f"{_f}(10000,DATE(2024,1,15),DATE(2024,12,31),500,1,1,0)", f"{_f}(10000,DATE(2024,1,15),DATE(2024,12,31),500,1,0.6,0)", f"{_f}(10000,DATE(2024,1,15),DATE(2024,12,31),500,1,0.1,0,1)",
                f"{_f}(10000.5,DATE(2024,1,15),DATE(2024,12,31),500.5,1,0.1,0)", f"{_f}(10000,DATE(2023,6,30),DATE(2023,12,31),500,2,0.1,1)", f"{_f}(10000,DATE(2023,6,30),DATE(2024,6,30),500,2,0.1,1)",
                f"{_f}(10000,DATE(2024,2,29),DATE(2024,12,31),500,2,0.1,1)", f"{_f}(10000,DATE(2024,2,29),DATE(2025,2,28),500,2,0.1,0)", f"{_f}(\"x\",DATE(2024,1,15),DATE(2024,12,31),500,1,0.1,0)",
                f"{_f}(100,DATE(2024,1,15),DATE(2024,12,31),1,3,0.5,0)", f"{_f}(100,DATE(2024,1,15),DATE(2024,12,31),1,0,0.5,0)", f"{_f}(100,DATE(2024,1,15),DATE(2024,12,31),1,1,0.5,0)", f"{_f}(100,DATE(2024,1,15),DATE(2024,12,31),1,2,0.5,0)"]
for _cost, _salv, _life in [(2400, 300, 10), (2400, 300, 120), (1000, 0, 5), (1000, 100, 3), (50000, 5000, 7), (100, 100, 4), (100, 150, 4), (0, 0, 5), (1000, 0, 1), (500, 50, 2.5), (1000, 100, 0.5)]:
    for _start, _end in [(0, 1), (1, 2), (0, 5), (0.5, 1.5), (2, 3), (0, _life), (_life - 1, _life), (3.5, 4.2), (0, 0), (1, 1), (2, 1), (0, _life + 1), (-1, 1), (0.25, 0.75), (_life, _life), (1, _life)]:
        for _factor in ("", ",1.5", ",2", ",3", ",0", ",-1"):
            BATCH13 += [f"VDB({_cost},{_salv},{_life},{_start},{_end}{_factor})"]
        BATCH13 += [f"VDB({_cost},{_salv},{_life},{_start},{_end},2,TRUE)", f"VDB({_cost},{_salv},{_life},{_start},{_end},1.5,TRUE)", f"VDB({_cost},{_salv},{_life},{_start},{_end},2,FALSE)", f"VDB({_cost},{_salv},{_life},{_start},{_end},,TRUE)"]
BATCH13 += ["VDB(2400,300,10,0,1,2,1)", "VDB(2400,300,10,0,1,2,0)", "VDB(2400,300,10,0,1,2,\"x\")", "VDB(2400,300,10,0,1,2,TRUE,1)", "VDB(2400,300,10,0,1)", "VDB(2400,300,10,0)", "VDB(\"x\",300,10,0,1)", "VDB(2400,300,0,0,1)", "VDB(2400,300,-1,0,1)",
            "VDB(2400,-300,10,0,1)", "VDB(-2400,300,10,0,1)", "VDB(2400,300,10,NA(),1)", "VDB(2400,300,10,0,1,Data!D1)", "VDB(2400,300,10,0,1,2,Data!D1)", "VDB(Data!A1:A2,300,10,0,1)"]
FORMULAS += BATCH13

# Securities with an odd first or last period (batch 14).
_ODDF = [((2024, 2, 15), (2027, 7, 1), (2024, 1, 1), (2024, 7, 1)), ((2023, 11, 11), (2028, 7, 1), (2023, 10, 15), (2024, 7, 1)), ((2024, 3, 1), (2026, 8, 31), (2024, 1, 31), (2024, 8, 31)),
         ((2024, 3, 20), (2025, 4, 15), (2024, 3, 15), (2024, 4, 15)), ((2024, 5, 1), (2026, 1, 15), (2024, 4, 1), (2024, 7, 15)), ((2024, 2, 29), (2027, 2, 28), (2024, 1, 15), (2024, 8, 31)),
         ((2024, 6, 15), (2029, 12, 31), (2024, 1, 1), (2024, 12, 31)), ((2024, 1, 10), (2026, 5, 20), (2023, 12, 20), (2024, 5, 20))]
_ODDL = [((2024, 3, 1), (2024, 6, 30), (2024, 1, 15)), ((2024, 2, 1), (2025, 1, 15), (2024, 1, 15)), ((2024, 3, 15), (2024, 11, 1), (2024, 1, 15)), ((2024, 2, 29), (2024, 8, 31), (2024, 2, 15)),
         ((2024, 4, 20), (2024, 7, 15), (2024, 1, 15)), ((2024, 8, 1), (2025, 8, 15), (2024, 7, 15)), ((2024, 1, 31), (2024, 7, 31), (2023, 7, 31)), ((2024, 5, 15), (2024, 5, 31), (2024, 4, 30))]
BATCH14 = []
for (_s, _m, _i, _f) in _ODDF:
    for _freq in (1, 2, 4):
        for _basis in (0, 1, 2, 3, 4):
            _a = f"{_d(*_s)},{_d(*_m)},{_d(*_i)},{_d(*_f)}"
            BATCH14 += [f"ODDFPRICE({_a},0.05,0.08,100,{_freq},{_basis})", f"ODDFPRICE({_a},0.1,0.03,105,{_freq},{_basis})", f"ODDFPRICE({_a},0,0.05,100,{_freq},{_basis})",
                        f"ODDFYIELD({_a},0.05,95,100,{_freq},{_basis})", f"ODDFYIELD({_a},0.1,101.5,100,{_freq},{_basis})", f"ODDFYIELD({_a},0.05,100,100,{_freq},{_basis})"]
for (_s, _m, _l) in _ODDL:
    for _freq in (1, 2, 4):
        for _basis in (0, 1, 2, 3, 4):
            _a = f"{_d(*_s)},{_d(*_m)},{_d(*_l)}"
            BATCH14 += [f"ODDLPRICE({_a},0.05,0.08,100,{_freq},{_basis})", f"ODDLPRICE({_a},0.1,0.03,105,{_freq},{_basis})", f"ODDLPRICE({_a},0,0.05,100,{_freq},{_basis})",
                        f"ODDLYIELD({_a},0.05,95,100,{_freq},{_basis})", f"ODDLYIELD({_a},0.1,101.5,100,{_freq},{_basis})", f"ODDLYIELD({_a},0.05,100,100,{_freq},{_basis})"]
_o = "DATE(2024,2,15),DATE(2027,7,1),DATE(2024,1,1),DATE(2024,7,1)"
BATCH14 += [f"ODDFPRICE({_o},-0.01,0.08,100,2,0)", f"ODDFPRICE({_o},0.05,-0.01,100,2,0)", f"ODDFPRICE({_o},0.05,0.08,0,2,0)", f"ODDFPRICE({_o},0.05,0.08,100,3,0)", f"ODDFPRICE({_o},0.05,0.08,100,2,5)",
            f"ODDFPRICE({_o},0.05,0.08,100,2)", f"ODDFPRICE({_o},0.05,0.08,100)", f"ODDFPRICE({_o},0.05,0.08,100,2,0,1)", f"ODDFPRICE({_o},0.05,0,100,2,0)",
            "ODDFPRICE(DATE(2024,2,15),DATE(2027,7,1),DATE(2024,3,1),DATE(2024,7,1),0.05,0.08,100,2,0)", "ODDFPRICE(DATE(2024,2,15),DATE(2027,7,1),DATE(2024,1,1),DATE(2024,1,20),0.05,0.08,100,2,0)",
            "ODDFPRICE(DATE(2024,2,15),DATE(2024,6,1),DATE(2024,1,1),DATE(2024,7,1),0.05,0.08,100,2,0)", "ODDFPRICE(DATE(2024,2,15),DATE(2027,7,1),DATE(2024,1,1),DATE(2024,2,15),0.05,0.08,100,2,0)",
            "ODDFPRICE(DATE(2024,2,15),DATE(2027,7,1),DATE(2024,1,1),DATE(2024,2,16),0.05,0.08,100,2,0)", "ODDFPRICE(DATE(2024,2,15),DATE(2024,7,1),DATE(2024,1,1),DATE(2024,7,1),0.05,0.08,100,2,0)",
            "ODDFPRICE(DATE(2024,2,15),DATE(2027,7,2),DATE(2024,1,1),DATE(2024,7,1),0.05,0.08,100,2,0)", "ODDFPRICE(DATE(2024,2,15),DATE(2027,7,1),DATE(2024,2,15),DATE(2024,7,1),0.05,0.08,100,2,0)",
            f"ODDFYIELD({_o},0.05,0,100,2,0)", f"ODDFYIELD({_o},0.05,95,0,2,0)", f"ODDFYIELD({_o},-0.05,95,100,2,0)", f"ODDFYIELD({_o},0.05,1000000,100,2,0)", f"ODDFYIELD({_o},0.05,95,100,3,0)",
            "ODDLPRICE(DATE(2024,3,1),DATE(2024,6,30),DATE(2024,1,15),-0.05,0.08,100,2,0)", "ODDLPRICE(DATE(2024,3,1),DATE(2024,6,30),DATE(2024,1,15),0.05,-0.08,100,2,0)", "ODDLPRICE(DATE(2024,3,1),DATE(2024,6,30),DATE(2024,1,15),0.05,0.08,0,2,0)",
            "ODDLPRICE(DATE(2024,3,1),DATE(2024,6,30),DATE(2024,1,15),0.05,0.08,100,3,0)", "ODDLPRICE(DATE(2024,3,1),DATE(2024,6,30),DATE(2024,1,15),0.05,0.08,100,2,5)", "ODDLPRICE(DATE(2024,3,1),DATE(2024,6,30),DATE(2024,1,15),0.05,0.08,100,2)",
            "ODDLPRICE(DATE(2024,3,1),DATE(2024,6,30),DATE(2024,1,15),0.05,0.08,100)", "ODDLPRICE(DATE(2024,3,1),DATE(2024,6,30),DATE(2024,3,1),0.05,0.08,100,2,0)", "ODDLPRICE(DATE(2024,3,1),DATE(2024,6,30),DATE(2024,4,1),0.05,0.08,100,2,0)",
            "ODDLPRICE(DATE(2024,6,30),DATE(2024,6,30),DATE(2024,1,15),0.05,0.08,100,2,0)", "ODDLPRICE(DATE(2024,7,30),DATE(2024,6,30),DATE(2024,1,15),0.05,0.08,100,2,0)", "ODDLPRICE(DATE(2024,3,1),DATE(2024,6,30),DATE(2024,1,15),0.05,0,100,2,0)",
            "ODDLYIELD(DATE(2024,3,1),DATE(2024,6,30),DATE(2024,1,15),0.05,0,100,2,0)", "ODDLYIELD(DATE(2024,3,1),DATE(2024,6,30),DATE(2024,1,15),0.05,95,0,2,0)", "ODDLYIELD(DATE(2024,3,1),DATE(2024,6,30),DATE(2024,1,15),-0.05,95,100,2,0)",
            "ODDLYIELD(DATE(2024,3,1),DATE(2024,6,30),DATE(2024,1,15),0.05,95,100,3,0)", "ODDLYIELD(DATE(2024,3,1),DATE(2024,6,30),DATE(2024,1,15),0.05,95,100,2,0,1)"]
FORMULAS += BATCH14

# Byte and width functions (batch 15).
BATCH15 = [
    'LENB("abc")', 'LENB("日本")', 'LENB("é")', 'LENB("")', 'LENB("ｱ")', 'LENB("ア")', 'LENB("ＡＢ")', 'LENB("한국")', 'LENB("中文")', 'LENB("€")', 'LENB("ñ")', 'LENB("¥")', 'LENB("£")', 'LENB("😀")', 'LENB("😀a")',
    'LENB(CHAR(200))', 'LENB(123)', 'LENB(TRUE)', 'LENB(Data!A1)', 'LENB(NA())', 'LENB(Data!D1)', 'LENB("a"&CHAR(10)&"日")', 'LENB("ก")', 'LENB("é日ñ")', 'LENB("ｶﾞ")', 'LENB("①")', 'LENB("ⅷ")', 'LENB("☆")', 'LENB("★")', 'LENB("±×÷")',
    'LEFTB("日本語",1)', 'LEFTB("日本語",2)', 'LEFTB("日本語",3)', 'LEFTB("日本語",4)', 'LEFTB("日本語",5)', 'LEFTB("日本語",6)', 'LEFTB("日本語",7)', 'LEFTB("日本語",-1)', 'LEFTB("日本語")', 'LEFTB("abc",2)', 'LEFTB("ab",0)', 'LEFTB("a日b",2)', 'LEFTB("a日b",3)',
    'LEFTB("a日b",4)', 'LEFTB("日a",1)', 'LEFTB(12345,2)', 'LEFTB("日本語","x")', 'LEFTB("日本語",1.9)', 'LEFTB(Data!D1,1)', 'LEFTB("😀",2)', 'FINDB("b","a😀b")',
    'RIGHTB("日本語",1)', 'RIGHTB("日本語",2)', 'RIGHTB("日本語",3)', 'RIGHTB("日本語",4)', 'RIGHTB("日本語",5)', 'RIGHTB("日本語",7)', 'RIGHTB("a日b",2)', 'RIGHTB("a日b",3)', 'RIGHTB("日a",2)', 'RIGHTB("日本語")', 'RIGHTB("abc")', 'RIGHTB("abc",0)', 'RIGHTB("abc",-1)',
    'MIDB("日本語",1,2)', 'MIDB("日本語",2,2)', 'MIDB("日本語",2,1)', 'MIDB("日本語",2,3)', 'MIDB("日本語",3,1)', 'MIDB("日本語",3,2)', 'MIDB("日本語",4,2)', 'MIDB("日本語",5,10)', 'MIDB("日本語",7,1)', 'MIDB("日本語",8,1)', 'MIDB("日本語",0,2)', 'MIDB("日本語",1,0)',
    'MIDB("日本語",1,-1)', 'MIDB("a日b",2,2)', 'MIDB("a日b",3,2)', 'MIDB("a日b",2,1)', 'MIDB("a日b",3,1)', 'MIDB("abc",2,100)', 'MIDB("abcdef",2,3)',
    'FINDB("本","日本語",1)', 'FINDB("本","日本語",2)', 'FINDB("本","日本語",3)', 'FINDB("本","日本語",4)', 'FINDB("語","日本語")', 'FINDB("","日本語")', 'FINDB("","日本語",3)', 'FINDB("x","日本語")', 'FINDB("a","日a本a",3)', 'FINDB("a","日a本a",5)', 'FINDB("a","日a本a",6)',
    'FINDB("a","日a本a",7)', 'FINDB("a","日a本a",0)', 'FINDB("日","日本語",2)', 'FINDB("b","abc")', 'FINDB("B","abc")', 'FINDB("","abc")', 'FINDB("c","abc",4)', 'FINDB("c","abc",3)',
    'SEARCHB("本","日本語",2)', 'SEARCHB("?","日本語")', 'SEARCHB("*語","日本語")', 'SEARCHB("本?","日本語")', 'SEARCHB("B","a日b")', 'SEARCHB("","日本語")', 'SEARCHB("a","日a本a",3)', 'SEARCHB("~?","日?")', 'SEARCHB("B","abc")', 'SEARCHB("语","日本語")',
    'REPLACEB("日本語",2,2,"x")', 'REPLACEB("日本語",1,1,"x")', 'REPLACEB("日本語",3,4,"x")', 'REPLACEB("日本語",7,2,"x")', 'REPLACEB("日本語",8,2,"x")', 'REPLACEB("日本語",1,0,"x")', 'REPLACEB("日本語",0,1,"x")', 'REPLACEB("abc",2,1,"日")', 'REPLACEB("日本語",1,6,"")',
    'REPLACEB("日本語",1,100,"x")', 'REPLACEB("a日b",3,1,"x")', 'REPLACEB("a日b",2,1,"x")', 'REPLACEB("a日b",2,2,"x")', 'REPLACEB("abcdef",2,3,"X")', 'REPLACEB("abc",4,1,"x")', 'REPLACEB("abc",3,1,"x")', 'REPLACEB("abc",3,2,"x")', 'REPLACEB("",1,0,"x")',
    'ASC("ＡＢＣ１２３")', 'ASC("ｱｲｳ")', 'ASC("アイウ")', 'ASC("ａ　ｂ")', 'ASC("ガ")', 'ASC("abc")', 'ASC("")', 'ASC("￥")', 'ASC(123)', 'ASC(NA())', 'ASC(Data!D1)', 'ASC(TRUE)', 'ASC("ｳﾞ")', 'ASC("ヴ")', 'ASC("　")', 'ASC("～")', 'ASC("￠")', 'ASC("―")', 'ASC("゛")',
    'ASC("パピプペポ")', 'ASC("ー")', 'ASC("日本語ＡＢＣ")', 'ASC("「」、。・")', 'ASC("ヷ")', 'ASC("ヺ")', 'ASC("ヰ")',
    'JIS("ABC123")', 'JIS("ｱｲｳ")', 'JIS("ｶﾞ")', 'JIS("a b")', 'JIS("abc")', 'JIS("")', 'JIS("ﾊﾟ")', 'JIS("¥")', 'JIS(123)', 'JIS(NA())', 'JIS("ｳﾞ")', 'JIS("ﾞ")', 'JIS("ｱﾞ")', 'JIS("ﾊﾞﾊﾟ")', 'JIS(" ")', 'JIS("~")', 'JIS("ｦﾞ")', 'JIS("ﾜﾞ")', 'JIS("ﾊﾞﾟ")',
    'JIS("ｶﾟ")', 'JIS("ｱｲｳｴｵ")', 'JIS("ｶｷｸｹｺ")', 'JIS("ﾊﾋﾌﾍﾎ")', 'JIS("ﾏﾐﾑﾒﾓ")', 'JIS("ｰ｡｢｣､･")', 'JIS("ｧｨｩｪｫｬｭｮｯ")', 'JIS("日本語abc")', 'JIS("!")', 'JIS("`")', 'JIS("é")',
]
FORMULAS += BATCH15

# BAHTTEXT and REGEX (batch 16).
def _q(t):
    return '"' + t.replace('"', '""') + '"'
BATCH16 = [f"_xlfn.BAHTTEXT({v})" for v in [
    0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 19, 20, 21, 22, 29, 30, 31, 41, 51, 61, 71, 81, 91, 99, 100, 101, 102, 110, 111, 112, 120, 121, 200, 201, 210, 211, 1000, 1001, 1010, 1011, 1021, 1100, 1101, 1111, 2000, 2001, 9999,
    10000, 10001, 10011, 10101, 11000, 100000, 100001, 100011, 110000, 111111, 200000, 999999, 1000000, 1000001, 1000011, 1000021, 1000100, 1010000, 1100000, 1500000, 2000000, 2000001, 10000000, 12345678, 100000000, 123456789,
    1000000000, 1000001000, 1234567890, 1000000000000, 1234567890123, 999999999999999, 1e15, 1e16, 1e18, 1e21, 0.01, 0.02, 0.1, 0.11, 0.21, 0.99, 0.999, 0.995, 0.994, 1.005, 1.234, 1.5, 2.5, 21.21, 100.01, 1000.1, 1234.56,
    -1, -0.01, -0.5, -21, -1234.56, -0.004, -0.005, 0.004, 0.005, 0.0049, 999999.999, 99.995, 1e-10, 1e300, '"5"', '"x"', '""', 'TRUE', 'FALSE', 'Data!D1', 'Data!A1', 'NA()', 'DATE(2024,1,15)', '1/3', '2/3', '1E+20', '123456.789',
]]
BATCH16 += ["_xlfn.BAHTTEXT()", "_xlfn.BAHTTEXT(1,2)", "_xlfn.BAHTTEXT(Data!A1:A3)"]
_R = "_xlfn.ORG.LIBREOFFICE.REGEX"
for _t, _p, _rep, _fl in [
    ("abc", "b", None, None), ("abc", "B", None, None), ("ABC", "b", None, None), ("abc", "b", "X", None), ("abc", "B", "X", None), ("abc123", "[0-9]+", None, None), ("abc123", "[0-9]+", "X", None), ("abc123abc", "abc", "X", 2),
    ("abc123abc", "abc", "X", "g"), ("abc", "(a)(b)", "$2$1", None), ("abc", "(a)(b)", None, 2), ("abc", "x", None, None), ("abc", "(b)", "[$0]", None), ("abc", "(b)", "[$1]", None), ("abc", "(b)", "[$2]", None),
    ("abc", "x", "[$2]", None), ("abc", "(b)", r"[\$1]", None), ("abc", "b", r"\$1", None), ("abc", "b", "\\", None), ("abc", "(b)", "[$$]", None), ("abc", "b", "$", None), ("abc", "(b)", "$a", None), ("abc", "(b)", "x$", None),
    ("abc", "(b)", "$10", None), ("abc", "(b)", "$0$0", None), ("abc", "(b)", "$1$1", None), ("abcabc", "b", "X", 1), ("abcabc", "b", "X", 2), ("abcabc", "b", "X", 3), ("abcabc", "b", "X", 0), ("abcabc", "b", "X", -1), ("abcabc", "b", None, 2),
    ("abcabc", "b", None, 1), ("abcabc", "b", None, 0), ("abcabc", "b", None, 3), ("abcabc", "b", "X", "g"), ("abcabc", "b", "X", "G"), ("abcabc", "b", "X", "x"), ("abcabc", "b", "X", ""), ("abcabc", "b", "", "g"), ("abcabc", "b", "", None),
    ("abcabc", "b", None, "g"), ("", "a", None, None), ("", "a", "X", None), ("a.c", r"\.", None, None), ("a.c", ".", "X", "g"), ("aaa", "a", "b", "g"),
    ("a1b22", r"\d+", "<$0>", "g"), ("hello world", r"(\w+) (\w+)", "$2 $1", None), ("123", "2", None, None), ("abc", "b", "5", None), ("abc", "[a-b]+", None, None),
    ("abc", "a|c", "X", "g"), ("abc", "(a)|(c)", "[$1$2]", "g"), ("日本語", "本", "X", None), ("abcabc", "(a)(b)(c)", "$3$2$1", "g"), ("abc", "a{1,2}", "X", None), ("aaa", "a{2}", "X", None), ("aaa", "a+?", "X", "g"),
    ("abc", "[^a]", "X", "g"), ("é", "é", "X", None), ("É", "é", "X", None), ("a b", "a b", "X", None), ("xay", "a", "X", 2), ("xay", "a", "X", 1), ("a,b", ",", "$0$0", None), ("abc", "b", "$0$0$0", "g"), ("a1b2c3", r"[a-c]\d", "<$0>", "g"), ("a1b2c3", r"(?:[a-c])(\d)", "$1", "g"), ("aaa", "a", "b", 2), ("aaa", "aa", "b", "g"), ("abab", "(ab)+", "X", None), ("abab", "(ab)+?", "X", None),
    ("foo.bar", r"\.", "_", None), ("foo bar", r"\s", "_", None), ("x", "x", "\\$1", None),
]:
    _args = [_q(_t), _q(_p)]
    if _rep is not None or _fl is not None:
        _args.append(_q(_rep) if _rep is not None else "")
    if _fl is not None:
        _args.append(_q(_fl) if isinstance(_fl, str) else str(_fl))
    BATCH16.append(f"{_R}({','.join(_args)})")
BATCH16 += [f"{_R}(NA(),\"a\")", f"{_R}(\"abc\",NA())", f"{_R}(123,\"2\")", f"{_R}(\"abc\",\"b\",5)", f"{_R}(\"abc\",\"b\",TRUE)", f"{_R}(\"abc\",\"b\",\"X\",TRUE)", f"{_R}(\"abc\",\"b\",\"X\",1.5)", f"{_R}(\"abc\",\"b\",\"X\",\"1\")",
            f"{_R}(\"abc\")", f"{_R}(\"abc\",\"b\",\"X\",1,2)", f"{_R}(Data!B1,\"p+\")", f"{_R}(Data!A1,\"5\")", f"{_R}(Data!D1,\"a\",\"X\")", f"{_R}(Data!B1:B3,\"a\")"]
FORMULAS += BATCH16

# Complex numbers (batch 9).
_IM_INPUTS = ["3+4i", "3-4i", "4i", "-4i", "i", "-i", "3", "0", "1+i", "1-i", "-1-i", "2.5+3.5i", "1E2+1E-2i", "3+4j", "j", "-j", "2j", "1+2k", "+3+4i",
              "3+ 4i", " 3+4i", "3+4I", "3+4i ", "1+i+j", "abc", "", "i3", "3i+4", "(1+i)", "1.5E300+1E300i", "0+0i", "-0-0i", "1/2+i", "1e1+i", "1+1e1i", ".5+.5i",
              "5.", "1,5+2i", "3++4i", "3+-4i", "--3", "3i4", "0.1+0.2i", "-2.5", "1E+2i", "1E-2i", "12345678901234567890+1i", "1+i2", "ii", "3 + 4i", "4i+3",
              "1+0i", "0+1i", "0-1i", "-1+0i", "1e-10+1e-10i", "2+0.0000001i", "TRUE", "1+Ii"]
BATCH9 = []
for _fn in ["IMABS", "IMAGINARY", "IMREAL", "IMARGUMENT", "IMCONJUGATE", "IMSQRT", "IMEXP", "IMLN", "IMLOG10", "IMLOG2", "IMSIN", "IMCOS", "IMTAN", "IMSEC", "IMCSC",
            "IMCOT", "IMSINH", "IMCOSH", "IMSECH", "IMCSCH"]:
    _pre = ""
    for _z in _IM_INPUTS:
        BATCH9.append(f'{_pre}{_fn}("{_z}")')
    BATCH9 += [f"{_pre}{_fn}(5)", f"{_pre}{_fn}(TRUE)", f"{_pre}{_fn}(Data!D1)", f"{_pre}{_fn}(Data!A1)", f"{_pre}{_fn}(Data!B1)", f"{_pre}{_fn}(NA())", f"{_pre}{_fn}(\"3+4i\",1)",
               f"{_pre}{_fn}()", f"{_pre}{_fn}(Data!A1:A3)", f"{_pre}{_fn}(-2.5)", f"{_pre}{_fn}(1E300)", f"{_pre}{_fn}(1E-300)"]
_PAIRS = [("3+4i", "1-2i"), ("1+i", "1-i"), ("3+4i", "3+4i"), ("0", "1+i"), ("1+i", "0"), ("2", "3"), ("4i", "2i"), ("1+i", "1+j"), ("1+i", "abc"), ("1+i", ""),
          ("1E200+1E200i", "1E200+1E200i"), ("i", "i"), ("5", "i"), ("3+4i", "2"), ("1.5+2.5i", "-0.5-1.5i"), ("1E-200+1i", "1E-200+1i"), ("j", "j"), ("1+i", "1")]
for _fn in ["IMSUM", "IMSUB", "IMPRODUCT", "IMDIV"]:
    for _a, _b in _PAIRS:
        BATCH9.append(f'{_fn}("{_a}","{_b}")')
    BATCH9 += [f'{_fn}("1+i")', f'{_fn}("1+i","2+i","3-i")', f'{_fn}(Data!A1:A3)', f'{_fn}(Data!A1:A2,Data!A3)', f'{_fn}(1,2)', f'{_fn}(TRUE,1)', f'{_fn}("1+i",NA())', f'{_fn}("1+i",Data!D1)',
               f'{_fn}("1+i",Data!B1)', f'{_fn}("1+i",Data!B1:B2)', f'{_fn}()']
BATCH9 += ['IMSUM("1+i","2+i","3-i","1","i")', 'IMPRODUCT("1+i","2+i","3-i","1","i")', 'IMSUM(Data!A1:A3,"i")', 'IMPRODUCT(Data!A1:A3,"i")', 'IMPRODUCT(Data!D1:D3)', 'IMSUM(Data!D1:D3)',
           'IMPRODUCT("1+i",Data!D1:D3)', 'IMSUM("1+i",Data!D1:D3)', 'IMSUM(Data!B1:B3)', 'IMPRODUCT(Data!M1:M4)', 'IMSUM(Data!M1:M4)', 'IMSUM(Data!K1:K3,Data!K4)', 'IMPRODUCT({"1+i","2";"i","3"})', 'IMSUM({"1+i","2";"i","3"})']
for _z in ["3+4i", "1+i", "i", "2", "0", "-1", "1-i", "4i", "abc", "", "1+2k"]:
    for _n in ["2", "0", "1", "-1", "0.5", "-0.5", "3", "1.5", "10", "TRUE", "\"2\"", "\"x\"", "Data!D1", "NA()", "1E10", "-3", "0.1", "Data!A1:A2", "2.5", "100"]:
        BATCH9.append(f'IMPOWER("{_z}",{_n})')
for _re, _im, _sfx in [("3", "4", None), ("3", "4", '"i"'), ("3", "4", '"j"'), ("3", "4", '"k"'), ("3", "4", '"I"'), ("3", "4", '"J"'), ("3", "4", '""'), ("0", "1", None), ("1", "0", None), ("0", "0", None),
                       ("-1", "-1", None), ("1.5", "-2.5", None), ("1E20", "1E-20", None), ("3", "1", None), ("3", "-1", None), ("0", "-1", None), ("0", "5", None), ("5", "0", '"j"'), ("0", "1", '"j"'),
                       ('"3"', "4", None), ("3", '"4"', None), ('"x"', "4", None), ("3", '"x"', None), ("TRUE", "1", None), ("3", "4", "1"), ("3", "4", '"ij"'), ("3", "4", "Data!D1"), ("Data!D1", "Data!D1", None),
                       ("NA()", "1", None), ("1", "NA()", None), ("3", "4", "NA()"), ("0.1", "0.2", None), ("1E-300", "1E300", None), ("1E15", "1", None), ("123456789.123456789", "1", None),
                       ("1", "1", '" "'), ("1", "1", '"i "'), ("1", "-1", '"j"'), ("-0", "-0", None)]:
    BATCH9.append(f"COMPLEX({_re},{_im}" + (f",{_sfx}" if _sfx else "") + ")")
BATCH9 += ["COMPLEX(1)", "COMPLEX()", "COMPLEX(1,2,\"i\",1)", "COMPLEX(Data!A1:A2,1)", "COMPLEX(1,Data!A1:A2)"]
FORMULAS += BATCH9

# A seeded set of complex numbers over every function (the digits of the results depend on how LibreOffice computes them).
import random as _random
_rng = _random.Random(20240615)
def _rand_z():
    kind = _rng.choice(["both", "both", "both", "real", "imag"])
    def num():
        sc = _rng.choice([1, 10, 100, 0.01, 1000])
        return round(_rng.uniform(-1, 1) * sc, _rng.choice([0, 1, 2, 3, 4]))
    a, b = num(), num()
    sfx = _rng.choice(["i", "i", "i", "j"])
    if kind == "real":
        return f"{a}"
    if kind == "imag":
        return f"{b}{sfx}"
    return f"{a}{'+' if b >= 0 else ''}{b}{sfx}"
_ZS = [_rand_z() for _ in range(60)]
BATCH17 = []
for _fn in ["IMABS", "IMAGINARY", "IMREAL", "IMARGUMENT", "IMCONJUGATE", "IMSQRT", "IMEXP", "IMLN", "IMLOG10", "IMLOG2", "IMSIN", "IMCOS", "IMTAN", "IMSEC", "IMCSC", "IMCOT", "IMSINH", "IMCOSH", "IMSECH", "IMCSCH"]:
    BATCH17 += [f'{_fn}("{_z}")' for _z in _ZS]
for _i in range(0, 60, 2):
    BATCH17 += [f'IMSUM("{_ZS[_i]}","{_ZS[_i + 1]}")', f'IMSUB("{_ZS[_i]}","{_ZS[_i + 1]}")', f'IMPRODUCT("{_ZS[_i]}","{_ZS[_i + 1]}")', f'IMDIV("{_ZS[_i]}","{_ZS[_i + 1]}")']
for _z in _ZS[:30]:
    BATCH17 += [f'IMPOWER("{_z}",{_n})' for _n in (2, 3, 0.5, -1, 1.5, 7, -2.5, 0.25)]
BATCH17 += [f"COMPLEX({_rng.choice([-1, 1]) * round(_rng.uniform(0, 100), 3)},{_rng.choice([-1, 1]) * round(_rng.uniform(0, 100), 3)})" for _ in range(30)]
FORMULAS += BATCH17

BATCH18 = [
    "BESSELI(1.5,1)", "BESSELI(1.5,0)", "BESSELI(-2.5,3)", "BESSELI(30,2)", "BESSELI(1.5,1.9)", "BESSELI(1.5,-1)",
    "BESSELI(0,0)", "BESSELI(0,1)", "BESSELI(1,-0.5)", "BESSELI(2,3000000000)", "BESSELI(1,200)", "BESSELI(5,2000)",
    "BESSELI(700,1)", "BESSELI(800,1)", "BESSELI(1000,2)", "BESSELI(10000,0)", "BESSELI(-1.5,3)", "BESSELI(TRUE,1)",
    'BESSELI("a",1)', 'BESSELI(2,"x")', "BESSELJ(1.9,2)", "BESSELJ(0,0)", "BESSELJ(0,1)", "BESSELJ(0,0.5)", "BESSELJ(10,0)",
    "BESSELJ(-3,1)", "BESSELJ(-1.5,2)", "BESSELJ(-1.5,3)", "BESSELJ(100,5)", "BESSELJ(1.5,-1)", "BESSELJ(1,-0.5)",
    "BESSELJ(1E5,1)", "BESSELJ(30,25)", "BESSELJ(1000,1)", "BESSELJ(500,2)", "BESSELJ(5000000,2)", "BESSELJ(10000000,2)",
    "BESSELJ(10000000,1000)", "BESSELJ(2,2147483647)", "BESSELJ(2,2147483648)", "BESSELJ(1,1000)", "BESSELJ(1E-300,2)",
    'BESSELJ(1,"1")', "BESSELJ(2,TRUE)", "BESSELK(1.5,1)", "BESSELK(0.1,0)", "BESSELK(0,1)", "BESSELK(-1,1)", "BESSELK(10,3)",
    "BESSELK(800,1)", "BESSELK(1E-5,2)", "BESSELK(2,100)", "BESSELK(1,-0.5)", "BESSELK(2,3)", "BESSELK(2,2)", "BESSELK(2.5,0)",
    "BESSELK(1E-300,0)", "BESSELK(1E-300,3)", "BESSELY(1.5,1)", "BESSELY(0.1,0)", "BESSELY(0,1)", "BESSELY(-1,1)",
    "BESSELY(10,3)", "BESSELY(1000,2)", "BESSELY(2.5,0)", "BESSELY(1E-5,2)", "BESSELY(2,100)", "BESSELY(1,-0.5)",
    "BESSELY(10000000,1)", "BESSELY(6000000,0)", "BESSELY(2E19,0)", "BESSELY(1,50)", "BESSELY(2,3)",
]
for _fn in ["BESSELI", "BESSELJ", "BESSELK", "BESSELY"]:
    for _ in range(40):
        BATCH18.append(f"{_fn}({round(_rng.uniform(0.01, 60), _rng.choice([0, 1, 3]))},{_rng.choice([0, 1, 2, 3, 4, 7, 12])})")
FORMULAS += BATCH18

_CLASSES = {
    "mass": ["g", "kg", "mg", "sg", "lbm", "u", "ozm", "stone", "ton", "grain", "pweight", "hweight", "shweight", "brton", "cwt", "uk_cwt", "lcwt", "uk_ton", "LTON", "Gg", "ug"],
    "length": ["m", "km", "cm", "mm", "mi", "Nmi", "in", "ft", "yd", "ang", "Pica", "picapt", "pica", "ell", "parsec", "pc", "lightyear", "ly", "survey_mi", "nm", "um", "Mpc", "dam", "hm"],
    "time": ["yr", "day", "d", "hr", "mn", "min", "sec", "s", "ms", "us", "ns", "ks"],
    "pressure": ["Pa", "kPa", "atm", "at", "mmHg", "Torr", "psi", "hPa", "mPa", "GPa"],
    "force": ["N", "kN", "dyn", "dy", "lbf", "pond", "mN", "mdyn"],
    "energy": ["J", "kJ", "e", "c", "cal", "kcal", "eV", "ev", "keV", "HPh", "hh", "Wh", "kWh", "wh", "flb", "BTU", "btu", "mJ", "uJ", "Mcal"],
    "power": ["W", "kW", "w", "HP", "h", "PS", "mW", "MW", "GW"],
    "magnetism": ["T", "ga", "mT", "uT", "kga", "nT"],
    "temperature": ["C", "cel", "F", "fah", "K", "kel", "Reau", "Rank", "mK", "kK", "uK"],
    "volume": ["tsp", "tbs", "oz", "cup", "pt", "us_pt", "uk_pt", "qt", "gal", "l", "L", "lt", "m3", "mi3", "Nmi3", "in3", "ft3", "yd3", "ang3", "Pica3", "picapt3", "pica3", "barrel", "bushel", "regton", "GRT", "Schooner", "Middy", "Glass", "Sixpack", "Humpen", "ly3", "MTON", "tspm", "uk_gal", "uk_qt", "ml", "cl", "dl", "kl", "cm3", "mm3", "km3", "dam3", "hl", "m^3"],
    "area": ["m2", "mi2", "Nmi2", "in2", "ft2", "yd2", "ang2", "Pica2", "picapt2", "pica2", "Morgen", "ar", "acre", "uk_acre", "us_acre", "ly2", "ha", "cm2", "km2", "mm2", "m^2", "dam2", "kar"],
    "speed": ["m/s", "m/sec", "m/h", "m/hr", "mph", "kn", "admkn", "ludicrous speed", "ridiculous speed", "km/h", "cm/s", "mm/s", "km/s"],
    "information": ["bit", "byte", "kbit", "Mbit", "Gbit", "kibit", "Mibit", "Gibit", "Tibit", "kibyte", "Mibyte", "Gibyte", "kbyte", "Mbyte", "Gbyte", "Tbyte", "Pibyte", "Eibyte", "Zibyte", "Yibyte", "mbit", "mibit", "Kibit", "kiBit"],
}
_UNITS = [u for us in _CLASSES.values() for u in us]
BATCH19 = ['CONVERT(1,"m","km")', 'CONVERT(1,"km","m")', 'CONVERT(1,"m","mi")', 'CONVERT(100,"C","F")', 'CONVERT(32,"F","C")', 'CONVERT(0,"C","K")', 'CONVERT(1,"m","kg")',
    'CONVERT(1,"x","m")', 'CONVERT(1,"m","x")', 'CONVERT(1,"","m")', 'CONVERT(1,"M","m")', 'CONVERT(1,"km","KM")', 'CONVERT(1,"cm3","m3")', 'CONVERT(1,"cm^3","m^3")',
    'CONVERT(1,"m^2","cm^2")', 'CONVERT(1,"m^","cm")', 'CONVERT(1,"^m","m")', 'CONVERT(1,"da","m")', 'CONVERT(1,"dam","m")', 'CONVERT(1,"dag","g")', 'CONVERT(1,"kibit","bit")',
    'CONVERT(1,"kibyte","kbyte")', 'CONVERT(1,"kbit","kibit")', 'CONVERT(1,"mibit","bit")', 'CONVERT(1,"kg","lbm")', 'CONVERT(1,"u","g")', 'CONVERT(1,"Tm","m")', 'CONVERT(1E308,"Ym","m")',
    'CONVERT(1E-308,"ym","Ym")', 'CONVERT(1,5,"m")', 'CONVERT("1","m","km")', 'CONVERT("a","m","km")', 'CONVERT(TRUE,"m","km")', 'CONVERT(-40,"C","F")', 'CONVERT(-273.15,"C","K")',
    'CONVERT(1,"Rank","K")', 'CONVERT(1,"Reau","C")', 'CONVERT(1,"kC","C")', 'CONVERT(1,"mK","K")', 'CONVERT(1,"cel","kel")', 'CONVERT(1,"g","g")', 'CONVERT(1,"hh","h")', 'CONVERT(1,"e","c")', 'CONVERT(1,"c","e")']
for _cls, _us in _CLASSES.items():
    for _ in range(40):
        _a, _b = _rng.choice(_us), _rng.choice(_us)
        BATCH19.append(f'CONVERT({round(_rng.uniform(-50, 500), _rng.choice([0, 1, 3]))},"{_a}","{_b}")')
for _ in range(60):
    BATCH19.append(f'CONVERT({round(_rng.uniform(0, 100), 2)},"{_rng.choice(_UNITS)}","{_rng.choice(_UNITS)}")')
FORMULAS += BATCH19

BATCH20 = [
    "CHISQDIST(3,2,1)", "CHISQDIST(3,2,0)", "CHISQDIST(3,2)", "CHISQDIST(-1,2,1)", "CHISQDIST(0,2,0)", "CHISQDIST(0,1,0)", "CHISQDIST(3,0,1)",
    "CHISQDIST(3,1.9,1)", "CHISQDIST(3,2,5)", 'CHISQDIST(3,2,"a")', "CHISQDIST(3,2,TRUE)", "CHISQDIST(3,1E10,1)", "CHISQDIST(3,2E10,1)",
    "CHISQDIST(3,0.5,1)", "CHISQDIST(3,1E10,0)", "CHISQDIST(1E10,1E10,1)", "CHISQDIST(1E10,1E10,0)", "CHISQDIST(1E5,1E5,0)", "CHISQDIST(1E5,1E5,1)",
    "CHISQDIST(1E6,1E6,0)", "CHISQDIST(1E3,1E3,0)", "CHISQDIST(1E7,2E7,1)", "CHISQDIST(0.5,1,0)", "CHISQDIST(1,1,0)", "CHISQDIST(100,3,0)",
    "CHISQDIST(1000,3,0)", "CHISQDIST(1500,3,0)", "CHISQDIST(1400,3,0)", "CHISQDIST(-0.0001,3,0)", "CHISQINV(0.5,2)", "CHISQINV(0,2)", "CHISQINV(1,2)",
    "CHISQINV(1.1,2)", "CHISQINV(-0.1,2)", "CHISQINV(0.5,0)", "CHISQINV(0.5,1.5)", "CHISQINV(0.99999,3)", "CHISQINV(1E-10,3)", "CHISQINV(0.5,1E10)",
    "CHISQINV(0.5,2E10)", "CHISQINV(0.5,1E6)", "CHISQINV(0.001,50)", "CHISQINV(0.999,50)", "CHISQINV(0.9999999,1)", 'CHISQINV(0.5,"a")', "CHISQINV(0.5,2.999)",
]
for _ in range(40):
    BATCH20.append(f"CHISQDIST({round(_rng.uniform(0, 80), 2)},{_rng.choice([1, 2, 3, 5, 10, 30, 100])},{_rng.choice([0, 1])})")
    BATCH20.append(f"CHISQINV({round(_rng.uniform(0.001, 0.999), 3)},{_rng.choice([1, 2, 3, 5, 10, 30, 100])})")
FORMULAS += BATCH20

BATCH21 = []
_O = "_xlfn.ORG.LIBREOFFICE."
for _ty in [0, 1, 2, 3, 10, 11, 12, 13, 14, 15, 16, 17, 18, 21, 22, 1.9, -1, "1", "x"]:
    for _d in ["DATE(2024,1,7)", "DATE(2024,1,1)", "DATE(2024,12,30)", "DATE(2021,1,1)", "DATE(2023,1,1)", "DATE(2022,12,31)", "DATE(2026,12,31)", "DATE(1900,1,1)", "DATE(1900,3,1)"]:
        BATCH21.append(f"{_O}WEEKNUM_OOO({_d},{_ty if not isinstance(_ty, str) or _ty == 'x' and False else repr(_ty).replace(chr(39), chr(34))})")
BATCH21 += [f"{_O}WEEKNUM_OOO(DATE(2024,1,7))", f"{_O}WEEKNUM_OOO(\"x\",1)", f"{_O}WEEKNUM_OOO(-5,1)", f"{_O}WEEKNUM_OOO(DATE(2024,1,7),)",
            f"{_O}RANDBETWEEN.NV(5,5)", f"{_O}RANDBETWEEN.NV(5,4)", f"{_O}RANDBETWEEN.NV(1.2,1.9)", f"{_O}RANDBETWEEN.NV(-3,-3)",
            f"{_O}RANDBETWEEN.NV(1,3)<=3", f"{_O}RAND.NV()<1", f"{_O}RAND.NV()>=0", f"{_O}RANDBETWEEN.NV(\"a\",2)"]
for _i in range(60):
    BATCH21.append(f"{_O}WEEKNUM_OOO(DATE({_rng.randint(1900, 2100)},{_rng.randint(1, 12)},{_rng.randint(1, 28)}),{_rng.choice([1, 2, 11, 12, 13, 14, 15, 16, 17, 21])})")
FORMULAS += BATCH21
BATCH22 = []
for _n in ["RANDBETWEEN", _O + "RANDBETWEEN.NV"]:
    for _lim in ["1.2,1.9", "1.5,1.5", "-1.5,-1.5", "-1.2,-1.9", "0.2,0.9", "0.2,0.4", "1.9,1.2", "-1.9,-1.2", "2.5,2.4", "0,0", "-0.5,-0.5", "7,7", "1E3,1E3", "1.2,1.2", "TRUE,TRUE", "-1,-2"]:
        BATCH22.append(f"{_n}({_lim})")
    for _lim in ["1.2,2.9", "2.5,3.5", "-0.5,0.5", "1,1.99", "0.2,0.9", "-7,3.5"]:
        BATCH22.append(f"AND({_n}({_lim})>=INT({_lim.split(',')[0]}),{_n}({_lim})<=INT({_lim.split(',')[1]})+1)")
FORMULAS += BATCH22

BATCH23 = [f"{_O}COLOR(255,0,0)", f"{_O}COLOR(0,255,0)", f"{_O}COLOR(0,0,255)", f"{_O}COLOR(1,2,3)", f"{_O}COLOR(1,2,3,4)", f"{_O}COLOR(255,255,255,255)",
    f"{_O}COLOR(256,0,0)", f"{_O}COLOR(0,256,0)", f"{_O}COLOR(0,0,256)", f"{_O}COLOR(0,0,0,256)", f"{_O}COLOR(-1,0,0)", f"{_O}COLOR(-0.5,0,0)", f"{_O}COLOR(1.9,2.9,3.9)",
    f"{_O}COLOR(255.9,0,0)", f"{_O}COLOR(\"1\",2,3)", f"{_O}COLOR(\"x\",2,3)", f"{_O}COLOR(TRUE,FALSE,TRUE)", f"{_O}COLOR(1/0,2,3)", f"{_O}COLOR(1/0,NA(),3)", f"{_O}COLOR(300,NA(),1/0)",
    f"{_O}COLOR(300,0,0,NA())", f"{_O}COLOR(1,2)", "CURRENT()", "CURRENT()+1", "ISERROR(CURRENT())", "IFERROR(CURRENT(),7)",
    'EUROCONVERT(1,"DEM","EUR")', 'EUROCONVERT(1,"EUR","DEM")', 'EUROCONVERT(100,"FRF","DEM")', 'EUROCONVERT(100,"DEM","FRF")', 'EUROCONVERT(100,"DEM","FRF",TRUE)',
    'EUROCONVERT(100,"DEM","FRF",FALSE,3)', 'EUROCONVERT(100,"DEM","FRF",TRUE,3)', 'EUROCONVERT(100,"DEM","FRF",TRUE,2)', 'EUROCONVERT(100,"DEM","FRF",TRUE,5.9)', 'EUROCONVERT(100,"DEM","FRF",,4)',
    'EUROCONVERT(1,"eur","dem")', 'EUROCONVERT(1,"Dem","Dem")', 'EUROCONVERT(1.234,"DEM","DEM")', 'EUROCONVERT(1.234,"EUR","EUR")', 'EUROCONVERT(1.2345,"EUR","BEF")', 'EUROCONVERT(1.2345,"EUR","ITL")',
    'EUROCONVERT(1,"EUR","USD")', 'EUROCONVERT(1,"USD","EUR")', 'EUROCONVERT(1,"","EUR")', 'EUROCONVERT(1,"EUR","")', 'EUROCONVERT("a","EUR","DEM")', 'EUROCONVERT(1,"EUR","DEM","x")',
    'EUROCONVERT(1,"EUR","DEM",1)', 'EUROCONVERT(1,"EUR","DEM",0)', 'EUROCONVERT(1,5,"DEM")', 'EUROCONVERT(-1.005,"EUR","DEM")', 'EUROCONVERT(1E10,"ITL","BEF")', 'EUROCONVERT(1E300,"EUR","ITL")',
    'EUROCONVERT(1/0,"EUR","DEM",TRUE,NA())', 'EUROCONVERT(NA(),"EUR","DEM",1/0)', 'EUROCONVERT(1,NA(),"DEM",TRUE,1/0)', 'EUROCONVERT(1,"EUR",NA(),1/0)', 'EUROCONVERT(1,"EUR","DEM",NA(),2)', 'EUROCONVERT(1,"EUR","XXX",NA())',
    'EUROCONVERT(1,"EUR","DEM",1/0)', 'EUROCONVERT(1/0,"EUR","DEM")', 'EUROCONVERT(1,"EUR","DEM",TRUE,NA())']
for _ in range(60):
    _c = ["EUR", "ATS", "BEF", "DEM", "ESP", "FIM", "FRF", "IEP", "ITL", "LUF", "NLG", "PTE", "GRD", "SIT", "MTL", "CYP", "SKK", "EEK", "LVL", "LTL", "HRK", "BGN"]
    BATCH23.append(f'EUROCONVERT({round(_rng.uniform(-100, 10000), _rng.choice([0, 2, 4]))},"{_rng.choice(_c)}","{_rng.choice(_c)}",{_rng.choice(["TRUE", "FALSE"])},{_rng.choice([3, 4, 6, 10])})')
    BATCH23.append(f'EUROCONVERT({round(_rng.uniform(-100, 10000), _rng.choice([0, 2, 4]))},"{_rng.choice(_c)}","{_rng.choice(_c)}")')
FORMULAS += BATCH23

BATCH24 = []
_F = _O + "FOURIER"
def _fo(rng_formula, *opts):
    return f"{_F}({rng_formula}{''.join(',' + o for o in opts)})"
_inputs = ["Data!A1:A8", "Data!A1:A5", "Data!A1:A3", "Data!A1:A2", "Data!A1:A1", "Data!A1:A4", "Data!A1:A6", "Data!A1:A7", "{1;2;3;4}", "{1;2;3}", "{1;2;3;4;5}", "{1;2;3;4;5;6;7;8;9;10;11;12}",
    "{1,2;3,4}", "{1,2;3,4;5,6}", "{1,2;3,4;5,6;7,8}", "{1,2;3,4;5,6;7,8;9,10;11,12;13,14;15,16}", "{1,2,3;4,5,6}", "{1,2,3,4}", "{1,2,3,4,5,6,7,8}", "{1,2,3,4,5,6}", "{1,2,3,4,5,6,7}", "{1,2,3,4,5,6;7,8,9,10,11,12}",
    "{1,2,3,4,5,6,7,8;9,10,11,12,13,14,15,16}", "5", "TRUE", "{1;TRUE;3;4}", "{0.5;-0.25;1E3;7}", "{1E10;1E-10;3;4;5}", "Data!C1:C4"]
for _in0 in _inputs:
    for _by in ["TRUE", "FALSE"]:
        # LibreOffice transposes an array written in the formula, in the formula, when it takes the
        # series by rows (see FIDELITY.md): a computed array stands for it.
        _in = f"TRANSPOSE({_in0})" if _by == "FALSE" and _in0.startswith("{") else _in0
        BATCH24.append(f"ROWS({_fo(_in, _by)})")
        BATCH24.append(f"COLUMNS({_fo(_in, _by)})")
        for _extra in [(), ("TRUE",), ("FALSE", "TRUE"), ("TRUE", "TRUE"), ("FALSE", "TRUE", "1"), ("TRUE", "TRUE", "5")]:
            _f = _fo(_in, _by, *_extra)
            BATCH24.append(f"SUM({_f})")
            BATCH24.append(f"INDEX({_f},1,1)")
            BATCH24.append(f"INDEX({_f},2,1)")
            BATCH24.append(f"INDEX({_f},2,2)")
            BATCH24.append(f"INDEX({_f},3,2)")
            BATCH24.append(f"INDEX({_f},ROWS({_f}),1)")
            BATCH24.append(f"INDEX({_f},ROWS({_f}),2)")
BATCH24 += [_fo("Data!Z1:Z4", "TRUE"), "SUM(" + _fo("Data!Z1:Z4", "TRUE") + ")", "SUM(" + _fo("Data!B1:B4", "TRUE") + ")", "SUM(" + _fo('{1;"a";3;4}', "TRUE") + ")",
    "SUM(" + _fo("1/0", "TRUE") + ")", "SUM(" + _fo("{1;2;#N/A;4}", "TRUE") + ")", "SUM(" + _fo('{1;2;""}', "TRUE") + ")", "SUM(" + _fo("{1;2;3}", "TRUE", "FALSE", "TRUE", '"x"') + ")",
    "SUM(" + _fo("{1;2;3}", "TRUE", '"x"') + ")", "SUM(" + _fo("{1;2;3}", "TRUE", '"1"') + ")", "SUM(" + _fo("{1;2;3}", "TRUE", "FALSE", '"x"') + ")", "SUM(" + _fo("{1;2;3}", '"x"') + ")",
    "SUM(" + _fo("{1;2;3}", '"1"') + ")", "SUM(" + _fo("{1;2;3}", "TRUE", "1/0") + ")", "SUM(" + _fo("{1;2;3}", "1/0") + ")", "SUM(" + _fo("{1;2;3}", "TRUE", "FALSE", "TRUE", "1/0") + ")",
    "SUM(" + _fo("{1;2;3}", "TRUE", "NA()", "TRUE", "1/0") + ")", "SUM(" + _fo("1/0", "TRUE", "NA()") + ")", "SUM(" + _fo("{1;2;3}", "TRUE", "FALSE", "TRUE", "-1") + ")",
    "SUM(" + _fo("{1;2;3}", "TRUE", "FALSE", "TRUE", "1.5") + ")", "SUM(" + _fo("{1;2;3}", "TRUE", "FALSE", "TRUE", "2") + ")", "SUM(" + _fo("Data!A1:A3", "2") + ")",
    "SUM(" + _fo("Data!A1:C3", "TRUE") + ")", "SUM(" + _fo("Data!A1:A5", "FALSE") + ")", "SUM(" + _fo("Data!A1:C1", "FALSE") + ")"]
for _ in range(25):
    _n = _rng.choice([1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 12, 16, 17, 31, 32])
    _arr = ";".join(str(round(_rng.uniform(-100, 100), _rng.choice([0, 2, 5]))) for _ in range(_n))
    _arr2 = ",".join(str(round(_rng.uniform(-100, 100), _rng.choice([0, 2, 5]))) for _ in range(_n))
    _opts = _rng.choice([(), ("FALSE", "TRUE"), ("TRUE",), ("TRUE", "TRUE"), ("FALSE", "TRUE", "1"), ("TRUE", "TRUE", "10")])
    BATCH24.append(f"SUM({_fo('{' + _arr + '}', 'TRUE', *_opts)})")
    BATCH24.append(f"INDEX({_fo('{' + _arr + '}', 'TRUE', *_opts)},{_rng.randint(1, _n)},{_rng.randint(1, 2)})")
    BATCH24.append(f"INDEX({_fo('TRANSPOSE({' + _arr + '})', 'FALSE', *_opts)},{_rng.randint(1, _n)},{_rng.randint(1, 2)})")
    _pairs = ";".join(f"{round(_rng.uniform(-100, 100), 2)},{round(_rng.uniform(-100, 100), 2)}" for _ in range(_n))
    BATCH24.append(f"INDEX({_fo('{' + _pairs.replace(';', ';') + '}', 'TRUE', *_opts)},{_rng.randint(1, _n)},{_rng.randint(1, 2)})")
FORMULAS += BATCH24

BATCH25 = []
_greeks = ["", ',"value"', ',"price"', ',"delta"', ',"gamma"', ',"theta"', ',"vega"', ',"volga"', ',"vanna"', ',"rho"', ',"rhof"', ',"d"', ',"g"', ',"e"', ',"o"', ',"a"', ',"f"', ',"x"', ',""', ",1", ",TRUE"]
for _ in range(500):
    _spot = _rng.choice([100, 95, 80, 120, 100.5, 60, 150])
    _vol = _rng.choice([0.2, 0.3, 0.1, 0.5, 0.05])
    _rd = _rng.choice([0.05, 0.0, 0.02, -0.01, 0.1])
    _rf = _rng.choice([0.02, 0.0, 0.05, -0.02])
    _t = _rng.choice([1, 0.5, 2, 0.25, 0, 5, 0.01])
    _k = _rng.choice([100, 90, 110, 0, 80, 120, 100])
    _bl, _bu = _rng.choice([(90, 0), (0, 120), (90, 120), (0, 0), (85, 115), (95, 105), (120, 90), (70, 130)])
    _reb = _rng.choice([0, 0, 2, 5.5])
    _pc, _io, _bc = _rng.choice(['"call"', '"put"', '"c"', '"p"', '"Call"']), _rng.choice(['"in"', '"out"', '"i"', '"o"']), _rng.choice(['"cont"', '"end"', '"c"', '"e"', '"x"'])
    BATCH25.append(f"OPT_BARRIER({_spot},{_vol},{_rd},{_rf},{_t},{_k},{_bl},{_bu},{_reb},{_pc},{_io},{_bc}{_rng.choice(_greeks)})")
for _ in range(400):
    _spot = _rng.choice([100, 95, 80, 120, 100.5, 60, 150])
    _vol = _rng.choice([0.2, 0.3, 0.1, 0.5, 0.05])
    _rd = _rng.choice([0.05, 0.0, 0.02, -0.01, 0.1])
    _rf = _rng.choice([0.02, 0.0, 0.05, -0.02])
    _t = _rng.choice([1, 0.5, 2, 0.25, 0, 5, 0.01])
    _bl, _bu = _rng.choice([(90, 0), (0, 120), (90, 120), (0, 0), (85, 115), (95, 105), (120, 90), (70, 130)])
    _fd, _io, _bc = _rng.choice(['"for"', '"dom"', '"f"', '"d"', '"x"']), _rng.choice(['"in"', '"out"', '"i"', '"o"']), _rng.choice(['"cont"', '"end"', '"c"', '"e"'])
    BATCH25.append(f"OPT_TOUCH({_spot},{_vol},{_rd},{_rf},{_t},{_bl},{_bu},{_fd},{_io},{_bc}{_rng.choice(_greeks)})")
for _ in range(150):
    _spot = _rng.choice([100, 95, 80, 120, 60, 150])
    _vol = _rng.choice([0.2, 0.3, 0.1, 0.5])
    _mu = _rng.choice([0.05, 0.0, -0.02, 0.1])
    _t = _rng.choice([1, 0.5, 2, 0.25, 0, 5])
    _bl, _bu = _rng.choice([(90, 0), (0, 120), (90, 120), (0, 0), (85, 115), (120, 90), (70, 130)])
    BATCH25.append(f"OPT_PROB_HIT({_spot},{_vol},{_mu},{_t},{_bl},{_bu})")
    _k = _rng.choice([100, 90, 110, 0, -1, 80])
    _pcx = _rng.choice(['', ',"call"', ',"put"', ',"c"', ',"p"', ',"x"', ',1', ',""'])
    _kx = _rng.choice(['', f",{_k}", f",{_k}", ',"a"', ",TRUE"])
    BATCH25.append(f"OPT_PROB_INMONEY({_spot},{_vol},{_mu},{_t},{_bl},{_bu}{_kx if _kx else ',' if _pcx else ''}{_pcx})")
BATCH25 += ["OPT_BARRIER(0,0.2,0.05,0.02,1,100,90,120,0,\"call\",\"in\",\"cont\")", "OPT_BARRIER(100,0,0.05,0.02,1,100,90,120,0,\"call\",\"in\",\"cont\")", "OPT_BARRIER(100,0.2,0.05,0.02,-1,100,90,120,0,\"call\",\"in\",\"cont\")",
    "OPT_BARRIER(100,0.2,0.05,0.02,1,-100,90,120,0,\"call\",\"in\",\"cont\")", "OPT_BARRIER(\"x\",0.2,0.05,0.02,1,100,90,120,0,\"call\",\"in\",\"cont\")", "OPT_BARRIER(1/0,0.2,0.05,0.02,1,100,90,120,0,\"call\",NA(),\"cont\")",
    "OPT_BARRIER(100,0.2,0.05,0.02,1,100,90,120,0,1,\"in\",\"cont\")", "OPT_BARRIER(100,0.2,0.05,0.02,1,100,90,120,0,\"call\",\"in\")", "OPT_PROB_HIT(100,0.2,0.05,1,90)",
    "OPT_PROB_INMONEY(100,0.2,0.05,1,90,120,100,\"put\")", "OPT_PROB_INMONEY(100,0.2,0.05,1,90,120,,\"put\")", "OPT_PROB_INMONEY(100,0.2,0.05,1,90,120,100,)", "OPT_PROB_INMONEY(100,0.2,0.05,1,90,120)",
    "OPT_PROB_INMONEY(100,0.2,0.05,1,90,120,\"c\")", "OPT_PROB_INMONEY(100,0.2,0.05,1,90,120,\"c\",\"c\")"]
FORMULAS += BATCH25

BATCH26 = []
import datetime as _dt
def _serial(y, m, d):
    return (_dt.date(y, m, d) - _dt.date(1899, 12, 30)).days
def _fmt(v):
    return str(int(v)) if float(v) == int(v) else repr(round(v, 4))
def _arr(vs):
    return "{" + ";".join(_fmt(v) for v in vs) + "}"
def _series(n):
    kind = _rng.choice(["trend", "add", "mult", "flat", "rand"])
    p = _rng.choice([2, 3, 4, 5, 6, 7])
    if kind in ("add", "mult"):
        n = p * _rng.randint(2, 5)
    a, b = _rng.uniform(5, 50), _rng.uniform(-2, 5)
    season = [_rng.uniform(-8, 8) for _ in range(p)]
    factor = [_rng.uniform(0.6, 1.5) for _ in range(p)]
    ys = []
    for t in range(n):
        noise = _rng.uniform(-1.5, 1.5)
        if kind == "trend":
            ys.append(a + b * t + noise)
        elif kind == "add":
            ys.append(a + b * t + season[t % p] + noise)
        elif kind == "mult":
            ys.append(max(1, (a + b * t)) * factor[t % p] + noise)
        elif kind == "flat":
            ys.append(a)
        else:
            ys.append(_rng.uniform(1, 100))
    ys = [round(y, _rng.choice([0, 1, 2])) for y in ys]
    return ys, p
def _timeline(n):
    kind = _rng.choice(["unit", "unit", "step", "daily", "monthly", "month31", "frac"])
    if kind == "unit":
        return [i + 1 for i in range(n)], 1
    if kind == "step":
        st = _rng.choice([2, 5, 7, 10])
        return [st * i + 3 for i in range(n)], st
    if kind == "frac":
        return [0.5 * i + 1 for i in range(n)], 0.5
    if kind == "daily":
        return [45000 + i for i in range(n)], 1
    day = 31 if kind == "month31" else _rng.choice([1, 15, 28])
    out = []
    y, m = 2022, _rng.randint(1, 12)
    for i in range(n):
        d = day
        while True:
            try:
                out.append(_serial(y, m, d))
                break
            except ValueError:
                d -= 1
        m += 1
        if m > 12:
            m = 1
            y += 1
    return out, 30
for _ in range(700):
    _ys, _p = _series(_rng.randint(4, 24))
    _xs, _st = _timeline(len(_ys))
    _mode = _rng.choice(["plain", "plain", "plain", "gap", "dup", "shuffle"])
    _pairs = list(zip(_xs, _ys))
    if _mode == "gap" and len(_pairs) > 6:
        del _pairs[_rng.randint(2, len(_pairs) - 3)]
    elif _mode == "dup":
        _k = _rng.randint(1, len(_pairs) - 1)
        _pairs.insert(_k, (_pairs[_k - 1][0], round(_rng.uniform(1, 50), 1)))
    elif _mode == "shuffle":
        _rng.shuffle(_pairs)
    _xa, _ya = _arr([p[0] for p in _pairs]), _arr([p[1] for p in _pairs])
    _last = max(p[0] for p in _pairs)
    _targets = [_last + _st * _rng.choice([1, 2, 3, 0.5, 1.5, 4]) for _ in range(_rng.choice([1, 1, 2, 3]))]
    _tgt = _fmt(_targets[0]) if len(_targets) == 1 and _rng.random() < 0.6 else _arr(_targets) if _rng.random() < 0.5 else "{" + ",".join(_fmt(v) for v in _targets) + "}"
    _inside = _rng.random() < 0.15
    if _inside:
        _tgt = _fmt(_pairs[_rng.randint(0, len(_pairs) - 1)][0] + _st * _rng.choice([0, 0, 0.5, 0.25]))
    _opts = _rng.choice([[], [_rng.choice([0, 1, 2, 3, 4, _p])], [_rng.choice([0, 1, _p]), _rng.choice([0, 1])], [_rng.choice([1, _p]), _rng.choice([0, 1]), _rng.randint(1, 7)], [_rng.choice([1, _p, 0]), "", _rng.randint(1, 7)]])
    _o = "".join("," + str(x) for x in _opts)
    _fn = _rng.choice(["_xlfn.FORECAST.ETS", "_xlfn.ORG.LIBREOFFICE.FORECAST.ETS.MULT"])
    BATCH26.append(f"{_fn}({_tgt},{_ya},{_xa}{_o})")
    _fn2 = _rng.choice(["_xlfn.FORECAST.ETS.STAT", "_xlfn.ORG.LIBREOFFICE.FORECAST.ETS.STAT.MULT"])
    _types = _arr(_rng.sample(range(1, 10), _rng.choice([1, 2, 9])))
    BATCH26.append(f"{_fn2}({_ya},{_xa},{_types}{_o})")
    _so = "".join("," + str(x) for x in _rng.choice([[], [_rng.choice([0, 1])], [_rng.choice([0, 1]), _rng.randint(1, 7)]]))
    BATCH26.append(f"_xlfn.FORECAST.ETS.SEASONALITY({_ya},{_xa}{_so})")
    if _rng.random() < 0.3 and _st != 30:
        _fn3 = _rng.choice(["_xlfn.FORECAST.ETS.CONFINT", "_xlfn.ORG.LIBREOFFICE.FORECAST.ETS.PI.MULT"])
        _lv = _rng.choice(["0.95", "0.9", "0.5", "0.99", "", "1", "0"])
        _tail = _rng.choice([f",{_lv},0", f",{_lv},0,{_rng.choice([0, 1])}", f",{_lv},0,1,{_rng.randint(1, 7)}"])
        _tg = _fmt(_last + _st * _rng.choice([1, 2, 3])) if _rng.random() < 0.5 else _arr([_last + _st * k for k in (1, 2, 3)])
        BATCH26.append(f"{_fn3}({_tg},{_ya},{_xa}{_tail})")
BATCH26 += [
    "_xlfn.FORECAST.ETS(9,Data!A1:A8,Data!A1:A8)", "_xlfn.FORECAST.ETS.SEASONALITY(Data!A1:A8,Data!A1:A8)", "_xlfn.FORECAST.ETS(9,{1;2},{1;2})", "_xlfn.FORECAST.ETS(3,{1},{1})",
    "_xlfn.FORECAST.ETS(3,{1;2;3},{1;2})", '_xlfn.FORECAST.ETS(3,{1;"a";3},{1;2;3})', "_xlfn.FORECAST.ETS(0,{1;2;3;4},{1;2;3;4})", "_xlfn.FORECAST.ETS(1,{1;2;3;4},{1;2;3;4})",
    "_xlfn.FORECAST.ETS(2.5,{1;2;3;4},{1;2;3;4})", "_xlfn.FORECAST.ETS(7,{1;2;3;4},{1;2;3;4},-1)", "_xlfn.FORECAST.ETS(7,{1;2;3;4},{1;2;3;4},1.5)", "_xlfn.FORECAST.ETS(7,{1;2;3;4},{1;2;3;4},1,2)",
    "_xlfn.FORECAST.ETS(7,{1;2;3;4},{1;2;3;4},1,1,0)", "_xlfn.FORECAST.ETS(7,{1;2;3;4},{1;2;3;4},1,1,8)", "_xlfn.FORECAST.ETS(7,{1;2;3;4},{1;2;3;#DIV/0!},1,1,8)", "_xlfn.FORECAST.ETS(7,{1;2;3;4},{1;2;3;4},1/0)",
    "_xlfn.FORECAST.ETS(7,{1;2;3;4},{1;2;4;8})", "_xlfn.FORECAST.ETS(7,{1;2;3;4},{1;3;5;8})", "_xlfn.FORECAST.ETS(7,{1;1;3;4},{1;1;2;3},1,1,0)", "_xlfn.FORECAST.ETS.SEASONALITY({1;2;3},{1;2;3})",
    "_xlfn.FORECAST.ETS.STAT({1;2;3;4},{1;2;3;4},0)", "_xlfn.FORECAST.ETS.STAT({1;2;3;4},{1;2;3;4},10)", "_xlfn.FORECAST.ETS.STAT({1;2;3;4},{1;2;3;4},9.9)", "_xlfn.FORECAST.ETS.STAT({1;2;3;4},{1;2;3;4},{1,2;3,4})",
    "_xlfn.FORECAST.ETS.CONFINT(7,{1;2;3;5},{1;2;3;4},0.95,0)", "_xlfn.FORECAST.ETS.CONFINT(7,{1;2;3;5},{1;2;3;4},1.1,0)", "_xlfn.FORECAST.ETS.CONFINT(7,{1;2;3;5},{1;2;3;4},-0.1,0)",
    "_xlfn.FORECAST.ETS.CONFINT(3,{1;2;3;5},{1;2;3;4},0.95,0)", "_xlfn.FORECAST.ETS.CONFINT(4,{1;2;3;5},{1;2;3;4},0.95,0)", "_xlfn.FORECAST.ETS.CONFINT(5,{1;2;3;5},{1;2;3;4},0.95,0)",
]
FORMULAS += BATCH26

BATCH27 = ["ATAN(1)", "ATAN(-0.5)", "ATAN(1E300)", "COS(1)", "COS(0)", "COS(1E10)", "TAN(1)", "TAN(0.5)", "TAN(1E15)", "RADIANS(180)", "RADIANS(-45.5)", "RADIANS(1E300)",
    "ISERR(NA())", "ISERR(1/0)", "ISERR(1)", 'ISNONTEXT("a")', "ISNONTEXT(1)", "ISNONTEXT(Data!Z9)", "ISREF(Data!A1)", "ISREF(1)", "ISREF(Data!A1:B2)", "ISREF(INDEX(Data!A1:B2,1,1))",
    'AVERAGEIFS(Data!A1:A8,Data!B1:B8,"apple")', 'AVERAGEIFS(Data!A1:A8,Data!B1:B8,"apple",Data!A1:A8,">3")', 'AVERAGEIFS(Data!A1:A8,Data!B1:B8,"zzz")',
    "ERROR.TYPE(NA())", "ERROR.TYPE(1/0)", "ERROR.TYPE(1)", '_xlfn.NUMBERVALUE("1.234,5",",",".")', '_xlfn.NUMBERVALUE("12%")', '_xlfn.NUMBERVALUE("x")', "_xlfn.SKEW.P(Data!A1:A8)",
    "_xlfn.SKEW.P({1,2,3})", "_xlfn.HYPGEOM.DIST(2,5,5,10,TRUE)", "_xlfn.HYPGEOM.DIST(2,5,5,10,FALSE)", "HYPGEOMDIST(2,5,5,10)"]
for _ in range(40):
    _v = round(_rng.uniform(-1000, 1000), _rng.choice([0, 2, 6]))
    BATCH27 += [f"ATAN({_v})", f"COS({_v})", f"TAN({_v})", f"RADIANS({_v})", f"DEGREES({_v})", f"SIN({_v})"]
FORMULAS += BATCH27

BATCH28 = []
_refs = ["Data!B3", "Data!B3:C5", "Data!A1", "Data!H8", "Data!Z9", "Data!1:1", "Data!A:A", "Data!C:D", "Calc!B3", "Calc!A1:A5", "A3", "$B$7", "Data!$A$2", "(Data!B3)", "INDEX(Data!A1:A8,2)", "OFFSET(Data!A1,2,1)",
         "Data!C1", "Data!E2", "IF(1,Data!B2)", "Data!AA100", "Calc!XFD5", "Data!A1:Z9"]
for _r in _refs:
    for _t in ["address", "ADDRESS", "col", "row", "sheet", "coord", "contents", "type", "Contents", "ROW"]:
        BATCH28.append(f'CELL("{_t}",{_r})')
BATCH28 += ['CELL("address")', 'CELL("row")', 'CELL("col")', 'CELL("sheet")', 'CELL("coord")', 'CELL("type")', 'CELL("contents")', 'CELL("addres",Data!B3)', 'CELL("a",Data!B3)', 'CELL("",Data!B3)',
    'CELL(1,Data!B3)', 'CELL(TRUE,Data!A1)', 'CELL(NA(),Data!A1)', 'CELL("x",1)', 'CELL("x",Data!A1)', 'CELL("row",NA())', 'CELL("row",1)', 'CELL("row","a")', 'CELL("contents",1+1)', 'CELL("contents",NA())',
    'CELL("type",1)', 'CELL("type",TRUE)', 'CELL("type",NA())', 'CELL("type",Calc!A1)', 'CELL("type",Calc!A2)', 'CELL("type",Calc!A5)', 'CELL("type",Calc!A40)', 'CELL("type",Calc!A300)', 'CELL("contents",Calc!A300)',
    'CELL("row",Calc!A300:B400)', 'CELL("address",1/0)', 'CELL("type",1/0)', 'CELL("type",A1/0)', 'CELL("contents",Calc!A24297)', 'CELL("type",Calc!A24297)',
    'INFO("recalc")', 'INFO("RECALC")', 'INFO("numfile")', 'INFO("origin")', 'INFO("memavail")', 'INFO("totmem")', 'INFO("memused")', 'INFO("x")', 'INFO("")', 'INFO(1)', 'INFO(NA())', 'INFO("sys")', 'INFO("osversion ")', 'INFO(" recalc")', 'INFO(1/0)']
FORMULAS += BATCH28
FORMULAS += ['STYLE("Good")', 'STYLE("Default")', 'STYLE("NoSuchStyle")', 'STYLE("Good",2,"Default")', 'STYLE("Good","x")', 'STYLE("Good",-1)', 'STYLE("Good",0)', 'STYLE("Good",1,"Nope")', 'STYLE(1)',
    'STYLE("")', 'STYLE("Good",1,"")', 'STYLE(1/0)', 'STYLE("Good",1/0)', 'STYLE("Good",1,1/0)', 'STYLE(NA(),1/0)', 'STYLE("Good",TRUE)', 'STYLE(Data!B1)', 'STYLE("Good",Data!A1)', 'STYLE("Good",Data!B1)']


def build(path, formulas=None):
    wb = openpyxl.Workbook()
    data = wb.active
    data.title = "Data"
    nums = [5, 3, 8, 1, 9, 2, 7, 4]
    labels = ["apple", "banana", "cherry", "apple", "date", "banana", "apple", "fig"]
    for i, (n, t) in enumerate(zip(nums, labels), start=1):
        data[f"A{i}"] = n
        data[f"B{i}"] = t
    base = date(2024, 1, 15)
    for i in range(8):
        data[f"C{i + 1}"] = date(2024, 1 + i, 15 + i)
    data["E1"], data["F1"], data["G1"] = "x", "y", "z"
    data["E2"], data["F2"], data["G2"] = 10, 20, 30
    for i, v in enumerate([1, 3, 5, 7, 9], start=1):
        data[f"H{i}"] = v
    for i, v in enumerate([9, 7, 5, 3, 1], start=1):
        data[f"I{i}"] = v
    for i, v in enumerate([-100, 30, 40, 50, 60], start=1):
        data[f"J{i}"] = v
    for i, v in enumerate([2, 4, 5, 4, 5, 7, 8, 9], start=1):
        data[f"K{i}"] = v
    # Rows 7 and 8 are hidden; L5 is a subtotal inside the range L1:L6.
    for rr in (7, 8):
        data.row_dimensions[rr].hidden = True
    for i, v in enumerate([10, 20, 30, 40], start=1):
        data[f"L{i}"] = v
    data["L5"] = "=SUBTOTAL(9,L1:L4)"
    data["L6"] = 5
    # A column with an error and a text, for the functions that take arrays with them.
    data["M1"], data["M2"], data["M3"], data["M4"] = 1, "=NA()", 3, "a"
    # The database of the D functions and its criteria blocks (see BATCH7).
    people = [("Ann", 30, 160, 80), ("Bob", 25, 175, 90), ("Cid", 30, 180, 70), ("Dee", 41, 165, 85),
              ("Eve", 25, 170, 60), ("Fay", 35, 158, 95), ("Bob", 52, 182, 77), ("Gus", "n/a", 170, None)]
    for j, h in enumerate(["Name", "Age", "Height", "Score"]):
        data.cell(1, 15 + j, h)
    for i, row in enumerate(people, start=2):
        for j, v in enumerate(row):
            if v is not None:
                data.cell(i, 15 + j, v)
    def crit(col, header, *rows):
        # A block starts at column `col` (1-based); `header` is a list of names; rows are lists.
        for j, h in enumerate(header):
            data.cell(1, col + j, h)
        for i, r in enumerate(rows, start=2):
            for j, v in enumerate(r):
                if v is not None:
                    data.cell(i, col + j, v)
    crit(20, ["Age"], [30])
    crit(22, ["Age", "Name"], [">25", "B*"])
    crit(25, ["Name"], ["Ann"], ["Bob"])
    crit(27, ["Age", "Name"], [">30", None], ["<26", None])
    crit(30, ["Age"], [None])
    crit(32, ["Nope"], [1])
    crit(34, ["Score", "Score"], [">=80", "<=90"])
    crit(37, ["Name"], ["="])
    crit(39, ["Age"], ["=25"])
    crit(41, ["Name"], ["bob"])
    crit(43, ["Name"], ["*e"])
    crit(45, ["Name"], ["<>Bob"])
    crit(47, ["Name"], [None], [None])
    crit(49, ["Score"], [">0"])
    crit(51, ["Age", "Name", "Score"], [">26", None, None], [None, "B*", "<80"])
    for i, (u, v) in enumerate(zip([1, 4, 2, 6, 3, 5, 9, 7], [2, 8, 4, 12, 6, 10, 18, 14]), start=1):
        data.cell(i, 58, u)   # BF
        data.cell(i, 59, v if False else u * 3 % 7 + 1)   # BG, not a multiple of BF
        data.cell(i, 60, v)   # BH = 2 * BF (collinear)
    crit(55, ["Name"], ["Gus"])
    crit(57, ["Name"], ["Eve"])
    calc = wb.create_sheet("Calc")
    for i, f in enumerate(FORMULAS if formulas is None else formulas, start=1):
        calc[f"A{i}"] = "=" + f
        calc[f"B{i}"] = f  # label (as text) for failure messages
        calc[f"B{i}"].data_type = "s"
    wb.save(path)


def _profile(tmp, locale):
    """A LibreOffice user profile with the given system locale (None: the default)."""
    if locale is None:
        return []
    prof = os.path.join(tmp, "profile_" + locale)
    os.makedirs(os.path.join(prof, "user"), exist_ok=True)
    with open(os.path.join(prof, "user", "registrymodifications.xcu"), "w") as f:
        f.write(
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
            'xmlns:xs="http://www.w3.org/2001/XMLSchema" '
            'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            '<item oor:path="/org.openoffice.Setup/L10N"><prop oor:name="ooSetupSystemLocale" '
            f'oor:op="fuse"><value>{locale}</value></prop></item></oor:items>'
        )
    return [f"-env:UserInstallation=file://{prof}"]


def main(only=()):
    """The semantics fixtures; with `only` (locale tags), just those of the locales."""
    with tempfile.TemporaryDirectory() as tmp:
        fixtures = [] if only else [("semantics", None, FORMULAS)]
        fixtures += [(f"semantics_{tag}", tag, locale_formulas(tag)) for tag in LOCALE_FIXTURES if not only or tag in only]
        for name, locale, formulas in fixtures:
            src = os.path.join(tmp, f"{name}.xlsx")
            build(src, formulas)
            outdir = os.path.join(tmp, "out")
            for fmt in ("xlsx", "ods"):
                subprocess.run(
                    [SOFFICE, *_profile(tmp, locale), "--headless", "--convert-to", fmt, "--outdir", outdir, src],
                    check=True, capture_output=True,
                )
                os.replace(os.path.join(outdir, f"{name}.{fmt}"), os.path.join(OUT, f"{name}.{fmt}"))
            print(f"{name}: {len(formulas)} formulas")


# The same engine under a Spanish locale: separators and names of the user's
# regional settings.
ES_FORMULAS = _text_rows() + [
    "1234.5678&\"\"", "0.5&\"\"", "VALUE(\"1,5\")", "VALUE(\"1.5\")", "VALUE(\"1.234,5\")",
    "TEXT(1234.5,\"#.##0,00\")", "FIXED(1234.5678,2)", "FIXED(1234.5678,1,TRUE)",
    "TEXT(45292,\"dddd d mmmm\")", "TEXT(TRUE,\"0\")", "1/3&\"\"", "DATEVALUE(\"15/01/2024\")",
    "TEXT(1234.5678,\"#.##0\")", "_xlfn.NUMBERVALUE(\"1.234,5\",\",\",\".\")", "3>2&\"\"",
    "DOLLAR(1234.567,2)", "DOLLAR(-1234.567,2)", "DOLLAR(1234.567,0)", "DOLLAR(0.5,1)", "DOLLAR(-0.004,2)",
    "TEXT(0.5,\"0%\")", "TEXT(1234567.891,\"#,##0.00\")", "ROUND(1234.5678,2)&\"\"",
] + [f'DATEVALUE("{d}")' for d in _NAMED_DATES_ES]


# The locales with a fixture of their own (besides English of the United States, `semantics`): each
# is evaluated by a LibreOffice running with that regional setting, over rows built from its entry
# of the engine's table. A locale is added to the table (scripts/gen_locale_table.py) only with its
# fixture here and once the engine reproduces every row.
LOCALE_FIXTURES = ["es-ES", "es-CO", "es-MX", "en-GB"]
TABLE = os.path.join(os.path.dirname(__file__), "..", "..", "sheetrs-engine", "src", "locale_data.rs")


def _unescape(t):
    t = re.sub(r"\\u\{([0-9a-f]+)\}", lambda m: chr(int(m.group(1), 16)), t)
    return t.replace('\\"', '"').replace("\\\\", "\\")


def locale_row(tag):
    """The entry of the engine's table (locale_data.rs) for `tag`, as a dict."""
    src = open(TABLE, encoding="utf-8").read()
    block = re.search(r'LocaleRow \{\s*tag: "%s",(.*?)\n    \},' % re.escape(tag), src, re.S).group(1)
    row = {}
    for key in ("months", "months_abbr", "days", "days_abbr", "date_patterns"):
        body = re.search(r"%s: &?\[(.*?)\]," % key, block, re.S).group(1)
        row[key] = [_unescape(x) for x in re.findall(r'"((?:[^"\\]|\\.)*)"', body)]
    row["date_sep"] = _unescape(re.search(r"date_sep: '(.*?)',", block).group(1))
    return row


# Day 13, month 3, year 2024 (25 for the two-digit form): see scripts/gen_locale_table.py.
_NAMED_FORMS = ["M A B", "A M B", "B M A", "M A", "A M", "M B", "B M", "M A b", "A M b"]


def locale_formulas(tag):
    r = locale_row(tag)
    rows = ["TODAY()"]
    q = lambda f: f.replace('"', '""')
    # Dates of digits: every pattern the locale lists and the common ones, with valid and invalid days.
    patterns = list(dict.fromkeys(r["date_patterns"] + ["D/M/Y", "M/D/Y", "Y/M/D", "D.M.Y", "D-M-Y", "M-D-Y", "Y-M-D", "D.M.", "D/M", "M/D", "D-M"]))
    for p in patterns:
        for d, m, y in ((13, 3, 2024), (3, 1, 24), (31, 12, 99), (1, 13, 2024)):
            rows.append('DATEVALUE("%s")' % p.replace("D", str(d)).replace("M", str(m)).replace("Y", str(y)))
    rows += ['DATEVALUE("2024-03-13")', 'DATEVALUE("24-03-13")', 'DATEVALUE("99-12-31")', 'DATEVALUE("31-12-99")',
             'DATEVALUE("2024-3-1 10:00")', 'DATEVALUE("3/1/2024 10:00")', 'DATEVALUE("3.1.2024 10:00")']
    # Dates with a month name: each form, each separator, the name in full, abbreviated and in capitals.
    for m in (2, 8):
        full, abbr = r["months"][m], r["months_abbr"][m]
        for name in dict.fromkeys([full, abbr, abbr.rstrip("."), full.upper(), abbr + "." if not abbr.endswith(".") else abbr, "Sept"]):
            for sep in (" ", "-", "/"):
                for form in _NAMED_FORMS:
                    subs = {"M": name, "A": "13", "B": "2024", "b": "25"}
                    rows.append('DATEVALUE("%s")' % sep.join(subs[t] for t in form.split()))
            for text in (f"13. {name} 2024", f"13. {name}", f"{name} 13, 2024", f"Monday {name} 13 2024", f"{name} 13 2024 {r['days'][1]}"):
                rows.append('DATEVALUE("%s")' % text)
    rows += [f'DATEVALUE("13 {r["months"][2]} 2024")']
    # Numbers and times as text.
    rows += ['VALUE("%s")' % v for v in ("1,5", "1.5", "1.234,5", "1,234.5", "1 234,5", "1\'234.5", "12%", "13:30", "1:30 PM", "1,234", "1.234")]
    # TEXT with the number and date formats, the language tags included.
    for f in ("0", "0.00", "#,##0", "#,##0.00", "0.0%", "0.00E+00", "General", "[$$-409]#,##0.00", "[$€-407]#,##0.00",
              "#,##0;(#,##0)", "#.##0,00", "0,00", "# ?/?", "00000", "#,##0.00;[Red]-#,##0.00"):
        for v in ("1234.5678", "-1234.5678", "0.5"):
            rows.append(f'TEXT({v},"{q(f)}")')
    # Language tags: only the languages of the table are read (`[$-407]`, German, is not: LibreOffice
    # writes German names and the engine keeps the locale's).
    for f in ("dddd d mmmm yyyy", "ddd mmm", "mmmm", "dddd", "h:mm AM/PM", "dd/mm/yyyy", "d-mmm-yy", "mmmmm", "dd.mm.yyyy",
              "[$-409]dddd d mmmm", "[$-C0A]dddd d mmmm", "[$-40A]dddd", "[$-809]ddd mmm", "[$-409]h:mm AM/PM"):
        for v in ("45367.75", "45658.5", "45292.25", "45000.9"):
            rows.append(f'TEXT({v},"{q(f)}")')
    rows += ["DOLLAR(1234.567,2)", "DOLLAR(-1234.567,2)", "DOLLAR(0.5,1)", "DOLLAR(-0.004,2)", "DOLLAR(1234567.891,0)",
             "FIXED(1234.5678,2)", "FIXED(1234.5678,1,TRUE)", "FIXED(-1234.5678,0)", '1234.5678&""', '0.5&""', '1/3&""',
             '(1=1)&""', '(1=0)&""', '3>2&""', "1234567.891&\"\"", "1E+20&\"\"",
             '"ä"<"z"', '"ñ"<"o"', '"ñ"<"n"', '"é"<"f"', '"ß"="ss"']
    # Spanish keeps the rows the first fixtures of this locale had (the whole TEXT matrix, named dates).
    return ES_FORMULAS + rows if tag.startswith("es-") else rows


def build_loan(path):
    """A small, realistic model: loan inputs → payment, totals, schedule."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Loan"
    ws["A1"], ws["B1"] = "Principal", 200000
    ws["A2"], ws["B2"] = "Annual rate", 0.05
    ws["A3"], ws["B3"] = "Years", 30
    ws["A4"], ws["B4"] = "Monthly payment", "=PMT(B2/12,B3*12,-B1)"
    ws["A5"], ws["B5"] = "Total paid", "=B4*B3*12"
    ws["A6"], ws["B6"] = "Total interest", "=B5-B1"
    ws["A7"], ws["B7"] = "Month to inspect", 12
    ws["A8"], ws["B8"] = "Balance at that month", '=INDIRECT("D"&(10+B7))'
    ws["A9"], ws["B9"], ws["C9"], ws["D9"] = "Schedule", "Payment", "Interest", "Balance"
    ws["A10"], ws["D10"] = 0, "=B1"
    for m in range(1, 361):
        r = 10 + m
        ws[f"A{r}"] = m
        ws[f"B{r}"] = "=$B$4"
        ws[f"C{r}"] = f"=D{r - 1}*$B$2/12"
        ws[f"D{r}"] = f"=D{r - 1}+C{r}-B{r}"
    ws["A372"], ws["B372"] = "Unused note", "=SUM(B11:B370)*0+LEN(A1)"
    # The settings sheet declares the inputs, outputs and metadata: `key` is
    # the JSON name, `value` the range (or the literal of a metadata field).
    st = wb.create_sheet("settings")
    st.append(["key", "value", "description", "meta_type"])
    for key, ref, desc in [("principal", "Loan!B1", "Loan amount"), ("rate", "Loan!B2", "Annual rate"),
                           ("years", "Loan!B3", "Term in years"), ("month", "Loan!B7", "Month to inspect")]:
        st.append([key, ref, desc, "input"])
    for key, ref, desc in [("payment", "Loan!B4", "Monthly payment"), ("interest", "Loan!B6", "Total interest"),
                           ("balance", "Loan!B8", "Balance at that month")]:
        st.append([key, ref, desc, "output"])
    st.append(["version", "1.0", "Model version", "info"])
    wb.save(path)


def main_loan():
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "loan.xlsx")
        build_loan(src)
        outdir = os.path.join(tmp, "out")
        for fmt in ("xlsx", "ods"):
            subprocess.run(
                [SOFFICE, "--headless", "--convert-to", fmt, "--outdir", outdir, src],
                check=True, capture_output=True,
            )
            os.replace(os.path.join(outdir, f"loan.{fmt}"), os.path.join(OUT, f"loan.{fmt}"))


def main_macro():
    """loan.xlsx plus an (empty) VBA project: what the reader takes for macros."""
    import zipfile
    src, dst = os.path.join(OUT, "loan.xlsx"), os.path.join(OUT, "loan_macro.xlsx")
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("xl/vbaProject.bin", b"\xd0\xcf\x11\xe0")


if __name__ == "__main__":
    import sys

    if sys.argv[1:]:
        # `gen_semantics_fixture.py es-ES de-DE`: only the fixtures of those locales
        main(sys.argv[1:])
    else:
        main()
        main_loan()
        main_macro()
